import os
import io
import re
import uuid
import time
import json
import hmac
import hashlib
import inspect
import threading
import urllib.request
import urllib.parse
import xml.etree.ElementTree as ET
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, date
from zoneinfo import ZoneInfo
import streamlit as st
import pandas as pd
from pypdf import PdfReader
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from google import genai
from google.genai import types

# PDF 직접 생성을 위한 ReportLab 모듈
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# 1. 웹 페이지 기본 설정
st.set_page_config(
    page_title="세무·법령·회계 자율 에이전트",
    page_icon="⚖️",
    layout="wide"
)

# 백그라운드 스레드 및 하위 모듈(law_api, agent) 임포트 전에 st.secrets의 모든 API 키를 환경변수로 최우선 동기화
try:
    if hasattr(st, "secrets"):
        for k, v in st.secrets.items():
            if isinstance(v, (str, int, float)) and k not in os.environ:
                os.environ[k] = str(v).strip()
except Exception:
    pass

# 🔒 ImportError 원천 방지 및 환경변수 동기화 후 안전 임포트
import agent
import law_api
from agent import run_agent

# ---------------------------------------------------------
# 📁 전역 저장소 경로 최우선 선언 (NameError 원천 차단)
# ---------------------------------------------------------
ROOM_IDS = [f"대화방 {i}" for i in range(1, 11)]
STORAGE_DIR = ".workspace_storage"
WORKSPACE_STORAGE_DIR = STORAGE_DIR
os.makedirs(STORAGE_DIR, exist_ok=True)
EXPIRY_FILE_PATH = os.path.join(STORAGE_DIR, "api_expiry_dates.json")
ROOM_LAW_ACK_PATH = os.path.join(STORAGE_DIR, "room_law_ack.json")
MANAGER_AUTH_FILE = os.path.join(STORAGE_DIR, "_manager_google_auth.json")
ANON_WS_PREFIX = "__anon_session__:"

# ---------------------------------------------------------
# 💎 프리미엄 UI/UX CSS (좌측 사이드바 겹침 원천 해결 + 옵션 C 오로라 로딩 LED + 우측 본문 15.5px 가독성 최적화)
# ---------------------------------------------------------
st.markdown("""
<style>
@import url('https://cdn.jsdelivr.net/gh/orioncactus/pretendard/dist/web/static/pretendard.css');

html {
    font-size: 13.8px !important;
}
html, body, [class*="css"] {
    font-family: 'Pretendard', 'Malgun Gothic', 'Apple SD Gothic Neo', sans-serif;
}

/* 📐 우측 본문 상단 여백 확보 (상단 상태바가 Streamlit 헤더에 가려지지 않도록 3.2rem 보장) */
div.block-container {
    padding-top: 3.2rem !important;
    padding-bottom: 2.0rem !important;
}

/* 🎨 사이드바 전체 폭(300px) 및 여백 */
section[data-testid="stSidebar"] {
    width: 300px !important;
    min-width: 300px !important;
    background-color: #f1f5f9 !important;
    border-right: 1px solid #cbd5e1 !important;
}
section[data-testid="stSidebar"] div[data-testid="stSidebarUserContent"] {
    padding-top: 0.75rem !important;
    padding-bottom: 0.6rem !important;
}

/* 📦 사이드바 내부 화이트 카드 패딩 및 간격 (겹침 방지 여백 최적화) */
section[data-testid="stSidebar"] div[data-testid="stVerticalBlockBorderWrapper"] {
    background-color: #ffffff !important;
    border: 1px solid #e2e8f0 !important;
    border-radius: 10px !important;
    padding: 11px 12px !important;
    margin-bottom: 8px !important;
    box-shadow: 0 1px 2px rgba(15, 23, 42, 0.03) !important;
}

section[data-testid="stSidebar"] div[data-testid="stVerticalBlock"] {
    gap: 0.48rem !important;
}
/* Streamlit 기본 마크다운 음수 마진으로 인해 소제목과 아래 토글/입력창이 겹치는 현상 원천 차단 */
section[data-testid="stSidebar"] div[data-testid="stMarkdown"],
section[data-testid="stSidebar"] div[data-testid="stMarkdownContainer"] {
    overflow: visible !important;
    margin-bottom: 0px !important;
}

/* 사이드바 카드 소제목 스타일 (아래 위젯과 8px 안전 간격 확보) */
.sb-card-title {
    font-size: 0.88rem;
    font-weight: 800;
    color: #0f172a;
    line-height: 1.45 !important;
    margin-top: 1px !important;
    margin-bottom: 8px !important;
    padding-bottom: 2px !important;
    display: flex;
    align-items: center;
    justify-content: space-between;
    letter-spacing: -0.3px;
}
.sb-badge-green {
    background-color: #dcfce7;
    color: #15803d;
    font-size: 0.72rem;
    font-weight: 700;
    padding: 2px 7px;
    border-radius: 999px;
}
.sb-badge-amber {
    background-color: #fef3c7;
    color: #b45309;
    font-size: 0.72rem;
    font-weight: 700;
    padding: 2px 7px;
    border-radius: 999px;
}
.sb-badge-pro {
    background-color: #dbeafe;
    color: #1d4ed8;
    font-size: 0.72rem;
    font-weight: 700;
    padding: 2px 7px;
    border-radius: 999px;
}

/* 🟢 사이드바 최상단 모드 전환 버튼 (자율 에이전트AI / 통합고용세액공제) */
.st-key-btn_top_mode_agent button,
.st-key-btn_top_mode_credit button {
    border-radius: 10px !important;
    min-height: 40px !important;
    height: 40px !important;
    padding: 6px 10px !important;
    font-size: 0.90rem !important;
    font-weight: 800 !important;
    letter-spacing: -0.3px !important;
    box-shadow: 0 2px 5px rgba(15, 23, 42, 0.06) !important;
    transition: all 0.15s ease-in-out !important;
}
div.stButton > button[kind="primary"],
div.stButton > button[data-testid="stBaseButton-primary"],
div.stfFormSubmitButton > button[kind="primary"],
div[data-testid="stFormSubmitButton"] > button[kind="primary"],
.st-key-btn_top_mode_agent button[kind="primary"],
.st-key-btn_top_mode_credit button[kind="primary"] {
    background-color: #16a34a !important;
    border: 1.5px solid #15803d !important;
    color: #ffffff !important;
    font-weight: 800 !important;
}
div.stButton > button[kind="primary"]:hover,
div.stButton > button[data-testid="stBaseButton-primary"]:hover,
div[data-testid="stFormSubmitButton"] > button[kind="primary"]:hover {
    background-color: #15803d !important;
    border-color: #166534 !important;
    color: #ffffff !important;
}
div.stButton > button[kind="primary"] p,
div.stButton > button[data-testid="stBaseButton-primary"] p,
div[data-testid="stFormSubmitButton"] > button[kind="primary"] p {
    color: #ffffff !important;
    font-weight: 800 !important;
}

/* 🗑️ 입체감이 살아있는 [현재 방 비우기] 전용 3D 버튼 및 하단 간격 분리 */
.st-key-btn_clear_current_room {
    margin-top: 2px !important;
    margin-bottom: 10px !important;
    padding-bottom: 10px !important;
    border-bottom: 1px dashed #cbd5e1 !important;
}
.st-key-btn_clear_current_room button {
    background: linear-gradient(180deg, #ffffff 0%, #f1f5f9 100%) !important;
    border: 1px solid #94a3b8 !important;
    border-bottom: 2.5px solid #64748b !important;
    border-radius: 8px !important;
    min-height: 2.25rem !important;
    font-size: 0.86rem !important;
    font-weight: 700 !important;
    color: #1e293b !important;
    box-shadow: 0 2px 4px rgba(15, 23, 42, 0.08) !important;
    transition: all 0.12s ease !important;
}
.st-key-btn_clear_current_room button:hover {
    background: linear-gradient(180deg, #fff1f2 0%, #ffe4e6 100%) !important;
    border-color: #f43f5e !important;
    border-bottom-color: #e11d48 !important;
    color: #be123c !important;
    transform: translateY(-1px) !important;
}
.st-key-btn_clear_current_room button:active {
    transform: translateY(1px) !important;
    border-bottom-width: 1px !important;
}

/* 📋 1~10번 대화방 세로 리스트 (왼쪽 텍스트 정렬 + 오른쪽 끝 동그라미 LED 고정) */
div[class*="st-key-room_btn_"] button {
    position: relative !important;
    display: flex !important;
    justify-content: flex-start !important;
    align-items: center !important;
    text-align: left !important;
    width: 100% !important;
    min-height: 2.1rem !important;
    height: 2.1rem !important;
    padding: 0 28px 0 11px !important;
    border-radius: 7px !important;
    font-size: 0.90rem !important;
    font-weight: 600 !important;
    background-color: #ffffff !important;
    border: 1px solid #e2e8f0 !important;
    color: #475569 !important;
    box-shadow: none !important;
}
div[class*="st-key-room_btn_"] button > div,
div[class*="st-key-room_btn_"] button p {
    display: flex !important;
    justify-content: flex-start !important;
    text-align: left !important;
    width: 100% !important;
    margin: 0 !important;
    overflow: hidden !important;
    text-overflow: ellipsis !important;
    white-space: nowrap !important;
    font-size: 0.90rem !important;
}
div[class*="st-key-room_btn_"] button:hover {
    background-color: #f8fafc !important;
    border-color: #94a3b8 !important;
    color: #0f172a !important;
}

/* 상태 1: 빈 방 (우측 빈 동그라미 — 작성 중이 아닐 때만 적용) */
div[class*="st-key-room_btn_"][class*="_empty"]:not([class*="_running"]) button::after {
    content: "" !important;
    position: absolute !important;
    right: 10px !important;
    top: 50% !important;
    transform: translateY(-50%);
    width: 8px !important;
    height: 8px !important;
    border-radius: 50% !important;
    border: 1.5px solid #cbd5e1 !important;
    background-color: transparent;
}

/* 상태 2: 대화 기록이 있는 방 (우측 초록불 점등 — 작성 중이 아닐 때만 초록불 적용) */
div[class*="st-key-room_btn_"][class*="_hasmsg"] button {
    color: #0f172a !important;
    font-weight: 700 !important;
}
div[class*="st-key-room_btn_"][class*="_hasmsg"]:not([class*="_running"]) button::after {
    content: "" !important;
    position: absolute !important;
    right: 10px !important;
    top: 50% !important;
    transform: translateY(-50%);
    width: 8px !important;
    height: 8px !important;
    border-radius: 50% !important;
    background-color: #22c55e;
    box-shadow: 0 0 5px rgba(34, 197, 94, 0.55);
}

/* 상태 3: 현재 선택된 방 (칸 전체 연녹색 강조 + 작성 중이 아닐 때만 우측 진초록 LED) */
div[class*="st-key-room_btn_"][class*="_active"] button {
    background-color: #dcfce7 !important;
    border: 1.5px solid #16a34a !important;
    color: #14532d !important;
    font-weight: 800 !important;
}
div[class*="st-key-room_btn_"][class*="_active"] button p {
    color: #14532d !important;
    font-weight: 800 !important;
}
div[class*="st-key-room_btn_"][class*="_active"]:not([class*="_running"]) button::after {
    content: "" !important;
    position: absolute !important;
    right: 10px !important;
    top: 50% !important;
    transform: translateY(-50%);
    width: 9px !important;
    height: 9px !important;
    border-radius: 50% !important;
    background-color: #16a34a;
    box-shadow: 0 0 6px rgba(22, 163, 74, 0.7);
}

/* 상태 4: 🌌 [옵션 C] 작성 중(분석 중)인 방 — 제미나이 오로라 블루·인디고·바이올렛 회전/펄스 LED
   (CSS 표준상 !important 속성이 @keyframes 애니메이션을 덮어쓰지 않도록 :not([class*="_running"]) 분리 완료) */
@keyframes geminiAuroraPulse {
    0% {
        background-color: #2563eb;
        box-shadow: 0 0 0 0 rgba(59, 130, 246, 0.85), 0 0 9px 2px #3b82f6;
        transform: translateY(-50%) scale(0.92);
    }
    33% {
        background-color: #6366f1;
        box-shadow: 0 0 0 4px rgba(99, 102, 241, 0.30), 0 0 12px 3px #6366f1;
        transform: translateY(-50%) scale(1.22);
    }
    66% {
        background-color: #a855f7;
        box-shadow: 0 0 0 3px rgba(217, 70, 239, 0.35), 0 0 12px 3px #ec4899;
        transform: translateY(-50%) scale(1.10);
    }
    100% {
        background-color: #2563eb;
        box-shadow: 0 0 0 0 rgba(59, 130, 246, 0.85), 0 0 9px 2px #3b82f6;
        transform: translateY(-50%) scale(0.92);
    }
}
div[class*="st-key-room_btn_"][class*="_running"] button::after {
    content: "" !important;
    position: absolute !important;
    right: 10px !important;
    top: 50% !important;
    width: 10px !important;
    height: 10px !important;
    border-radius: 50% !important;
    border: none !important;
    background-color: #6366f1;
    animation: geminiAuroraPulse 1.4s infinite ease-in-out !important;
}

/* 상단 슬림 미니 상태바 */
.top-mini-status-bar {
    display: flex;
    justify-content: space-between;
    align-items: center;
    background-color: #f8fafc;
    border: 1px solid #e2e8f0;
    border-left: 4px solid #16a34a;
    border-radius: 8px;
    padding: 7px 14px;
    margin-bottom: 10px;
    font-size: 0.92rem;
    color: #334155;
}

/* 🤖 답변 로봇 아바타 초록색 배경(#22c55e) & 흰색 아이콘 */
[data-testid="chatAvatarIcon-assistant"],
div[data-testid="stChatMessageAvatarAssistant"] {
    background-color: #22c55e !important;
    color: #ffffff !important;
}
[data-testid="chatAvatarIcon-assistant"] svg,
div[data-testid="stChatMessageAvatarAssistant"] svg {
    fill: #ffffff !important;
    color: #ffffff !important;
}

/* 📖 우측 보고서 본문 폰트 가독성 확대 (15.5px / 줄간격 1.68) */
div[data-testid="stChatMessageContent"] p,
div[data-testid="stChatMessageContent"] li {
    font-size: 1.12rem !important;
    line-height: 1.68 !important;
    color: #1e293b !important;
}

/* 보고서 대제목 (### 1. 결론, ### 2. 내용, ### 3. 관련법령) */
div[data-testid="stChatMessageContent"] h3 {
    color: #1e3a8a !important;
    background: linear-gradient(90deg, #f1f5f9 0%, #ffffff 100%);
    border-left: 5px solid #1e3a8a;
    padding: 8px 14px !important;
    border-radius: 0 6px 6px 0;
    margin-top: 20px !important;
    margin-bottom: 10px !important;
    font-size: 1.28rem !important;
    font-weight: 800 !important;
}

/* 보고서 소제목 (#### (1) 상세 법리 등) */
div[data-testid="stChatMessageContent"] h4 {
    color: #1e40af !important;
    margin-top: 14px !important;
    margin-bottom: 6px !important;
    font-size: 1.14rem !important;
    font-weight: 700 !important;
}

/* 3. 관련법령 전용 조밀한 줄간격 컨테이너 */
.related-laws-compact-box ul,
.related-laws-compact-box ol {
    margin-top: 2px !important;
    margin-bottom: 4px !important;
}
.related-laws-compact-box li {
    font-size: 1.04rem !important;
    line-height: 1.42 !important;
    margin-bottom: 2px !important;
}

/* 본문 근거 태그 및 3. 관련법령 조문 직링크 스타일 */
div[data-testid="stChatMessageContent"] a {
    color: #1d4ed8 !important;
    font-weight: 600 !important;
    text-decoration: none !important;
}
div[data-testid="stChatMessageContent"] a:hover {
    text-decoration: underline !important;
    color: #15803d !important;
}

/* 마크다운 표(Table) 엔터프라이즈 네이비 스타일링 */
div[data-testid="stChatMessageContent"] table,
.law-diff-banner-box table,
.api-key-status-table-box table {
    width: 100%;
    border-collapse: collapse;
    margin: 10px 0 14px 0;
    font-size: 0.98rem !important;
    border-radius: 8px;
    overflow: hidden;
    box-shadow: 0 1px 3px rgba(0,0,0,0.05);
}
div[data-testid="stChatMessageContent"] th,
.law-diff-banner-box th,
.api-key-status-table-box th {
    background-color: #1e293b !important;
    color: #ffffff !important;
    font-weight: 700 !important;
    padding: 9px 12px !important;
    border: 1px solid #334155 !important;
    text-align: left !important;
}
div[data-testid="stChatMessageContent"] td,
.law-diff-banner-box td,
.api-key-status-table-box td {
    padding: 8px 12px !important;
    border: 1px solid #cbd5e1 !important;
    color: #1e293b;
    background-color: #ffffff;
    line-height: 1.5 !important;
    vertical-align: middle !important;
}
div[data-testid="stChatMessageContent"] tr:nth-child(even) td,
.law-diff-banner-box tr:nth-child(even) td,
.api-key-status-table-box tr:nth-child(even) td {
    background-color: #f8fafc;
}
.api-key-status-table-box a {
    display: inline-block;
    background-color: #eff6ff;
    color: #1d4ed8 !important;
    border: 1px solid #bfdbfe;
    padding: 3px 10px;
    border-radius: 6px;
    font-weight: 700 !important;
    font-size: 0.85rem !important;
    text-decoration: none !important;
}
.api-key-status-table-box a:hover {
    background-color: #dbeafe;
    color: #1e40af !important;
}

/* 상단 모드 안내 인용구 박스 (> ⚡ / > 🧠) */
div[data-testid="stChatMessageContent"] blockquote {
    background-color: #f8fafc;
    border-left: 4px solid #22c55e;
    padding: 7px 13px;
    border-radius: 6px;
    color: #334155;
    font-size: 0.95rem !important;
    margin-bottom: 10px;
}

footer {visibility: hidden;}
</style>
""", unsafe_allow_html=True)

# ---------------------------------------------------------
# 0. 작업공간 · API 만료일 · 🎯 맥락 연동형 법령 레이더 및 신·구법 비교 배너 엔진
# ---------------------------------------------------------
WATCHED_TAX_LAWS = [
    "국세기본법",
    "법인세법",
    "소득세법",
    "부가가치세법",
    "조세특례제한법",
    "상속세 및 증여세법",
    "지방세법",
    "농어촌특별세법",
    "근로기준법"
]

_CITED_ARTICLE_TAG_RE = re.compile(
    r"\[조문(?:\s*[:：]\s*|\s*)(?=[가-힣A-Za-z0-9「])([^\]]+)\]"
)

def _get_first_env_or_secret(key_names: list[str]) -> str:
    """지정된 키 후보들 중 st.secrets 또는 환경변수에 설정된 첫 번째 실제 값을 반환합니다."""
    for k in key_names:
        try:
            if hasattr(st, "secrets") and k in st.secrets:
                val = str(st.secrets[k]).strip().strip('"').strip("'")
                if val:
                    return val
        except Exception:
            pass
        val = os.getenv(k, "").strip().strip('"').strip("'")
        if val:
            return val
    return ""

def _get_law_oc() -> str:
    """agent.py, law_api.py, config.py, st.secrets, 환경변수에서 법제처 OC 키(LAW_OPEN_API_KEY 최우선)를 자동 추출합니다."""
    if hasattr(agent, "get_law_oc_auto"):
        val = agent.get_law_oc_auto()
        if val:
            return val
    val = _get_first_env_or_secret(["LAW_OPEN_API_KEY", "LAW_API_OC", "OC", "LAW_OC", "LAW_API_KEY", "MOLEG_OC", "DRF_OC"])
    if val:
        return val
    for attr in ["LAW_OPEN_API_KEY", "OC", "LAW_API_OC", "LAW_OC", "LAW_API_KEY", "API_KEY", "USER_OC", "oc"]:
        val = getattr(law_api, attr, None)
        if isinstance(val, str) and val.strip():
            return val.strip()
    return ""

def _fmt_ymd_str(raw_ymd: str) -> str:
    s = re.sub(r"[^0-9]", "", str(raw_ymd or ""))
    if len(s) == 8:
        return f"{s[:4]}-{s[4:6]}-{s[6:8]}"
    return str(raw_ymd or "-")

@st.cache_data(ttl=43200, show_spinner=False)
def _fetch_latest_tax_law_promulgations_cached() -> dict[str, dict]:
    """
    [속도 개선 ①] 9개 주요 세법 및 각 시행령(총 18건)의 최신 공포·시행 정보를
    ThreadPoolExecutor(6워커)로 병렬 조회하여 초기 구동 시간을 8~12초 -> 1~1.5초로 단축합니다.
    """
    today_d = datetime.now(ZoneInfo("Asia/Seoul")).date()
    today_int = int(today_d.strftime("%Y%m%d"))
    target_names = [base_name + suffix for base_name in WATCHED_TAX_LAWS for suffix in ("", " 시행령")]

    def _fetch_one_law(law_name: str):
        try:
            raw_res = law_api.search_law_data(law_name, "law")
            if not raw_res:
                return None
            if isinstance(raw_res, dict):
                raw_res = [raw_res]

            target_clean = law_name.replace(" ", "")
            exact_matches = [
                it for it in raw_res
                if str(it.get("법령명한글", "")).replace(" ", "") == target_clean
            ]
            candidates = exact_matches if exact_matches else raw_res[:1]
            candidates.sort(
                key=lambda x: (str(x.get("공포일자", "")), str(x.get("시행일자", ""))),
                reverse=True
            )
            top = candidates[0]
            prom_raw = str(top.get("공포일자", "")).strip()
            enf_raw = str(top.get("시행일자", "")).strip()
            prom_no = str(top.get("공포번호", "")).strip()
            rev_type = str(top.get("제개정구분명", "개정")).strip()
            law_id = str(top.get("법령ID") or top.get("법령일련번호") or "").strip()
            mst_id = str(top.get("법령일련번호") or top.get("MST") or law_id).strip()

            days_since_prom = 9999
            if len(prom_raw) == 8 and prom_raw.isdigit():
                try:
                    prom_d = datetime.strptime(prom_raw, "%Y%m%d").date()
                    days_since_prom = (today_d - prom_d).days
                except Exception:
                    pass

            is_future_enforcement = bool(len(enf_raw) == 8 and enf_raw.isdigit() and int(enf_raw) > today_int)
            is_recent_promulgation = (0 <= days_since_prom <= 60)

            return target_clean, {
                "law_name": law_name,
                "clean_name": target_clean,
                "prom_date": _fmt_ymd_str(prom_raw),
                "prom_raw": prom_raw,
                "enf_date": _fmt_ymd_str(enf_raw),
                "enf_raw": enf_raw,
                "prom_no": prom_no,
                "rev_type": rev_type,
                "law_id": law_id,
                "mst_id": mst_id,
                "days_since_prom": days_since_prom,
                "is_future_enforcement": is_future_enforcement,
                "is_recent_promulgation": is_recent_promulgation,
                "version_key": f"{prom_raw}_{prom_no}_{enf_raw}"
            }
        except Exception:
            return None

    fetched_dict = {}
    with ThreadPoolExecutor(max_workers=6) as executor:
        futures = {executor.submit(_fetch_one_law, ln): ln for ln in target_names}
        for fut in as_completed(futures):
            res = fut.result()
            if res:
                fetched_dict[res[0]] = res[1]

    results_map = {}
    for ln in target_names:
        tc = ln.replace(" ", "")
        if tc in fetched_dict:
            results_map[tc] = fetched_dict[tc]
    return results_map

def fetch_latest_tax_law_promulgations() -> dict[str, dict]:
    """일시적 네트워크 오류로 빈 결과가 반환될 경우 12시간 동안 빈 딕셔너리가 캐싱되지 않도록 즉시 캐시를 비웁니다."""
    res = _fetch_latest_tax_law_promulgations_cached()
    if not res:
        _fetch_latest_tax_law_promulgations_cached.clear()
    return res

fetch_latest_tax_law_promulgations.clear = _fetch_latest_tax_law_promulgations_cached.clear

@st.cache_data(ttl=43200, show_spinner=False)
def fetch_official_revision_reason_zero_token(
    law_id: str,
    law_name: str,
    mst_id: str = "",
    enf_raw: str = ""
) -> str:
    oc = _get_law_oc()
    if not oc:
        return "❌ 법제처 API 인증키(OC)를 확인할 수 없습니다."

    eff_mst = mst_id or law_id
    eff_yd = re.sub(r"[^0-9]", "", str(enf_raw or ""))
    urls_to_try = []
    if eff_mst:
        urls_to_try.append(f"http://www.law.go.kr/DRF/lawService.do?OC={oc}&target=law&MST={eff_mst}&type=XML")
        urls_to_try.append(f"http://www.law.go.kr/DRF/lawService.do?OC={oc}&target=lsRsn&MST={eff_mst}&type=XML")
        if len(eff_yd) == 8:
            urls_to_try.append(f"http://www.law.go.kr/DRF/lawService.do?OC={oc}&target=eflaw&MST={eff_mst}&efYd={eff_yd}&type=XML")
    if law_id:
        urls_to_try.append(f"http://www.law.go.kr/DRF/lawService.do?OC={oc}&target=law&ID={law_id}&type=XML")
        urls_to_try.append(f"http://www.law.go.kr/DRF/lawService.do?OC={oc}&target=lsRsn&ID={law_id}&type=XML")

    reason_texts = []
    amend_texts = []
    addenda_texts = []

    def _parse_reason_xml(xml_bytes: bytes):
        root = ET.fromstring(xml_bytes)
        for elem in root.iter():
            tag_name = elem.tag.strip()
            txt = (elem.text or "").strip()
            if not txt:
                continue
            if tag_name in ("제개정이유내용", "제개정이유", "개정이유", "입법이유"):
                reason_texts.append(txt)
            elif tag_name in ("개정문내용", "개정문"):
                amend_texts.append(txt)
            elif tag_name == "부칙내용":
                addenda_texts.append(txt)

    for url in dict.fromkeys(urls_to_try):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req, timeout=12) as resp:
                _parse_reason_xml(resp.read())
            if reason_texts:
                break
        except Exception:
            continue

    if not reason_texts and law_name:
        try:
            enc_q = urllib.parse.quote(law_name.strip())
            s_url = f"http://www.law.go.kr/DRF/lawSearch.do?OC={oc}&target=eflaw&type=XML&query={enc_q}&display=20"
            req_s = urllib.request.Request(s_url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req_s, timeout=12) as resp_s:
                s_root = ET.fromstring(resp_s.read())
            target_clean = law_name.replace(" ", "")
            for node in s_root.findall(".//law"):
                nm = (node.findtext("법령명한글") or "").replace(" ", "")
                rev_nm = (node.findtext("제개정구분명") or "").strip()
                m_id = (node.findtext("법령일련번호") or node.findtext("MST") or "").strip()
                if nm == target_clean and m_id and "일부개정" in rev_nm:
                    u_hist = f"http://www.law.go.kr/DRF/lawService.do?OC={oc}&target=law&MST={m_id}&type=XML"
                    req_h = urllib.request.Request(u_hist, headers={"User-Agent": "Mozilla/5.0"})
                    with urllib.request.urlopen(req_h, timeout=12) as resp_h:
                        _parse_reason_xml(resp_h.read())
                    if reason_texts:
                        break
        except Exception:
            pass

    if reason_texts:
        body = "\n\n".join(dict.fromkeys(reason_texts))
        if amend_texts:
            body = f"📌 [법제처 공식 제·개정이유 및 주요내용]\n{body}\n\n---\n📌 [공포 개정문 발췌]\n{amend_texts[0][:1800]}"
        return body

    if amend_texts:
        return f"📌 [공포 개정문 원문]\n{amend_texts[0][:3500]}"

    if addenda_texts:
        return "📌 [최근 개정 부칙 원문]\n" + "\n\n".join(addenda_texts[:3])[:3500]

    return f"ℹ️ 「{law_name}」의 해당 공포본에서 별도 제·개정이유서 텍스트를 찾지 못했습니다. 우측의 '⚡ 신·구법 비교' 버튼을 눌러 조문별 비교표를 확인하세요."

def extract_cited_articles_map_in_room(messages: list[dict]) -> dict[str, list[str]]:
    """이 대화방의 본문 답변에서 실제 [조문:...] (콜론 누락 포함) 태그로 인용된 (정규화 법령명 -> 조문번호 리스트)만 정확히 추출합니다."""
    cited_map: dict[str, list[str]] = {}
    if not messages:
        return cited_map
    blob = "\n".join(m.get("content", "") for m in messages if m.get("role") == "assistant")
    for m in _CITED_ARTICLE_TAG_RE.finditer(blob):
        raw_ref = m.group(1).strip()
        m_art = re.search(r"(제\s*\d+\s*조(?:\s*의\s*\d+)?(?:\s*제\s*\d+\s*항)?)", raw_ref)
        law_part = re.sub(r"제\s*\d+\s*조.*", "", raw_ref).replace(" ", "").replace("「", "").replace("」", "").strip()
        if not law_part:
            continue
        art_str = re.sub(r"\s+", "", m_art.group(1)) if m_art else ""
        bucket = cited_map.setdefault(law_part, [])
        if art_str and art_str not in bucket:
            bucket.append(art_str)
    return cited_map

def generate_banner_law_diff_table(
    messages: list[dict],
    alert_items: list[dict],
    room_key: str = ""
) -> str:
    """
    본문 대화를 늘리지 않고 상단 배너 안에서만 열고 닫을 수 있는
    [조문명 / 구법내용 / 개정법내용 / 개정효과] 4열 마크다운 비교표를 생성합니다.
    (사용된 토큰을 대화방 요금 계량기에 정확히 반영합니다.)
    """
    pool = agent.get_live_gemini_key_pool() if hasattr(agent, "get_live_gemini_key_pool") else []
    api_key = pool[0][1] if pool else (os.getenv("GEMINI_PAID_KEY") or os.getenv("GEMINI_API_KEY") or "")
    if not api_key:
        return "❌ Gemini API 키가 설정되어 있지 않아 신·구법 비교표를 생성할 수 없습니다."

    cited_map = extract_cited_articles_map_in_room(messages)

    last_user_q = ""
    last_assistant_ans = ""
    for m in reversed(messages or []):
        if not last_assistant_ans and m.get("role") == "assistant":
            last_assistant_ans = m.get("content", "")[:6000]
        if not last_user_q and m.get("role") == "user":
            last_user_q = m.get("content", "")[:1200]
        if last_user_q and last_assistant_ans:
            break

    law_context_blocks = []
    for it in alert_items[:4]:
        lname = it["law_name"]
        cname = it["clean_name"]
        cited_arts = cited_map.get(cname, [])
        arts_label = ", ".join(cited_arts) if cited_arts else "본문 인용 조문"

        reason_txt = fetch_official_revision_reason_zero_token(
            it.get("law_id", ""),
            lname,
            it.get("mst_id", ""),
            it.get("enf_raw", "")
        )
        law_context_blocks.append(
            f"■ 법령명: 「{lname}」 (공포일: {it.get('prom_date')}, 시행일: {it.get('enf_date')})\n"
            f"- 이 대화방 본문에서 실제 인용한 조문: {arts_label}\n"
            f"- 법제처 개정이유서 참고자료:\n{reason_txt[:2200]}"
        )

    all_cited_summary = "\n".join(
        f"- {k}: {', '.join(v) if v else '(조문 전체)'}" for k, v in cited_map.items()
    )

    prompt = (
        "당신은 대한민국 세법·법령 개정 비교 전문 감수자입니다.\n"
        "아래 [이 대화방의 직전 검토 보고서 본문]과 [법령 공포 정보]를 바탕으로, "
        "상단 배너에 표시할 **신·구법 비교표**를 작성하세요.\n\n"
        "[절대 준수 규칙]\n"
        "1. **본문에서 실제로 인용한 핵심 조문(또는 본문에서 직접 다룬 개정 쟁점 조문)만** 표에 포함하세요. "
        "개정이유서에 적혀 있더라도 **본문 주제와 무관한 다른 조문(예: 통합고용세액공제를 묻는 방에서 임원 정의로 법인세법 시행령 제40조만 인용했는데 무관한 제92조의2 임대주택 양도세 조문을 넣는 행위 등)은 절대 표에 넣지 마세요!**\n"
        "2. 만약 부수적으로 인용된 시행령/법령의 해당 조문(예: 법인세법 시행령 제40조)에 개정사항이 없다면 그 조문은 표에서 생략하거나 '종전과 동일(임원 범위 변동 없음)'으로만 짧게 표기하고, 본문의 핵심 주제 법령(예: 조세특례제한법 제29조의8 등)의 신·구 비교를 중심으로 정리하세요.\n"
        "3. 본문 검토 보고서에 이미 정리된 신·구 비교 내용 및 개정 수치와 100% 일관성을 유지하세요.\n"
        "4. 인사말이나 부연 설명 없이 오직 아래 4개 컬럼의 마크다운 표(`| 조문명 | 구법내용 | 개정법내용 | 개정효과 |`)만 출력하세요.\n\n"
        f"[사용자 질문]\n{last_user_q}\n\n"
        f"[이 대화방 본문에서 인용된 조문 목록]\n{all_cited_summary}\n\n"
        f"[이 대화방의 직전 검토 보고서 본문 (발췌)]\n{last_assistant_ans}\n\n"
        f"[법제처 공포 정보]\n" + "\n\n".join(law_context_blocks)
    )

    models_to_try = [
        getattr(agent, "FLASH_MODEL_ID", "gemini-3.8-flash"),
        getattr(agent, "FLASH_BACKUP_ID", "gemini-3.6-flash"),
        getattr(agent, "FLASH_LITE_ID", "gemini-3.5-flash-lite"),
        "gemini-2.5-flash"
    ]
    for m_id in dict.fromkeys(models_to_try):
        try:
            client = genai.Client(api_key=api_key, http_options={"timeout": 60000})
            resp = client.models.generate_content(model=m_id, contents=prompt)
            out_txt = agent.extract_safe_text(resp) if hasattr(agent, "extract_safe_text") else (resp.text or "")
            um = getattr(resp, "usage_metadata", None)
            in_tok = int(getattr(um, "prompt_token_count", 0) or int(len(prompt) * 0.7))
            out_tok = (
                int(getattr(um, "candidates_token_count", 0) or int(len(out_txt) * 0.7))
                + int(getattr(um, "thoughts_token_count", 0) or 0)
            )
            _log_custom_tokens(m_id, in_tok, out_tok, count_call=True, room_key=room_key)
            if out_txt.strip():
                return out_txt.strip()
        except Exception:
            continue
    return "❌ 신·구법 비교표 생성 중 일시적인 오류가 발생했습니다. 잠시 후 다시 시도해 주세요."

def load_room_law_ack() -> dict:
    if os.path.exists(ROOM_LAW_ACK_PATH):
        try:
            with open(ROOM_LAW_ACK_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {}

def save_room_law_ack(ack_data: dict) -> None:
    try:
        with open(ROOM_LAW_ACK_PATH, "w", encoding="utf-8") as f:
            json.dump(ack_data, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

def extract_cited_laws_in_room(messages: list[dict], law_prom_map: dict[str, dict]) -> list[dict]:
    """
    콜론(:)이 빠진 [조문 ...] 태그까지 모두 추출하며,
    단순 가나다순(sorted)이 아니라 **본문에서 실제로 가장 많이 인용된 빈도순(동률 시 첫 등장 순서)**으로 정렬하여
    본문의 핵심 법령이 항상 대표(1순위)로 선정되도록 보장합니다.
    """
    if not messages or not law_prom_map:
        return []
    blob = "\n".join(m.get("content", "") for m in messages if m.get("role") == "assistant")
    if not blob:
        return []

    cite_counts: Counter = Counter()
    first_seen_order: dict[str, int] = {}

    for idx, m in enumerate(_CITED_ARTICLE_TAG_RE.finditer(blob)):
        val = m.group(1)
        val_clean = re.sub(r"제\s*\d+\s*조.*", "", val).replace(" ", "").replace("「", "").replace("」", "").strip()
        matched_cname = ""
        if val_clean in law_prom_map:
            matched_cname = val_clean
        elif val_clean.endswith("시행규칙"):
            base_c = val_clean.replace("시행규칙", "")
            if base_c in law_prom_map:
                matched_cname = base_c

        if matched_cname:
            cite_counts[matched_cname] += 1
            if matched_cname not in first_seen_order:
                first_seen_order[matched_cname] = idx

    ordered_cnames = sorted(
        cite_counts.keys(),
        key=lambda c: (-cite_counts[c], first_seen_order[c])
    )
    return [law_prom_map[cname] for cname in ordered_cnames]

def get_room_revision_alerts(ws_id: str, room_id: str, messages: list[dict], law_prom_map: dict[str, dict]) -> tuple[list[dict], bool]:
    cited_list = extract_cited_laws_in_room(messages, law_prom_map)
    if not cited_list:
        return [], False

    ack_all = load_room_law_ack()
    room_key = f"{ws_id}::{room_id}"
    room_ack = ack_all.get(room_key, {})

    post_rev_alerts = []
    recent_info_alerts = []

    for item in cited_list:
        cname = item["clean_name"]
        vkey = item["version_key"]
        acked_vkey = room_ack.get(cname)

        if acked_vkey == vkey:
            continue
        if acked_vkey is not None and acked_vkey != vkey:
            post_rev_alerts.append(item)
        elif item["is_recent_promulgation"] or item["is_future_enforcement"]:
            recent_info_alerts.append(item)

    if post_rev_alerts:
        return post_rev_alerts, True

    if recent_info_alerts and cited_list:
        primary_law = cited_list[0]
        p_cname = primary_law["clean_name"]
        if room_ack.get(p_cname) != primary_law["version_key"]:
            if all(x["clean_name"] != p_cname for x in recent_info_alerts):
                recent_info_alerts.insert(0, primary_law)

    return recent_info_alerts, False

def acknowledge_room_laws(ws_id: str, room_id: str, law_items: list[dict]) -> None:
    ack_all = load_room_law_ack()
    room_key = f"{ws_id}::{room_id}"
    room_ack = ack_all.get(room_key, {})
    for it in law_items:
        room_ack[it["clean_name"]] = it["version_key"]
    ack_all[room_key] = room_ack
    save_room_law_ack(ack_all)

DEFAULT_API_EXPIRY_CONFIG = {
    "VWORLD": {
        "name": "브이월드(공시지가) API",
        "org_name": "국토교통부 브이월드 (V-World)",
        "api_info": "부동산 개별공시지가 · 공동주택 공시가격 조회",
        "key_names": ["VWORLD_API_KEY"],
        "secret_key": "VWORLD_EXPIRY_DATE",
        "default_date": "2027-04-20",
        "cycle": "6개월",
        "url": "https://www.vworld.kr/dev/v4dv_apikey_s001.do"
    },
    "DATA_GO_KR": {
        "name": "공공데이터포털(국세청·실거래가) API",
        "org_name": "공공데이터포털 (국세청 · 국토교통부)",
        "api_info": "사업자등록 상태·진위확인 및 부동산 실거래가 조회",
        "key_names": ["NTS_API_KEY", "DATA_GO_KR_KEY", "PUBLIC_DATA_KEY", "MOLIT_API_KEY"],
        "secret_key": "DATA_GO_KR_EXPIRY_DATE",
        "default_date": "2028-10-20",
        "cycle": "2년",
        "url": "https://www.data.go.kr/iim/api/selectAcountList.do"
    },
    "BOK_ECOS": {
        "name": "한국은행(ECOS 환율·금리) API",
        "org_name": "한국은행 경제통계시스템 (ECOS)",
        "api_info": "공식 기준환율 · 재정환율 및 한국은행 기준금리 조회",
        "key_names": ["BOK_API_KEY", "ECOS_API_KEY"],
        "secret_key": "BOK_EXPIRY_DATE",
        "default_date": "2028-10-20",
        "cycle": "2년",
        "url": "https://ecos.bok.or.kr/api/#/MyPage/CertKeyInfo"
    }
}

def load_api_expiry_dates() -> dict:
    saved = {}
    if os.path.exists(EXPIRY_FILE_PATH):
        try:
            with open(EXPIRY_FILE_PATH, "r", encoding="utf-8") as f:
                saved = json.load(f)
        except Exception:
            pass
    result = {}
    for code, info in DEFAULT_API_EXPIRY_CONFIG.items():
        env_val = os.getenv(info["secret_key"], "").strip()
        result[code] = saved.get(code) or env_val or info["default_date"]
    return result

def save_api_expiry_dates(expiry_dict: dict) -> None:
    try:
        with open(EXPIRY_FILE_PATH, "w", encoding="utf-8") as f:
            json.dump(expiry_dict, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"⚠️ 만료일 저장 오류: {e}")

def check_api_expiry_alerts() -> tuple[list[dict], list[dict]]:
    expiry_map = load_api_expiry_dates()
    today_d = datetime.now(ZoneInfo("Asia/Seoul")).date()
    alerts = []
    all_status = []

    for code, info in DEFAULT_API_EXPIRY_CONFIG.items():
        d_str = expiry_map.get(code, info["default_date"])
        try:
            exp_d = datetime.strptime(d_str, "%Y-%m-%d").date()
        except Exception:
            exp_d = date(2027, 4, 20)
        days_left = (exp_d - today_d).days
        fmt_kor = f"{str(exp_d.year)[2:]}년 {exp_d.month:02d}월 {exp_d.day:02d}일"

        item_data = {
            "code": code,
            "name": info["name"],
            "org_name": info["org_name"],
            "api_info": info["api_info"],
            "key_names": info["key_names"],
            "date_str": exp_d.strftime("%Y-%m-%d"),
            "fmt_kor": fmt_kor,
            "days_left": days_left,
            "cycle": info["cycle"],
            "url": info["url"]
        }
        all_status.append(item_data)

        if days_left < 0:
            item_data["level"] = "expired"
            item_data["msg"] = (
                f"🚨 **[{info['name']} 만료됨]** 키가 **{fmt_kor}**자로 만료되었습니다(D+{-days_left}). "
                "새로운 API키를 발급받아주세요."
            )
            alerts.append(item_data)
        elif days_left <= 30:
            item_data["level"] = "warning"
            item_data["msg"] = (
                f"⚠️ **[API 만료 임박 알림]** **{info['name']}** 키가 **{fmt_kor}**자로 만료예정입니다(D-{days_left}). "
                "새로운 API키를 발급받아주세요."
            )
            alerts.append(item_data)

    return alerts, all_status

@st.cache_resource
def get_background_task_manager() -> dict:
    return {"tasks": {}, "lock": threading.Lock()}

def _log_custom_tokens(model_id: str, in_tok: int, out_tok: int, count_call: bool = True, room_key: str = "") -> None:
    if hasattr(agent, "_record_token_usage"):
        try:
            agent._record_token_usage(model_id, in_tok, out_tok, count_call=count_call, room_key=room_key)
            return
        except Exception:
            pass

def _check_api_configured(key_names: list[str]) -> bool:
    return bool(_get_first_env_or_secret(key_names))

def is_anonymous_workspace(workspace_id: str) -> bool:
    return str(workspace_id or "").startswith(ANON_WS_PREFIX)

def get_storage_path(workspace_id: str) -> str:
    """
    - 비밀번호를 입력한 작업공간: `ws_<hash>.json` (영구 보존, 다른 PC에서도 동일 암호 입력 시 이어서 사용 가능)
    - 비밀번호 없이 접속한 각 PC/브라우저 세션: `ws_anon_<hash>.json` (접속한 각 PC/브라우저마다 100% 독립된 1~10번 대화방 부여)
    """
    ws_clean = str(workspace_id or "").strip()
    safe_hash = hashlib.sha256(ws_clean.encode("utf-8")).hexdigest()[:16]
    if is_anonymous_workspace(ws_clean):
        return os.path.join(STORAGE_DIR, f"ws_anon_{safe_hash}.json")
    return os.path.join(STORAGE_DIR, f"ws_{safe_hash}.json")

def _cleanup_stale_anon_workspaces(max_age_sec: int = 86400) -> None:
    """
    비밀번호 없이 생성된 임시 개별 세션 파일(`ws_anon_*.json`) 중 24시간 이상 경과한 파일과
    과거 공용 임시 파일(`default_public`)을 정리하여 디스크 용량을 쾌적하게 유지합니다.
    (비밀번호를 설정한 `ws_*.json` 파일은 절대 건드리지 않습니다.)
    """
    now_ts = time.time()
    try:
        legacy_pub_hash = hashlib.sha256(b"default_public").hexdigest()[:16]
        legacy_pub_path = os.path.join(STORAGE_DIR, f"ws_{legacy_pub_hash}.json")
        if os.path.exists(legacy_pub_path):
            os.remove(legacy_pub_path)
    except Exception:
        pass

    try:
        for fname in os.listdir(STORAGE_DIR):
            if fname.startswith("ws_anon_") and fname.endswith(".json"):
                fpath = os.path.join(STORAGE_DIR, fname)
                if (now_ts - os.path.getmtime(fpath)) > max_age_sec:
                    os.remove(fpath)
    except Exception:
        pass

def _normalize_room_file_entry(entry: dict) -> dict:
    if not isinstance(entry, dict):
        return {"name": None, "text": "", "items": []}
    items = entry.get("items")
    if items is None:
        if entry.get("name") and entry.get("text"):
            items = [{"name": entry["name"], "text": entry["text"]}]
        else:
            items = []
    if not items:
        return {"name": None, "text": "", "items": []}
    summary_name = items[0]["name"] if len(items) == 1 else f"{items[0]['name']} 외 {len(items)-1}건"
    combined_text = "\n\n========================================\n\n".join(
        f"📄 [파일 {idx+1}: {it['name']}]\n{it['text']}" for idx, it in enumerate(items)
    )
    return {"name": summary_name, "text": combined_text, "items": items}

def get_default_workspace_data() -> dict:
    default_titles = {r_id: "빈 대화방" for r_id in ROOM_IDS}
    default_titles["대화방 1"] = "메인 대화"
    return {
        "chat_rooms": {r_id: [] for r_id in ROOM_IDS},
        "room_titles": default_titles,
        "room_files": {r_id: {"name": None, "text": "", "items": []} for r_id in ROOM_IDS},
        "room_diff_banners": {r_id: "" for r_id in ROOM_IDS}
    }

def load_workspace_from_disk(workspace_id: str) -> dict:
    path = get_storage_path(workspace_id)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
            default_data = get_default_workspace_data()
            for key in ["chat_rooms", "room_titles", "room_files", "room_diff_banners"]:
                if key not in data:
                    data[key] = default_data[key]
                else:
                    for r_id in ROOM_IDS:
                        if r_id not in data[key]:
                            data[key][r_id] = default_data[key][r_id]
            for r_id in ROOM_IDS:
                data["room_files"][r_id] = _normalize_room_file_entry(data["room_files"][r_id])
            return data
        except Exception:
            pass
    return get_default_workspace_data()

def save_workspace_to_disk(workspace_id: str, force_overwrite_room: str = ""):
    """
    [멀티룸 무손실 병합 저장]
    대화방 1~10번을 수시로 이동하며 동시 검색할 때, 특정 방(예: 대화방 1)에서 저장을 호출하더라도
    그사이 백그라운드 스레드가 디스크(`ws_xxxx.json`)에 먼저 기록해 둔 다른 방(예: 대화방 2~10)의 완료 답변이
    구버전 메모리(`st.session_state.chat_rooms`)에 의해 덮어써져 증발하지 않도록
    락(lock) 내부에서 디스크 데이터와 메모리 데이터를 방별로 비교·병합(Merge)하여 저장합니다.
    """
    if not workspace_id:
        return
    path = get_storage_path(workspace_id)
    mgr = get_background_task_manager()
    try:
        with mgr["lock"]:
            disk_data = load_workspace_from_disk(workspace_id)
            merged_rooms = {}
            for r_id in ROOM_IDS:
                mem_msgs = list(st.session_state.chat_rooms.get(r_id, []))
                dsk_msgs = list(disk_data["chat_rooms"].get(r_id, []))

                if r_id == force_overwrite_room:
                    merged_rooms[r_id] = mem_msgs
                elif len(dsk_msgs) > len(mem_msgs):
                    merged_rooms[r_id] = dsk_msgs
                    st.session_state.chat_rooms[r_id] = list(dsk_msgs)
                elif len(mem_msgs) == len(dsk_msgs) and dsk_msgs and mem_msgs:
                    if dsk_msgs[-1].get("role") == "assistant" and mem_msgs[-1].get("role") != "assistant":
                        merged_rooms[r_id] = dsk_msgs
                        st.session_state.chat_rooms[r_id] = list(dsk_msgs)
                    else:
                        merged_rooms[r_id] = mem_msgs
                else:
                    merged_rooms[r_id] = mem_msgs

                # 완료("done") 상태인 백그라운드 태스크가 남아 있으면 누락 없이 병합
                t_key = f"{workspace_id}::{r_id}"
                t_info = mgr["tasks"].get(t_key)
                if t_info and t_info.get("status") == "done" and r_id != force_overwrite_room:
                    ans = t_info.get("response", "")
                    q_msg = t_info.get("question_msg", "")
                    if q_msg and not merged_rooms[r_id]:
                        merged_rooms[r_id].append({"role": "user", "content": q_msg})
                    if ans and (not merged_rooms[r_id] or merged_rooms[r_id][-1].get("role") != "assistant" or merged_rooms[r_id][-1].get("content") != ans):
                        merged_rooms[r_id].append({"role": "assistant", "content": ans})
                    st.session_state.chat_rooms[r_id] = list(merged_rooms[r_id])

            data = {
                "chat_rooms": merged_rooms,
                "room_titles": st.session_state.room_titles,
                "room_files": st.session_state.room_files,
                "room_diff_banners": st.session_state.get("room_diff_banners", {r_id: "" for r_id in ROOM_IDS})
            }
            with open(path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"⚠️ 자동 저장 오류: {e}")

def sync_all_rooms_state(workspace_id: str) -> bool:
    """
    [1~10번 대화방 전수 무손실 동기화 엔진]
    매 렌더링 시점 및 폴링 시점마다 1~10번 모든 방의 (1) 디스크 저장본, (2) 백그라운드 완료 태스크, (3) 세션 메모리를
    하나로 동기화하여, 사용자가 검색 도중 어느 방으로 어떻게 이동하든 질문과 답변이 100% 보존되도록 보장합니다.
    """
    if not workspace_id:
        return False
    mgr = get_background_task_manager()
    any_updated = False
    try:
        with mgr["lock"]:
            disk_data = load_workspace_from_disk(workspace_id)
            if "chat_rooms" not in st.session_state:
                st.session_state.chat_rooms = disk_data["chat_rooms"]
            if "room_titles" not in st.session_state:
                st.session_state.room_titles = disk_data["room_titles"]
            if "room_files" not in st.session_state:
                st.session_state.room_files = disk_data["room_files"]
            if "room_diff_banners" not in st.session_state:
                st.session_state.room_diff_banners = disk_data.get("room_diff_banners", {r_id: "" for r_id in ROOM_IDS})

            disk_changed = False
            for r_id in ROOM_IDS:
                st.session_state.room_files[r_id] = _normalize_room_file_entry(st.session_state.room_files.get(r_id, {}))
                mem_msgs = list(st.session_state.chat_rooms.get(r_id, []))
                dsk_msgs = list(disk_data["chat_rooms"].get(r_id, []))

                if len(dsk_msgs) > len(mem_msgs):
                    mem_msgs = list(dsk_msgs)
                    st.session_state.chat_rooms[r_id] = mem_msgs
                    any_updated = True
                elif len(mem_msgs) > len(dsk_msgs):
                    dsk_msgs = list(mem_msgs)
                    disk_data["chat_rooms"][r_id] = dsk_msgs
                    disk_changed = True
                elif mem_msgs and dsk_msgs and mem_msgs[-1] != dsk_msgs[-1]:
                    if dsk_msgs[-1].get("role") == "assistant" and mem_msgs[-1].get("role") != "assistant":
                        mem_msgs = list(dsk_msgs)
                        st.session_state.chat_rooms[r_id] = mem_msgs
                        any_updated = True

                t_key = f"{workspace_id}::{r_id}"
                t_info = mgr["tasks"].get(t_key)
                if t_info and t_info.get("status") == "done":
                    ans = t_info.get("response", "")
                    q_msg = t_info.get("question_msg", "")
                    if q_msg and not mem_msgs:
                        mem_msgs.append({"role": "user", "content": q_msg})
                    if ans and (not mem_msgs or mem_msgs[-1].get("role") != "assistant" or mem_msgs[-1].get("content") != ans):
                        mem_msgs.append({"role": "assistant", "content": ans})
                        any_updated = True
                    st.session_state.chat_rooms[r_id] = list(mem_msgs)
                    if disk_data["chat_rooms"].get(r_id) != mem_msgs:
                        disk_data["chat_rooms"][r_id] = list(mem_msgs)
                        disk_changed = True
                    del mgr["tasks"][t_key]
                    any_updated = True

            if disk_changed:
                disk_data["room_titles"] = st.session_state.room_titles
                disk_data["room_files"] = st.session_state.room_files
                disk_data["room_diff_banners"] = st.session_state.get("room_diff_banners", {r_id: "" for r_id in ROOM_IDS})
                path = get_storage_path(workspace_id)
                with open(path, "w", encoding="utf-8") as f:
                    json.dump(disk_data, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"⚠️ 멀티룸 상태 동기화 오류: {e}")
    return any_updated

def delete_workspace_from_disk(workspace_id: str):
    if not workspace_id or is_anonymous_workspace(workspace_id):
        return
    path = get_storage_path(workspace_id)
    if os.path.exists(path):
        try:
            os.remove(path)
        except Exception as e:
            print(f"⚠️ 파일 삭제 오류: {e}")

def _background_agent_worker(
    ws_id: str,
    room_id: str,
    full_input: str,
    raw_query: str,
    has_file: bool,
    is_pro: bool,
    task_info: dict
) -> None:
    room_key = f"{ws_id}::{room_id}"

    def _on_live_step(msg: str):
        task_info["live_step"] = msg

    try:
        try:
            response_text = run_agent(
                user_input=full_input,
                raw_query=raw_query,
                has_file=has_file,
                room_key=room_key,
                status_callback=_on_live_step
            )
        except TypeError:
            response_text = run_agent(
                user_input=full_input,
                raw_query=raw_query,
                has_file=has_file,
                room_key=room_key
            )
    except Exception as e:
        response_text = f"❌ 에이전트 실행 중 오류가 발생했습니다: {e}"

    try:
        mgr = get_background_task_manager()
        with mgr["lock"]:
            disk_data = load_workspace_from_disk(ws_id)
            room_msgs = list(disk_data["chat_rooms"].get(room_id, []))
            q_msg = task_info.get("question_msg", "")
            if q_msg and (not room_msgs or room_msgs[-1].get("role") != "user"):
                if not any(m.get("role") == "user" and m.get("content") == q_msg for m in room_msgs[-2:]):
                    room_msgs.append({"role": "user", "content": q_msg})
            if not room_msgs or room_msgs[-1].get("role") != "assistant" or room_msgs[-1].get("content") != response_text:
                room_msgs.append({"role": "assistant", "content": response_text})
                disk_data["chat_rooms"][room_id] = room_msgs
                path = get_storage_path(ws_id)
                with open(path, "w", encoding="utf-8") as f:
                    json.dump(disk_data, f, ensure_ascii=False, indent=2)
            # 락(lock) 내부에서 response와 status="done"을 원자적(Atomic)으로 설정
            task_info["response"] = response_text
            task_info["status"] = "done"
    except Exception as e:
        print(f"⚠️ 백그라운드 디스크 저장 오류: {e}")
        task_info["response"] = response_text
        task_info["status"] = "done"

# ---------------------------------------------------------
# 헬퍼 함수 1: 파일 추출 + 📸 사진/스캔본 PDF 1회 자동 OCR (Gemini Flash 비전)
# ---------------------------------------------------------
MAX_DOC_CHARS_STANDARD = 50000
MAX_DOC_CHARS_PRO      = 150000

def _ocr_with_gemini_flash(file_bytes: bytes, mime_type: str, filename: str, room_key: str = "") -> str:
    pool = agent.get_live_gemini_key_pool() if hasattr(agent, "get_live_gemini_key_pool") else []
    api_key = pool[0][1] if pool else (os.getenv("GEMINI_PAID_KEY") or os.getenv("GEMINI_API_KEY") or "")
    if not api_key:
        return "❌ [OCR 오류] 설정된 Gemini API 키가 없어 이미지/스캔본 글자를 인식할 수 없습니다."

    ocr_prompt = (
        f"첨부된 문서/이미지({filename})의 모든 텍스트, 숫자, 사업자등록번호, 날짜, 금액, 계약 조항, "
        "그리고 표(Table) 내용을 단 한 글자도 누락하거나 요약하지 말고 원문 그대로 추출하세요. "
        "표가 있다면 마크다운 표(| 컬럼 | ... |) 형식으로 정확히 옮겨 적으세요."
    )
    models_to_try = [
        getattr(agent, "FLASH_MODEL_ID", "gemini-3.8-flash"),
        getattr(agent, "FLASH_BACKUP_ID", "gemini-3.6-flash"),
        getattr(agent, "FLASH_LITE_ID", "gemini-3.5-flash-lite"),
        "gemini-2.5-flash"
    ]
    last_err = ""
    for m_id in dict.fromkeys(models_to_try):
        try:
            client = genai.Client(api_key=api_key, http_options={"timeout": 90000})
            resp = client.models.generate_content(
                model=m_id,
                contents=[
                    types.Part.from_bytes(data=file_bytes, mime_type=mime_type),
                    ocr_prompt
                ]
            )
            extracted = agent.extract_safe_text(resp) if hasattr(agent, "extract_safe_text") else (resp.text or "")
            um = getattr(resp, "usage_metadata", None)
            in_tok = int(getattr(um, "prompt_token_count", 0) or 1500)
            out_tok = (
                int(getattr(um, "candidates_token_count", 0) or max(200, int(len(extracted) * 0.7)))
                + int(getattr(um, "thoughts_token_count", 0) or 0)
            )
            _log_custom_tokens(m_id, in_tok, out_tok, count_call=True, room_key=room_key)
            if extracted.strip():
                return f"[📸 AI 비전 OCR 자동 판독 완료: {filename}]\n{extracted.strip()}"
        except Exception as e:
            last_err = str(e)
            continue
    return f"❌ 이미지/스캔본 OCR 판독 실패 ({filename}): {last_err}"

def extract_file_content(uploaded_file, room_key: str = "") -> str:
    filename = uploaded_file.name
    lower_name = filename.lower()
    try:
        uploaded_file.seek(0)
        raw_bytes = uploaded_file.read()

        if lower_name.endswith((".xlsx", ".xls")):
            excel_file = pd.ExcelFile(io.BytesIO(raw_bytes))
            text_parts = [f"[첨부 엑셀 파일명: {filename}]"]
            for sheet in excel_file.sheet_names:
                df = excel_file.parse(sheet)
                text_parts.append(f"\n--- 시트명: {sheet} (총 {len(df)}행) ---")
                text_parts.append(df.head(1500).to_string(index=False))
            return "\n".join(text_parts)[:MAX_DOC_CHARS_PRO]

        elif lower_name.endswith(".csv"):
            for enc in ["utf-8-sig", "utf-8", "cp949", "euc-kr"]:
                try:
                    df = pd.read_csv(io.BytesIO(raw_bytes), encoding=enc)
                    return (f"[첨부 CSV 파일명: {filename} (총 {len(df)}행)]\n" + df.head(1500).to_string(index=False))[:MAX_DOC_CHARS_PRO]
                except Exception:
                    continue
            return "❌ CSV 파일의 한글 인코딩을 읽을 수 없습니다. 엑셀(.xlsx)로 저장 후 다시 올려주세요."

        elif lower_name.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(raw_bytes))
            text_parts = [f"[첨부 PDF 파일명: {filename} (총 {len(reader.pages)}페이지)]"]
            pure_body_chars = 0
            for idx, page in enumerate(reader.pages[:150]):
                page_text = (page.extract_text() or "").strip()
                pure_body_chars += len(page_text)
                text_parts.append(f"\n--- {idx+1}페이지 ---\n{page_text}")
            if pure_body_chars < 40:
                return _ocr_with_gemini_flash(raw_bytes, "application/pdf", filename, room_key=room_key)[:MAX_DOC_CHARS_PRO]
            return "\n".join(text_parts)[:MAX_DOC_CHARS_PRO]

        elif lower_name.endswith(".docx"):
            doc = Document(io.BytesIO(raw_bytes))
            text_parts = [f"[첨부 워드 파일명: {filename}]"]
            for p in doc.paragraphs:
                if p.text.strip():
                    text_parts.append(p.text.strip())
            if doc.tables:
                text_parts.append("\n--- [워드 문서 내 표(Table) 데이터] ---")
                for t_idx, table in enumerate(doc.tables):
                    text_parts.append(f"\n[표 {t_idx+1}]")
                    for row in doc.tables[t_idx].rows:
                        row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                        text_parts.append(" | ".join(row_cells))
            return "\n".join(text_parts)[:MAX_DOC_CHARS_PRO]

        elif lower_name.endswith((".png", ".jpg", ".jpeg")):
            mime = "image/png" if lower_name.endswith(".png") else "image/jpeg"
            return _ocr_with_gemini_flash(raw_bytes, mime, filename, room_key=room_key)[:MAX_DOC_CHARS_PRO]

        return "지원하지 않는 파일 형식입니다."
    except Exception as e:
        return f"❌ 파일 읽기 오류: {e}"

# ---------------------------------------------------------
# 공통 정제 도우미 & 화면 렌더링 전용 마크다운 보정기
# (URL 내부에 괄호 `()`가 포함된 판례·해석례 직링크까지 온전히 인식하도록 정규식 보강)
# ---------------------------------------------------------
_URL_BODY_PAT = r"(?:[^\s()<>\"']|\([^\s()<>\"']*\))+"
_URL_OR_MD_LINK_RE = re.compile(rf"(\]\(https?://{_URL_BODY_PAT}\)|https?://{_URL_BODY_PAT})")

def _escape_tilde_outside_urls(text: str) -> str:
    """
    일반 텍스트의 물결표(~)는 마크다운 취소선(<del>) 오작동 방지를 위해 &#126;으로 안전 변환하되,
    마크다운 링크 ](https://...) 및 순수 URL 내부에 포함된 물결표(~)는 주소가 깨지지 않도록 원형 그대로 보존합니다.
    """
    if not text or "~" not in text:
        return text or ""
    parts = _URL_OR_MD_LINK_RE.split(text)
    for idx in range(0, len(parts), 2):
        parts[idx] = parts[idx].replace("~", "&#126;")
    return "".join(parts)

def _format_markdown_for_screen(md_text: str) -> str:
    if not md_text:
        return ""

    s = _escape_tilde_outside_urls(md_text)
    s = re.sub(r"\*\*([^\*\n]+?)\*\*", r"<strong>\1</strong>", s)

    pattern_rel = re.compile(r"(###\s*3\.\s*관련\s*법령[^\n]*\n)(.*)", flags=re.DOTALL)
    m = pattern_rel.search(s)
    if m:
        heading_part = m.group(1)
        rest_body = m.group(2)
        if "\n---\n" in rest_body:
            laws_part, tail_part = rest_body.split("\n---\n", 1)
            s = (
                s[:m.start()]
                + heading_part
                + f"<div class='related-laws-compact-box'>\n\n{laws_part.strip()}\n\n</div>\n\n---\n"
                + tail_part
            )
        else:
            s = (
                s[:m.start()]
                + heading_part
                + f"<div class='related-laws-compact-box'>\n\n{rest_body.strip()}\n\n</div>"
            )
    return s

def _sanitize_math_text(text: str) -> str:
    """
    PDF/워드/복사용 텍스트 정제 시 판례 링크 주소(`https://www.law.go.kr/판례/(2005두9415)`)처럼
    URL 안에 괄호 `()`가 들어 있는 경우에도 닫는 괄호 `)`가 남지 않도록 완벽히 제거합니다.
    """
    if not text:
        return ""
    s = text
    s = re.sub(
        r"\[(판례|심판례|해석례|예규|조문)(?:\s*[:：]\s*|\s+)(?=[가-힣A-Za-z0-9「])([^\]]+)\]",
        r"[\1: \2]",
        s
    )
    s = re.sub(
        r"\[(판례|심판례|해석례|예규|조문)(?=[가-힣A-Za-z0-9「])([^\]:：\n]+)\]",
        r"[\1: \2]",
        s
    )
    s = re.sub(rf"\[(\[(?:판례|심판례|해석례|예규|조문)\s*:\s*[^\]]+\])\]\((https?://{_URL_BODY_PAT})\)", r"\1", s)
    s = re.sub(rf"\[([^\n]+?)\]\((https?://(?:www\.)?law\.go\.kr/{_URL_BODY_PAT})\)", r"\1", s)

    s = re.sub(r"\$\$(.*?)\$\$", r"\1", s, flags=re.DOTALL)
    s = re.sub(r"\$([^$\n]+)\$", r"\1", s)
    s = re.sub(r"\\text(?:bf|rm|it)?\{([^}]*)\}", r"\1", s)
    s = re.sub(r"\\frac\{([^}]*)\}\{([^}]*)\}", r"(\1/\2)", s)
    s = s.replace(r"\times", " × ").replace(r"\div", " ÷ ").replace(r"\cdot", "·")
    s = s.replace(r"\%", "%").replace(r"\_", "_").replace(r"\,", ",")
    s = re.sub(r"\\(left|right|quad|qquad|mathrm|mathbf)\b", "", s)
    s = s.replace("−", " - ").replace("–", " - ").replace("—", " - ").replace("∼", "~")
    s = s.replace("→", " -> ").replace("↔", " <-> ").replace("≧", ">=").replace("≦", "<=")
    return s

def _extract_clean_copy_text(report_text: str) -> str:
    cleaned = _sanitize_math_text(report_text)
    lines = []
    for line in cleaned.split("\n"):
        if line.strip().startswith(">"):
            continue
        lines.append(line)
    body = "\n".join(lines).strip()
    # 카톡·메신저 붙여넣기 시 별표(**)나 샵(###)이 노출되지 않도록 깔끔한 순수 텍스트로 정제
    body = re.sub(r"\*\*(.*?)\*\*", r"\1", body)
    body = re.sub(r"__(.*?)__", r"\1", body)
    body = re.sub(r"(?m)^#{1,6}\s*(.+)$", r"[\1]", body)
    return body.strip()

@st.cache_data(ttl=3600, show_spinner=False)
def create_excel_from_markdown_tables(report_text: str) -> bytes | None:
    cleaned = _sanitize_math_text(report_text)
    lines = cleaned.split("\n")
    tables = []
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if s.startswith("|") and s.endswith("|"):
            t_rows = []
            while i < len(lines):
                rs = lines[i].strip()
                if not (rs.startswith("|") and rs.endswith("|")):
                    break
                if not re.match(r"^\|[\s\-:|]+\|$", rs):
                    cells = [re.sub(r"\*\*(.*?)\*\*", r"\1", c.strip()) for c in rs.strip("|").split("|")]
                    t_rows.append(cells)
                i += 1
            if len(t_rows) >= 2:
                tables.append(t_rows)
            continue
        i += 1

    if not tables:
        return None

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for idx, t_rows in enumerate(tables):
            max_cols = max(len(r) for r in t_rows)
            padded = [r + [""] * (max_cols - len(r)) for r in t_rows]
            header = padded[0]
            seen_cols = {}
            safe_header = []
            for col_name in header:
                c_label = col_name or "항목"
                if c_label in seen_cols:
                    seen_cols[c_label] += 1
                    safe_header.append(f"{c_label}_{seen_cols[c_label]}")
                else:
                    seen_cols[c_label] = 1
                    safe_header.append(c_label)
            df = pd.DataFrame(padded[1:], columns=safe_header)
            df.to_excel(writer, sheet_name=f"표_{idx+1}", index=False)
    buffer.seek(0)
    return buffer.getvalue()

# ---------------------------------------------------------
# 헬퍼 함수 2: AI 답변을 워드(.docx)로 변환
# ---------------------------------------------------------
def _set_malgun_font(run):
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")

@st.cache_data(ttl=3600, show_spinner=False)
def create_word_docx(report_text: str) -> bytes:
    report_text = _sanitize_math_text(report_text)
    doc = Document()
    
    style = doc.styles["Normal"]
    style.font.name = "Malgun Gothic"
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    
    lines = report_text.split("\n")
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped or stripped.startswith(">"):
            i += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            table_rows = []
            while i < len(lines):
                row_s = lines[i].strip()
                if not (row_s.startswith("|") and row_s.endswith("|")):
                    break
                if not re.match(r"^\|[\s\-:|]+\|$", row_s):
                    cells = [re.sub(r"\*\*(.*?)\*\*", r"\1", c.strip()) for c in row_s.strip("|").split("|")]
                    table_rows.append(cells)
                i += 1
            if table_rows:
                max_cols = max(len(r) for r in table_rows)
                t = doc.add_table(rows=len(table_rows), cols=max_cols)
                t.style = "Table Grid"
                for r_idx, r_data in enumerate(table_rows):
                    for c_idx, val in enumerate(r_data):
                        cell = t.cell(r_idx, c_idx)
                        cell.text = val
                        for p in cell.paragraphs:
                            for run in p.runs:
                                _set_malgun_font(run)
                                if r_idx == 0:
                                    run.bold = True
            continue

        clean_line = re.sub(r"\*\*(.*?)\*\*", r"\1", stripped)
        
        if stripped.startswith("#### "):
            h = doc.add_heading(clean_line.replace("#### ", ""), level=2)
            for run in h.runs:
                _set_malgun_font(run)
        elif stripped.startswith("### ") or stripped.startswith("## "):
            heading_text = clean_line.replace("### ", "").replace("## ", "")
            h = doc.add_heading(heading_text, level=1)
            for run in h.runs:
                _set_malgun_font(run)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(clean_line[2:], style="List Bullet")
            p.paragraph_format.space_after = Pt(4)
        elif not re.match(r"^-{3,}$", stripped):
            p = doc.add_paragraph(clean_line)
            p.paragraph_format.space_after = Pt(6)
        i += 1
            
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

# ---------------------------------------------------------
# 헬퍼 함수 3: 정식 PDF 직접 생성 엔진 (+ 📕 하단 면책문구 여백 최적화 완료)
# ---------------------------------------------------------
def _ensure_korean_font() -> tuple[str, str]:
    reg_name = "KoreanFont"
    bold_name = "KoreanFontBold"
    if reg_name in pdfmetrics.getRegisteredFontNames():
        return reg_name, bold_name

    font_dir = STORAGE_DIR
    reg_path = os.path.join(font_dir, "NanumGothic.ttf")
    bold_path = os.path.join(font_dir, "NanumGothicBold.ttf")

    win_malgun = "C:/Windows/Fonts/malgun.ttf"
    win_malgun_bd = "C:/Windows/Fonts/malgunbd.ttf"

    if os.path.exists(win_malgun) and os.path.exists(win_malgun_bd):
        reg_path, bold_path = win_malgun, win_malgun_bd
    else:
        try:
            if not os.path.exists(reg_path):
                urllib.request.urlretrieve(
                    "https://github.com/google/fonts/raw/main/ofl/nanumgothic/NanumGothic-Regular.ttf",
                    reg_path
                )
            if not os.path.exists(bold_path):
                urllib.request.urlretrieve(
                    "https://github.com/google/fonts/raw/main/ofl/nanumgothic/NanumGothic-Bold.ttf",
                    bold_path
                )
        except Exception as e:
            print(f"⚠️ 한글 폰트 다운로드 오류: {e}")

    try:
        pdfmetrics.registerFont(TTFont(reg_name, reg_path))
        pdfmetrics.registerFont(TTFont(bold_name, bold_path if os.path.exists(bold_path) else reg_path))
        return reg_name, bold_name
    except Exception:
        return "Helvetica", "Helvetica-Bold"

def _strip_emojis_for_pdf(text: str) -> str:
    return re.sub(r"[\U00010000-\U0010ffff]", "", text or "")

def _md_to_rl_xml(text: str, bold_font: str) -> str:
    s = _strip_emojis_for_pdf(text).strip()
    s = s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    s = re.sub(r"\*\*(.+?)\*\*", rf'<font name="{bold_font}">\1</font>', s)
    s = re.sub(r"`([^`]+)`", r'<font backColor="#f1f5f9">\1</font>', s)
    return s

def _draw_page_footer(canvas, doc):
    canvas.saveState()
    reg_font, _ = _ensure_korean_font()

    canvas.setStrokeColor(colors.HexColor("#e2e8f0"))
    canvas.setLineWidth(0.6)
    canvas.line(40, 36, A4[0] - 40, 36)

    canvas.setFont(reg_font, 6.2)
    canvas.setFillColor(colors.HexColor("#64748b"))
    disclaimer_str = (
        "※ 본 보고서는 질의하신 사실관계와 현행 법령 기준의 참고자료이며, "
        "사실관계 변동이나 과세관청·법원의 해석에 따라 판단이 달라질 수 있습니다."
    )
    canvas.drawString(40, 24, disclaimer_str)

    canvas.setFont(reg_font, 8.2)
    canvas.setFillColor(colors.HexColor("#475569"))
    page_num_str = f"- {doc.page} -"
    canvas.drawRightString(A4[0] - 40, 24, page_num_str)
    canvas.restoreState()

@st.cache_data(ttl=3600, show_spinner=False)
def create_pdf_report(report_text: str) -> bytes:
    reg_font, bold_font = _ensure_korean_font()
    cleaned_text = _sanitize_math_text(report_text)
    raw_lines = cleaned_text.split("\n")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=40,
        rightMargin=40,
        topMargin=42,
        bottomMargin=48
    )
    usable_width = A4[0] - 80

    styles = getSampleStyleSheet()
    style_h2 = ParagraphStyle(
        "DocH2", parent=styles["Normal"], fontName=bold_font,
        fontSize=13, leading=19, textColor=colors.HexColor("#1f4e78"),
        spaceBefore=12, spaceAfter=6
    )
    style_h3 = ParagraphStyle(
        "DocH3", parent=styles["Normal"], fontName=bold_font,
        fontSize=11, leading=16, textColor=colors.HexColor("#2e5cb8"),
        spaceBefore=10, spaceAfter=4
    )
    style_body = ParagraphStyle(
        "DocBody", parent=styles["Normal"], fontName=reg_font,
        fontSize=9.8, leading=15.5, textColor=colors.HexColor("#1e293b"),
        spaceAfter=4, wordWrap="CJK"
    )
    style_bullet = ParagraphStyle(
        "DocBullet", parent=style_body, leftIndent=12, firstLineIndent=-8, spaceAfter=3.5
    )
    style_th = ParagraphStyle(
        "DocTH", parent=style_body, fontName=bold_font, fontSize=9.2, leading=13.5,
        textColor=colors.HexColor("#0f172a"), alignment=0
    )
    style_td = ParagraphStyle(
        "DocTD", parent=style_body, fontName=reg_font, fontSize=9.0, leading=13.5,
        textColor=colors.HexColor("#1e293b"), alignment=0
    )

    story = []

    i = 0
    while i < len(raw_lines):
        line = raw_lines[i]
        stripped = line.strip()

        if not stripped:
            story.append(Spacer(1, 4))
            i += 1
            continue

        if stripped.startswith(">"):
            i += 1
            continue

        if re.match(r"^-{3,}$", stripped):
            story.append(Spacer(1, 4))
            story.append(HRFlowable(width="100%", thickness=0.7, color=colors.HexColor("#cbd5e1")))
            story.append(Spacer(1, 6))
            i += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while i < len(raw_lines):
                r_str = raw_lines[i].strip()
                if not (r_str.startswith("|") and r_str.endswith("|")):
                    break
                table_lines.append(r_str)
                i += 1

            rows_data = []
            for r_str in table_lines:
                if re.match(r"^\|[\s\-:|]+\|$", r_str):
                    continue
                cells = [c.strip() for c in r_str.strip("|").split("|")]
                rows_data.append(cells)

            if rows_data:
                max_cols = max(len(r) for r in rows_data)
                pdf_table_data = []
                for r_idx, r in enumerate(rows_data):
                    padded_r = r + [""] * (max_cols - len(r))
                    row_paras = []
                    for c_text in padded_r:
                        cell_xml = _md_to_rl_xml(c_text, bold_font)
                        row_paras.append(Paragraph(cell_xml, style_th if r_idx == 0 else style_td))
                    pdf_table_data.append(row_paras)

                col_w = usable_width / max_cols
                pdf_table = Table(pdf_table_data, colWidths=[col_w] * max_cols, repeatRows=1)
                t_style = [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ]
                for r_idx in range(1, len(pdf_table_data)):
                    if r_idx % 2 == 0:
                        t_style.append(("BACKGROUND", (0, r_idx), (-1, r_idx), colors.HexColor("#f8fafc")))
                pdf_table.setStyle(TableStyle(t_style))
                story.append(Spacer(1, 4))
                story.append(pdf_table)
                story.append(Spacer(1, 8))
            continue

        if stripped.startswith("#### "):
            story.append(Paragraph(_md_to_rl_xml(stripped[5:], bold_font), style_h3))
        elif stripped.startswith("### "):
            story.append(Paragraph(_md_to_rl_xml(stripped[4:], bold_font), style_h2))
        elif stripped.startswith("## "):
            story.append(Paragraph(_md_to_rl_xml(stripped[3:], bold_font), style_h2))
        elif stripped.startswith("# "):
            story.append(Paragraph(_md_to_rl_xml(stripped[2:], bold_font), style_h2))
        elif stripped.startswith(("- ", "* ")):
            indent_level = (len(line) - len(line.lstrip(" "))) // 2
            b_style = ParagraphStyle(
                f"Bullet_{indent_level}", parent=style_bullet,
                leftIndent=12 + (indent_level * 10)
            )
            story.append(Paragraph("• " + _md_to_rl_xml(stripped[2:], bold_font), b_style))
        else:
            story.append(Paragraph(_md_to_rl_xml(stripped, bold_font), style_body))

        i += 1

    if not story:
        story.append(Paragraph("내용 없음", style_body))

    doc.build(story, onFirstPage=_draw_page_footer, onLaterPages=_draw_page_footer)
    buffer.seek(0)
    return buffer.getvalue()

# ---------------------------------------------------------
# 헬퍼 함수 4: 질문/원클릭 액션 통합 백그라운드 실행기
# ---------------------------------------------------------
def _trigger_agent_task(
    effective_ws_id: str,
    selected_room: str,
    raw_prompt: str,
    pro_toggle_on: bool,
    budget_exceeded: bool = False,
    force_flash: bool = False
) -> None:
    task_mgr = get_background_task_manager()
    sync_all_rooms_state(effective_ws_id)
    current_messages = st.session_state.chat_rooms[selected_room]
    current_file_info = _normalize_room_file_entry(st.session_state.room_files[selected_room])
    st.session_state.room_files[selected_room] = current_file_info

    effective_prompt = raw_prompt.strip()
    if force_flash or budget_exceeded:
        if effective_prompt.startswith("[프로]"):
            effective_prompt = effective_prompt.replace("[프로]", "", 1).strip()
    elif pro_toggle_on and not effective_prompt.startswith("[프로]"):
        effective_prompt = f"[프로] {effective_prompt}"

    has_attached_file = bool(current_file_info["name"] and current_file_info["text"])
    display_prompt = f"📎 **[첨부파일: {current_file_info['name']}]**\n\n{effective_prompt}" if has_attached_file else effective_prompt

    if st.session_state.room_titles[selected_room] in ["빈 대화방", "메인 대화"] and not current_messages:
        auto_title = raw_prompt.strip().replace("[프로]", "").strip().replace("\n", " ")[:14]
        st.session_state.room_titles[selected_room] = auto_title
        # 이미 인스턴스화된 사이드바 text_input 위젯 값을 직접 수정하면 StreamlitAPIException이 발생하므로
        # pop()으로 위젯 상태 키만 제거하여 다음 rerun 시 room_titles 값으로 안전하게 채워지도록 처리합니다.
        st.session_state.pop(f"title_input_{effective_ws_id}_{selected_room}", None)

    is_pro_request = (
        not force_flash and not budget_exceeded and (
            effective_prompt.startswith("[프로]") or (has_attached_file and len(current_file_info["text"]) >= 15000)
        )
    )
    msg_char_cap = 24000 if is_pro_request else 8000

    history_context = ""
    if current_messages:
        recent_msgs = (current_messages[:2] + current_messages[-6:]) if len(current_messages) > 6 else current_messages
        history_lines = [
            "[이 대화방의 이전 대화 맥락 — 사용자가 별도 수식어 없이 후속 질문이나 조건 변경, 카톡/이메일 초안 변환, 또는 '[프로] 다시 검토해 줘'라고 말하더라도 반드시 아래 직전 질문·답변의 사례 금액, 자산 정보, 계산 내역을 그대로 이어받아 답변할 것]"
        ]
        for m in recent_msgs:
            role_label = "사용자" if m["role"] == "user" else "에이전트"
            history_lines.append(f"- {role_label}: {m['content'][:msg_char_cap]}")
        history_context = "\n".join(history_lines) + "\n\n"

    current_messages.append({"role": "user", "content": display_prompt})

    if has_attached_file and not force_flash:
        doc_cap = MAX_DOC_CHARS_PRO if is_pro_request else MAX_DOC_CHARS_STANDARD
        full_input_for_agent = (
            f"{history_context}"
            f"[현재 사용자 질문]: {effective_prompt}\n\n"
            f"아래는 사용자가 현재 대화방에 업로드한 첨부파일({current_file_info['name']})의 데이터입니다. "
            f"직전 대화 맥락과 이 파일 데이터를 바탕으로 관련 법령·예규·판례·환율을 정밀 대조하여 답변하세요:\n"
            f"========================================\n"
            f"{current_file_info['text'][:doc_cap]}\n"
            f"========================================"
        )
    else:
        full_input_for_agent = f"{history_context}[현재 사용자 질문]: {effective_prompt}" if history_context else effective_prompt

    new_task_info = {
        "status": "running",
        "response": None,
        "question_msg": display_prompt,
        "is_pro": is_pro_request,
        "live_step": "",
        "started_at": time.time()
    }
    worker_thread = threading.Thread(
        target=_background_agent_worker,
        args=(
            effective_ws_id,
            selected_room,
            full_input_for_agent,
            effective_prompt,
            (has_attached_file and not force_flash),
            is_pro_request,
            new_task_info
        ),
        daemon=True
    )
    new_task_info["thread"] = worker_thread
    with task_mgr["lock"]:
        task_mgr["tasks"][f"{effective_ws_id}::{selected_room}"] = new_task_info
    save_workspace_to_disk(effective_ws_id)
    worker_thread.start()
    st.rerun()

# ---------------------------------------------------------
# 헬퍼 함수 5: 📜 법제처 공식 『제·개정이유서』 0토큰 원문 열람 뷰어 & 🔑 API키 통합 관리 모달
# ---------------------------------------------------------
_dialog_decorator = getattr(st, "dialog", getattr(st, "experimental_dialog", None))

def _render_official_reason_viewer_body(
    law_prom_map: dict[str, dict],
    effective_ws_id: str,
    selected_room: str,
    pro_toggle_on: bool,
    budget_exceeded: bool = False,
    preselected_clean_name: str = ""
):
    st.caption(
        "💡 **AI 토큰을 전혀 소모하지 않고(0원)** 국가법령정보센터(법제처 DRF) XML에서 "
        "정부가 직접 작성한 **공식 『제·개정이유 및 주요내용(입법이유)』** 원문을 즉시 불러옵니다."
    )
    items = list(law_prom_map.values())
    if not items:
        st.warning("법제처 공포 정보를 불러오지 못했습니다.")
        return

    options_clean = [it["clean_name"] for it in items]
    default_idx = options_clean.index(preselected_clean_name) if preselected_clean_name in options_clean else 0

    def _fmt_law_opt(cname: str) -> str:
        it = law_prom_map[cname]
        tag = "🚀시행예정" if it["is_future_enforcement"] else (f"🆕공포 D+{it['days_since_prom']}" if it["is_recent_promulgation"] else "현행")
        return f"📘 {it['law_name']} ({it['rev_type']} | 공포: {it['prom_date']} | 시행: {it['enf_date']} | {tag})"

    chosen_clean = st.selectbox(
        "열람할 법령·시행령 선택",
        options=options_clean,
        index=default_idx,
        format_func=_fmt_law_opt
    )
    chosen_item = law_prom_map[chosen_clean]

    st.markdown(
        f"**공포번호:** 제{chosen_item['prom_no']}호 &nbsp;|&nbsp; "
        f"**공포일:** `{chosen_item['prom_date']}` &nbsp;|&nbsp; "
        f"**시행일:** `{chosen_item['enf_date']}` &nbsp;|&nbsp; "
        f"**구분:** `{chosen_item['rev_type']}`"
    )

    with st.spinner("법제처 공식 『제·개정이유서(입법이유)』 원문 수신 중 (0토큰)..."):
        official_reason_text = fetch_official_revision_reason_zero_token(
            chosen_item["law_id"],
            chosen_item["law_name"],
            chosen_item.get("mst_id", ""),
            chosen_item.get("enf_raw", "")
        )

    st.text_area(
        "📜 법제처 공식 『제·개정이유 및 주요내용』 원문",
        value=official_reason_text,
        height=320
    )

    c_v1, c_v2 = st.columns(2)
    with c_v1:
        if st.button("🔄 법제처 최신 공포정보 새로고침", use_container_width=True):
            fetch_latest_tax_law_promulgations.clear()
            fetch_official_revision_reason_zero_token.clear()
            st.rerun()
    with c_v2:
        if st.button(f"⚡ 상단 배너에 「{chosen_item['law_name']}」 신·구법 비교 열기", use_container_width=True):
            curr_msgs = st.session_state.chat_rooms.get(selected_room, [])
            diff_md = generate_banner_law_diff_table(
                curr_msgs, [chosen_item], room_key=f"{effective_ws_id}::{selected_room}"
            )
            st.session_state.setdefault("room_diff_banners", {})[selected_room] = diff_md
            st.session_state[f"diff_banner_open_{effective_ws_id}_{selected_room}"] = True
            save_workspace_to_disk(effective_ws_id)
            st.rerun()

# ---------------------------------------------------------
# 🔐 [🔑 API키] 3단 자동 패스(1순위 Google ID 자동교차 · 2순위 내 작업공간 암호 연동 · 3순위 PIN 단독해제)
#    + [사이트 명 또는 단체 명 / 불러오는 API정보 명 / API코드 / 🟢🟡🔴 동그라미 불 / 사이트 링크] 통합 뷰어
# ---------------------------------------------------------
def _hash_manager_secret(val: str) -> str:
    return hashlib.sha256(val.strip().encode("utf-8")).hexdigest()

def _mask_google_email(email_str: str) -> str:
    email_str = (email_str or "").strip()
    if "@" not in email_str:
        return email_str
    local_part, domain_part = email_str.split("@", 1)
    if len(local_part) <= 2:
        masked_local = local_part[0] + "***" if local_part else "***"
    else:
        masked_local = local_part[:2] + "*" * max(3, len(local_part) - 2)
    return f"{masked_local}@{domain_part}"

def _format_api_code_display(raw_code: str, show_full: bool = False) -> str:
    """API 키 코드 표시 도우미: 기본은 앞 4자리·뒤 4자리 마스킹 표시, 전체 보기 토글 시 원문 표시."""
    s = (raw_code or "").strip()
    if not s:
        return "미설정"
    if show_full or len(s) <= 8:
        return s
    return f"{s[:4]}••••{s[-4:]}"

def _load_manager_google_config() -> dict:
    cfg = {"email": "", "pin_hash": "", "source": "none"}
    env_email = ""
    env_pin = ""
    for k in ["MANAGER_GOOGLE_EMAIL", "ADMIN_GOOGLE_EMAIL", "OWNER_GOOGLE_EMAIL", "MANAGER_EMAIL", "ADMIN_EMAIL"]:
        val = os.environ.get(k, "").strip()
        if not val and hasattr(st, "secrets"):
            try:
                val = str(st.secrets.get(k, "")).strip()
            except Exception:
                val = ""
        if val and "@" in val:
            env_email = val.lower()
            break

    for pk in ["MANAGER_SECURITY_PIN", "MANAGER_PASSWORD", "ADMIN_PASSWORD", "MANAGER_PIN"]:
        pval = os.environ.get(pk, "").strip()
        if not pval and hasattr(st, "secrets"):
            try:
                pval = str(st.secrets.get(pk, "")).strip()
            except Exception:
                pval = ""
        if pval:
            env_pin = pval
            break

    if os.path.exists(MANAGER_AUTH_FILE):
        try:
            with open(MANAGER_AUTH_FILE, "r", encoding="utf-8") as f:
                disk_cfg = json.load(f)
            if isinstance(disk_cfg, dict):
                cfg["email"] = str(disk_cfg.get("email", "")).strip().lower()
                cfg["pin_hash"] = str(disk_cfg.get("pin_hash", "")).strip()
                cfg["source"] = "disk"
        except Exception:
            pass

    if env_email:
        cfg["email"] = env_email
        cfg["source"] = "secrets"
    if env_pin:
        cfg["pin_hash"] = _hash_manager_secret(env_pin)
        cfg["source"] = "secrets"

    return cfg

def _save_manager_google_config(email_str: str, pin_plain: str) -> None:
    os.makedirs(STORAGE_DIR, exist_ok=True)
    payload = {
        "email": email_str.strip().lower(),
        "pin_hash": _hash_manager_secret(pin_plain),
        "updated_at": datetime.now(ZoneInfo("Asia/Seoul")).strftime("%Y-%m-%d %H:%M:%S")
    }
    with open(MANAGER_AUTH_FILE, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)

def _get_streamlit_logged_google_email() -> str:
    try:
        u_obj = getattr(st, "user", None) or getattr(st, "experimental_user", None)
        if u_obj is not None:
            is_logged = getattr(u_obj, "is_logged_in", True)
            raw_email = getattr(u_obj, "email", None)
            if not raw_email and isinstance(u_obj, dict):
                raw_email = u_obj.get("email")
            if is_logged and raw_email:
                clean_em = str(raw_email).strip().lower()
                if clean_em and clean_em != "test@example.com" and "@" in clean_em:
                    return clean_em
    except Exception:
        pass
    return ""

def _check_three_tier_manager_pass(registered_email: str, registered_pin_hash: str) -> tuple[bool, str, str]:
    """
    [3단 자동 패스 검증기]
    - 1단 (Google 계정 자동 교차검증): Streamlit/OAuth에 로그인된 구글 이메일이 `MANAGER_GOOGLE_EMAIL`과 일치하면 즉시 통과
    - 2단 (내 작업공간 암호 자동 교차검증): 좌측 사이드바 `🔐 내 작업공간 코드`에 입력한 암호가 `MANAGER_SECURITY_PIN`(또는 이메일+PIN 조합)과 일치하면 즉시 통과
    - 3단 (PIN 단독 인증 세션): 모달에서 `MANAGER_SECURITY_PIN` 1개만 입력해 통과한 경우 세션 유지
    """
    oauth_google_email = _get_streamlit_logged_google_email()
    session_authed_email = str(st.session_state.get("manager_google_authed_email", "")).strip().lower()
    active_ws_code = str(st.session_state.get("active_workspace_id", "")).strip()

    # 1순위: 구글 계정 자동 교차검증
    if registered_email and oauth_google_email and oauth_google_email == registered_email:
        st.session_state.manager_google_authed_email = registered_email
        return True, "1단 자동패스 (Google 계정 이메일 자동 교차검증 완료)", oauth_google_email

    # 2순위: 좌측 사이드바 '내 작업공간 코드'와 MANAGER_SECURITY_PIN 교차검증
    if active_ws_code and registered_pin_hash and not st.session_state.get("manager_manual_locked", False):
        ws_hash = _hash_manager_secret(active_ws_code)
        combo_ok = False
        if hmac.compare_digest(ws_hash, registered_pin_hash):
            combo_ok = True
        elif registered_email and ":" in active_ws_code:
            em_part, pin_part = active_ws_code.split(":", 1)
            if em_part.strip().lower() == registered_email and hmac.compare_digest(_hash_manager_secret(pin_part), registered_pin_hash):
                combo_ok = True
        if combo_ok:
            st.session_state.manager_google_authed_email = registered_email or "manager@verified"
            return True, "2단 자동패스 (좌측 '내 작업공간 코드' 보안키 교차검증 완료)", oauth_google_email

    # 3순위: 이미 현재 세션에서 PIN 단독 인증을 통과한 경우
    if session_authed_email and (session_authed_email == registered_email or (not registered_email and session_authed_email == "manager@verified")):
        return True, "3단 패스 (관리자 보안 PIN 교차검증 완료)", oauth_google_email

    return False, "", oauth_google_email

def _render_api_key_manager_body(expiry_all_status: list[dict], law_prom_map: dict):
    mgr_cfg = _load_manager_google_config()
    registered_email = mgr_cfg.get("email", "").strip().lower()
    registered_pin_hash = mgr_cfg.get("pin_hash", "").strip()

    # 최초 미등록 상태 (Secrets에도 없고 디스크에도 없을 때)
    if not registered_email and not registered_pin_hash:
        oauth_em = _get_streamlit_logged_google_email()
        st.warning("🔒 **[보안 설정 필요] 관리자(Manager) Google 계정과 보안 PIN을 최초 1회 등록해 주세요.**")
        st.caption(
            "Streamlit Cloud의 Secrets에 `MANAGER_GOOGLE_EMAIL`과 `MANAGER_SECURITY_PIN`을 넣어두시거나, "
            "아래에서 최초 1회 등록하시면 **3단 자동 패스(① 구글ID 자동대조 → ② 작업공간 암호 연동 → ③ PIN 단독 해제)**가 즉시 작동합니다."
        )
        with st.form("mgr_initial_register_form", clear_on_submit=False):
            st.markdown("##### 🛡️ 관리자(Manager) Google 계정 및 보안 PIN 최초 등록")
            reg_email = st.text_input(
                "관리자 Google 이메일 주소 (MANAGER_GOOGLE_EMAIL)",
                value=oauth_em if oauth_em else "",
                placeholder="예: myaccount@gmail.com"
            )
            reg_pin1 = st.text_input(
                "관리자 전용 보안 PIN 설정 (MANAGER_SECURITY_PIN, 4자리 이상)",
                type="password",
                placeholder="본인만 아는 보안 PIN 입력 (좌측 작업공간 암호와 동일하게 설정하시면 자동 연동됩니다)"
            )
            reg_pin2 = st.text_input(
                "관리자 전용 보안 PIN 확인",
                type="password",
                placeholder="보안 PIN 다시 입력"
            )
            submitted_reg = st.form_submit_button(
                "✅ 관리자 등록 및 3단 자동 패스 활성화",
                type="primary",
                use_container_width=True
            )
            if submitted_reg:
                clean_reg_em = (reg_email or "").strip().lower()
                p1 = (reg_pin1 or "").strip()
                p2 = (reg_pin2 or "").strip()
                if not clean_reg_em or "@" not in clean_reg_em:
                    st.error("유효한 Google 이메일 주소(예: example@gmail.com)를 입력해 주세요.")
                elif len(p1) < 4:
                    st.error("보안 PIN은 최소 4자리 이상 입력해 주세요.")
                elif p1 != p2:
                    st.error("입력하신 두 보안 PIN이 일치하지 않습니다.")
                else:
                    _save_manager_google_config(clean_reg_em, p1)
                    st.session_state.manager_google_authed_email = clean_reg_em
                    st.session_state.manager_manual_locked = False
                    registered_email = clean_reg_em
                    registered_pin_hash = _hash_manager_secret(p1)
                    st.success(f"✅ 관리자({clean_reg_em}) 등록 및 인증이 완료되었습니다!")
        if not str(st.session_state.get("manager_google_authed_email", "")).strip():
            return

    is_manager_verified, pass_stage_label, oauth_google_email = _check_three_tier_manager_pass(
        registered_email, registered_pin_hash
    )

    if not is_manager_verified:
        masked_mgr = _mask_google_email(registered_email) if registered_email else "미지정(PIN 전용)"
        st.info(
            f"🛡️ **관리자 3단 자동 패스 보안 구역** — 등록된 관리자 계정(`{masked_mgr}`)\n\n"
            "- **1단 (자동):** Google 로그인 세션이 감지되면 별도 입력 없이 즉시 열립니다.\n"
            "- **2단 (자동):** 좌측 사이드바 **`🔐 내 작업공간 코드`**에 관리자 보안 PIN(`MANAGER_SECURITY_PIN`)을 입력해 두시면 별도 로그인 없이 즉시 열립니다.\n"
            "- **3단 (수동):** 아래 칸에 **보안 PIN 번호 1개만** 입력하면 내부 `MANAGER_GOOGLE_EMAIL`과 자동 교차검증하여 즉시 열립니다."
        )
        if oauth_google_email and registered_email and oauth_google_email != registered_email:
            st.warning(
                f"🚫 현재 감지된 Google 계정(`{oauth_google_email}`)은 등록된 관리자(`{masked_mgr}`)와 다릅니다. 아래 보안 PIN으로 인증해 주세요."
            )

        with st.form("mgr_pin_quick_unlock_form", clear_on_submit=False):
            st.markdown("##### 🔑 3단 패스: 보안 PIN 단독 입력으로 즉시 열기 (이메일 입력 불필요)")
            input_g_pin = st.text_input(
                "관리자 보안 PIN (MANAGER_SECURITY_PIN)",
                type="password",
                placeholder=f"관리자({masked_mgr}) 보안 PIN 입력 후 Enter"
            )
            c_u1, c_u2 = st.columns([3, 1.4])
            with c_u1:
                submitted_pin = st.form_submit_button(
                    "🔓 보안 PIN 교차검증 및 API키 즉시 확인",
                    type="primary",
                    use_container_width=True
                )
            with c_u2:
                oauth_try = st.form_submit_button(
                    "🌐 Google OAuth 재확인",
                    use_container_width=True
                )

            if submitted_pin:
                cand_pin_plain = (input_g_pin or "").strip()
                cand_pin_hash = _hash_manager_secret(cand_pin_plain)
                if not registered_pin_hash:
                    st.error("🚫 관리자 보안 PIN(MANAGER_SECURITY_PIN)이 설정되어 있지 않습니다. Secrets에 MANAGER_SECURITY_PIN을 추가해 주세요.")
                elif not cand_pin_plain or not hmac.compare_digest(cand_pin_hash, registered_pin_hash):
                    st.error("🚫 보안 PIN이 일치하지 않습니다.")
                else:
                    st.session_state.manager_google_authed_email = registered_email or "manager@verified"
                    st.session_state.manager_manual_locked = False
                    st.rerun()
            elif oauth_try:
                if hasattr(st, "login"):
                    try:
                        st.login()
                    except Exception:
                        st.caption("💡 현재 Streamlit OIDC 설정이 없어 보안 PIN 또는 좌측 작업공간 코드로 즉시 인증하실 수 있습니다.")
                else:
                    st.caption("💡 현재 환경에서는 보안 PIN 또는 좌측 작업공간 코드로 즉시 인증하실 수 있습니다.")
        return

    # ---------------------------------------------------------
    # ✅ 인증 통과 화면: 요청하신 5열 순서
    #    [사이트 명 또는 단체 명 / 불러오는 API정보 명 / API코드 / 🟢(만료 1달 전 🟡) 동그라미 불 / 사이트 링크]
    # ---------------------------------------------------------
    with st.container(border=True):
        c_top1, c_top2, c_top3 = st.columns([3.0, 1.2, 1.0])
        with c_top1:
            st.markdown(f"🛡️ **관리자 교차검증 완료:** `{registered_email or 'Manager'}` &nbsp;|&nbsp; <span class='sb-badge-green'>{pass_stage_label}</span>", unsafe_allow_html=True)
        with c_top2:
            show_full_code = st.toggle("👁️ API코드 전체 보기", value=False, key="toggle_show_full_api_code")
        with c_top3:
            if st.button("🔒 잠금", use_container_width=True, key="btn_lock_api_modal"):
                st.session_state.manager_google_authed_email = ""
                st.session_state.manager_manual_locked = True
                st.rerun()

    # 1~3번 상시/무기한 API 데이터 구성
    gemini_raw = _get_first_env_or_secret(["GEMINI_PAID_KEY", "GEMINI_API_KEY"])
    law_raw = _get_law_oc()
    dart_raw = _get_first_env_or_secret(["DART_API_KEY", "OPENDART_API_KEY"])
    ind_ok = os.path.exists("industry_codes.xlsx")

    api_rows = [
        {
            "site_org": "1. Google AI Studio (Google Cloud)",
            "api_info": "Gemini 유료 에이전트 엔진 (Argon / 3.8 Flash / 3.1 Pro)",
            "api_code": f"`{_format_api_code_display(gemini_raw, show_full_code)}`" if gemini_raw else "`미설정 (GEMINI_PAID_KEY)`",
            "led_status": "🟢 정상 연동 (상시)" if gemini_raw else "⚪ 미설정",
            "link_md": "[🔗 사이트 이동](https://aistudio.google.com/apikey)"
        },
        {
            "site_org": "2. 법제처 국가법령정보센터 (MOLEG)",
            "api_info": "현행법령 · 연혁법령(구법) · 대법원판례 · 심판례 · 예규 · 개정이유서",
            "api_code": f"`{_format_api_code_display(law_raw, show_full_code)}`" if law_raw else "`미설정 (LAW_OPEN_API_KEY)`",
            "led_status": "🟢 정상 연동 (무기한)" if (law_raw or law_prom_map) else "⚪ 미설정",
            "link_md": "[🔗 사이트 이동](https://open.law.go.kr/LSO/openApi/guideList.do)"
        },
        {
            "site_org": "3. 금융감독원 전자공시시스템 (Open DART)",
            "api_info": "기업개황 · 요약재무제표 · 최대주주 지분 · 배당 공시",
            "api_code": f"`{_format_api_code_display(dart_raw, show_full_code)}`" if dart_raw else "`미설정 (DART_API_KEY)`",
            "led_status": "🟢 정상 연동 (무기한)" if dart_raw else "⚪ 미설정",
            "link_md": "[🔗 사이트 이동](https://opendart.fss.or.kr/uss/umt/EgovMberManage.do)"
        },
    ]

    # 4~6번 주기적 갱신 API (만료 1달 전 🟡 노란불, 정상 🟢 초록불, 만료 시 🔴 빨간불)
    for idx, st_item in enumerate(expiry_all_status, start=4):
        raw_k = _get_first_env_or_secret(st_item["key_names"])
        d_left = st_item["days_left"]
        if not raw_k:
            led_str = f"⚪ 미설정 (만료일: {st_item['date_str']})"
        elif d_left < 0:
            led_str = f"🔴 만료됨 (D+{-d_left} · {st_item['date_str']})"
        elif d_left <= 30:
            led_str = f"🟡 만료 임박 (D-{d_left} · {st_item['date_str']})"
        else:
            led_str = f"🟢 정상 연동 (D-{d_left} · {st_item['date_str']})"

        api_rows.append({
            "site_org": f"{idx}. {st_item['org_name']}",
            "api_info": f"{st_item['api_info']} *(갱신주기: {st_item['cycle']})*",
            "api_code": f"`{_format_api_code_display(raw_k, show_full_code)}`" if raw_k else f"`미설정 ({st_item['key_names'][0]})`",
            "led_status": led_str,
            "link_md": f"[🔗 사이트 이동]({st_item['url']})"
        })

    # 7번 국세청 업종코드 마스터 DB
    api_rows.append({
        "site_org": "7. 국세청 홈택스 (내장 마스터 DB)",
        "api_info": "국세청 6자리 업종코드 및 기준·단순경비율 (`industry_codes.xlsx`)",
        "api_code": "`industry_codes.xlsx` (로컬 내장)" if ind_ok else "`파일 없음`",
        "led_status": "🟢 정상 로드됨 (상시)" if ind_ok else "🟡 파일 없음",
        "link_md": "[🔗 사이트 이동](https://www.hometax.go.kr)"
    })

    # 요청하신 정확한 컬럼 순서로 마크다운 테이블 렌더링
    md_lines = [
        "| 사이트 명 또는 단체 명 | 불러오는 API정보 명 | API코드 | 상태 (🟢정상 / 🟡만료 1달 전) | 사이트 링크 |",
        "| :--- | :--- | :--- | :--- | :---: |"
    ]
    for r in api_rows:
        md_lines.append(
            f"| **{r['site_org']}** | {r['api_info']} | {r['api_code']} | **{r['led_status']}** | {r['link_md']} |"
        )

    st.markdown(
        "<div class='api-key-status-table-box'>\n\n" + "\n".join(md_lines) + "\n\n</div>",
        unsafe_allow_html=True
    )

    with st.expander("⏰ 주기적 갱신 API 만료 예정일 수정 및 저장 (브이월드 · 공공데이터포털 · 한국은행)"):
        current_expiry_map = load_api_expiry_dates()
        updated_expiry_map = {}
        e_cols = st.columns(len(expiry_all_status))
        for c_idx, st_item in enumerate(expiry_all_status):
            code = st_item["code"]
            with e_cols[c_idx]:
                val_str = st.text_input(
                    f"{st_item['name']} 만료일 (YYYY-MM-DD)",
                    value=current_expiry_map.get(code, st_item["date_str"]),
                    placeholder="YYYY-MM-DD",
                    key=f"modal_exp_input_{code}"
                )
                updated_expiry_map[code] = val_str.strip()
        if st.button("💾 수정한 API 만료일 저장", use_container_width=True):
            save_api_expiry_dates(updated_expiry_map)
            st.success("✅ API 만료일이 저장되었습니다!")
            st.rerun()

    with st.expander("⚙️ 관리자(Manager) Google 계정 또는 보안 PIN 변경"):
        with st.form("mgr_change_cred_form", clear_on_submit=False):
            new_mgr_em = st.text_input("변경할 관리자 Google 이메일", value=registered_email)
            new_mgr_pw = st.text_input("새 관리자 보안 PIN (4자리 이상)", type="password")
            if st.form_submit_button("🔄 관리자 보안 정보 업데이트"):
                if "@" not in (new_mgr_em or "").strip():
                    st.error("유효한 Google 이메일을 입력해 주세요.")
                elif len((new_mgr_pw or "").strip()) < 4:
                    st.error("새 보안 PIN을 4자리 이상 입력해 주세요.")
                else:
                    _save_manager_google_config(new_mgr_em.strip().lower(), new_mgr_pw.strip())
                    st.session_state.manager_google_authed_email = new_mgr_em.strip().lower()
                    st.session_state.manager_manual_locked = False
                    st.success("✅ 관리자 Google 계정 및 보안 PIN이 변경되었습니다.")

if _dialog_decorator:
    @_dialog_decorator("📜 법제처 공식 『제·개정이유 및 주요내용』 원문 뷰어", width="large")
    def open_official_reason_modal(
        law_prom_map: dict[str, dict],
        effective_ws_id: str,
        selected_room: str,
        pro_toggle_on: bool,
        budget_exceeded: bool = False,
        preselected_clean_name: str = ""
    ):
        _render_official_reason_viewer_body(
            law_prom_map, effective_ws_id, selected_room, pro_toggle_on, budget_exceeded, preselected_clean_name
        )

    @_dialog_decorator("🔑 연동 API 현황 및 보안키·만료일 통합 관리", width="large")
    def open_api_key_manager_modal(expiry_all_status: list[dict], law_prom_map: dict):
        _render_api_key_manager_body(expiry_all_status, law_prom_map)

# ---------------------------------------------------------
# 2. 왼쪽 사이드바 제어판 (접속 PC/브라우저별 100% 독립 세션 분리 + 1~10번 방 무손실 동기화 + 옵션 C 오로라 로딩 LED)
# ---------------------------------------------------------
task_mgr = get_background_task_manager()
expiry_alerts, expiry_all_status = check_api_expiry_alerts()
law_prom_map = fetch_latest_tax_law_promulgations()
if hasattr(agent, "register_prewarmed_law_ids") and law_prom_map:
    agent.register_prewarmed_law_ids(law_prom_map)
budget_exceeded = False

# 🛡️ [접속 PC/브라우저별 고유 익명 작업공간 ID 발급]
# 비밀번호를 입력하지 않고 접속하더라도 다른 PC나 다른 브라우저 접속자와 대화방(1~10번)이 절대 섞이지 않도록
# 각 브라우저 세션마다 고유한 독립 작업공간 ID(`__anon_session__:<uuid>`)를 부여합니다.
if "anon_workspace_id" not in st.session_state:
    st.session_state.anon_workspace_id = f"{ANON_WS_PREFIX}{uuid.uuid4().hex[:16]}"
    _cleanup_stale_anon_workspaces(max_age_sec=86400)

if "app_top_mode" not in st.session_state:
    st.session_state.app_top_mode = "AGENT"

def _detect_any_room_button_click() -> None:
    """
    [클릭 씹힘 원천 방지]
    백그라운드 작업 완료로 인해 버튼 key의 접미사(_running -> _hasmsg 등)가 바뀌는 순간에
    사용자가 대화방 버튼을 클릭하더라도, 브라우저가 보낸 이전 상태 key 클릭 신호를 100% 포착하여 즉시 방을 전환합니다.
    """
    for idx_r, r_id in enumerate(ROOM_IDS, start=1):
        for sfx in ("empty", "hasmsg", "active", "empty_running", "hasmsg_running", "active_running"):
            cand_k = f"room_btn_{idx_r}_{sfx}"
            if st.session_state.get(cand_k):
                st.session_state[cand_k] = False
                if st.session_state.get("current_room") != r_id:
                    st.session_state.current_room = r_id

_detect_any_room_button_click()

with st.sidebar:
    # 🔀 [최상단 고정: "자율 에이전트AI" / "통합고용세액공제" 2열 전환 버튼]
    sb_m1, sb_m2 = st.columns(2, gap="small")
    with sb_m1:
        if st.button(
            "자율 에이전트AI",
            key="btn_top_mode_agent",
            type=("primary" if st.session_state.app_top_mode == "AGENT" else "secondary"),
            use_container_width=True,
        ):
            st.session_state.app_top_mode = "AGENT"
            st.rerun()
    with sb_m2:
        if st.button(
            "통합고용세액공제",
            key="btn_top_mode_credit",
            type=("primary" if st.session_state.app_top_mode == "CREDIT_CALC" else "secondary"),
            use_container_width=True,
        ):
            st.session_state.app_top_mode = "CREDIT_CALC"
            st.rerun()

    st.markdown(
        "<div style='font-size:1.00rem; font-weight:800; color:#0f172a; line-height:1.4; margin-top:10px; margin-bottom:6px; padding-left:2px; letter-spacing:-0.3px;'>"
        "⚙️ 에이전트 제어판</div>",
        unsafe_allow_html=True
    )
    
    # 📦 [카드 1] 엔진 출력 모드 제어 (1번 아르곤 직결 방식)
    with st.container(border=True):
        is_pro_now = st.session_state.get("pro_toggle_state", False)
        mode_pill = (
            "<span class='sb-badge-pro'>👑 Argon 직결 (3배 개방)</span>"
            if is_pro_now else
            "<span class='sb-badge-green'>⚡ Flash 고속</span>"
        )
        st.markdown(f"<div class='sb-card-title'><span>🚀 엔진 출력 모드</span>{mode_pill}</div>", unsafe_allow_html=True)
        pro_toggle_on = st.toggle(
            "Pro 심층 모드 (Argon 직결 · 3배 개방)",
            value=is_pro_now,
            help="켜두면 최상위 Gemini 4 Argon 엔진이 1순위로 직결 가동되며 최대 15만 자 문서 독해 및 데이터·도구 한도 3배 개방 모드로 작동합니다."
        )
        st.session_state.pro_toggle_state = pro_toggle_on

    # 📦 [카드 2] 내 작업공간 (비밀번호 미입력 시 접속 PC/세션별 1~10번 방 완벽 분리 + 암호 입력 시 영구 동기화 및 API키 2단 자동패스 연동)
    with st.container(border=True):
        ws_widget_ver = st.session_state.get("ws_widget_ver", 0)
        saved_ws_val = st.session_state.get("active_workspace_id", "")
        ws_pill = (
            "<span class='sb-badge-green'>🟢 보안 저장중</span>"
            if saved_ws_val.strip() else
            "<span class='sb-badge-amber'>🛡️ 개별 분리(임시)</span>"
        )
        st.markdown(f"<div class='sb-card-title'><span>🔐 내 작업공간 코드</span>{ws_pill}</div>", unsafe_allow_html=True)
        workspace_input = st.text_input(
            "작업공간 코드",
            value=saved_ws_val,
            type="password",
            placeholder="나만의 암호 입력 후 Enter",
            label_visibility="collapsed",
            key=f"ws_input_box_{ws_widget_ver}",
            help="암호를 입력하지 않아도 각 접속 PC/브라우저마다 1~10번 대화방이 타인과 완벽히 분리됩니다. 여기에 관리자 보안 PIN(MANAGER_SECURITY_PIN)을 입력해 두시면 하단 [🔑 API키]도 별도 로그인 없이 즉시 열립니다."
        )
        effective_ws_id = (
            workspace_input.strip()
            if workspace_input.strip()
            else st.session_state.anon_workspace_id
        )

        if st.session_state.get("loaded_workspace_id") != effective_ws_id:
            ws_data = load_workspace_from_disk(effective_ws_id)
            st.session_state.chat_rooms = ws_data["chat_rooms"]
            st.session_state.room_titles = ws_data["room_titles"]
            st.session_state.room_files = ws_data["room_files"]
            st.session_state.room_diff_banners = ws_data.get("room_diff_banners", {r_id: "" for r_id in ROOM_IDS})
            st.session_state.loaded_workspace_id = effective_ws_id
            st.session_state.active_workspace_id = workspace_input.strip()
            st.session_state.manager_manual_locked = False
            if "current_room" not in st.session_state:
                st.session_state.current_room = ROOM_IDS[0]
            sync_all_rooms_state(effective_ws_id)
            st.rerun()

        # 매 렌더링 시점마다 현재 세션(개별 분리 세션 또는 암호 작업공간)의 1~10번 방을 무손실 동기화
        sync_all_rooms_state(effective_ws_id)

        if workspace_input.strip():
            if st.button("🔒 작업공간 잠금 (나가기)", use_container_width=True):
                st.session_state.active_workspace_id = ""
                st.session_state.loaded_workspace_id = None
                st.session_state.manager_google_authed_email = ""
                st.session_state.ws_widget_ver = ws_widget_ver + 1
                st.rerun()

    # 📦 [통합고용세액공제 모드 전용] 귀속사업연도 · 사업자정보 (내 작업공간 바로 아래)
    if st.session_state.app_top_mode == "CREDIT_CALC":
        import employment_credit
        employment_credit.render_credit_sidebar()

    # 📦 [카드 3] 멀티 대화방 (입체 비우기 버튼 + 1~10번 좌측 정렬 리스트 + 옵션 C 제미나이 오로라 로딩 LED)
    with st.container(border=True):
        if "current_room" not in st.session_state or st.session_state.current_room not in ROOM_IDS:
            st.session_state.current_room = ROOM_IDS[0]
        selected_room = st.session_state.current_room
        active_room_key = f"{effective_ws_id}::{selected_room}"

        rooms_with_rev_cnt = sum(
            1 for r_id in ROOM_IDS
            if get_room_revision_alerts(effective_ws_id, r_id, st.session_state.chat_rooms.get(r_id, []), law_prom_map)[0]
        )
        rev_header_pill = (
            f"<span class='sb-badge-amber'>🔄 {rooms_with_rev_cnt}개방 개정</span>"
            if rooms_with_rev_cnt > 0 else ""
        )
        st.markdown(f"<div class='sb-card-title'><span>🗂️ 멀티 대화방 (1~10번)</span>{rev_header_pill}</div>", unsafe_allow_html=True)

        room_search_kw = st.text_input(
            "전체 대화방 통합 검색",
            placeholder="🔎 1~10번 방 통합 키워드 검색",
            label_visibility="collapsed",
            key=f"room_search_{effective_ws_id}"
        )
        if room_search_kw.strip():
            kw_lower = room_search_kw.strip().lower()
            matched_rooms = []
            for r_id in ROOM_IDS:
                r_title = st.session_state.room_titles.get(r_id, "")
                r_fname = st.session_state.room_files[r_id].get("name") or ""
                r_msgs = st.session_state.chat_rooms.get(r_id, [])
                hit_cnt = sum(1 for m in r_msgs if kw_lower in m.get("content", "").lower())
                if kw_lower in r_title.lower() or kw_lower in r_fname.lower():
                    hit_cnt += 1
                if hit_cnt > 0:
                    matched_rooms.append((r_id, r_title, hit_cnt))
            if matched_rooms:
                for m_rid, m_title, m_cnt in matched_rooms:
                    if st.button(f"👉 {m_rid} [{m_title}] ({m_cnt}건 일치)", key=f"jump_{m_rid}", use_container_width=True):
                        st.session_state.current_room = m_rid
                        st.rerun()
            else:
                st.caption(f"'{room_search_kw}' 검색 결과가 없습니다.")

        title_widget_key = f"title_input_{effective_ws_id}_{selected_room}"
        # 위젯(st.text_input)이 렌더링되기 전에만 상태 동기화를 수행합니다.
        if (
            title_widget_key in st.session_state
            and st.session_state[title_widget_key] in ("빈 대화방", "메인 대화")
            and st.session_state.room_titles[selected_room] not in ("빈 대화방", "메인 대화")
            and st.session_state.chat_rooms.get(selected_room)
        ):
            st.session_state[title_widget_key] = st.session_state.room_titles[selected_room]

        new_title = st.text_input(
            "방 주제 메모",
            value=st.session_state.room_titles[selected_room],
            placeholder="✏️ 현재 방 주제 메모 (예: A사 세무조사)",
            label_visibility="collapsed",
            key=title_widget_key
        )
        if new_title != st.session_state.room_titles[selected_room]:
            st.session_state.room_titles[selected_room] = new_title
            save_workspace_to_disk(effective_ws_id)
            st.rerun()

        if st.button(f"🗑️ [{selected_room}] 현재 방 비우기", key="btn_clear_current_room", use_container_width=True):
            with task_mgr["lock"]:
                task_mgr["tasks"].pop(active_room_key, None)
            st.session_state.chat_rooms[selected_room] = []
            st.session_state.room_files[selected_room] = {"name": None, "text": "", "items": []}
            st.session_state.room_titles[selected_room] = "빈 대화방"
            # 이미 인스턴스화된 위젯 키는 값을 대입하지 않고 pop()으로 제거하여 StreamlitAPIException 원천 방지
            st.session_state.pop(title_widget_key, None)
            st.session_state.setdefault("room_diff_banners", {})[selected_room] = ""
            st.session_state[f"diff_banner_open_{effective_ws_id}_{selected_room}"] = False
            if hasattr(agent, "reset_room_usage"):
                agent.reset_room_usage(active_room_key)
            st.session_state[f"reset_cnt_{selected_room}"] = st.session_state.get(f"reset_cnt_{selected_room}", 0) + 1
            save_workspace_to_disk(effective_ws_id, force_overwrite_room=selected_room)
            st.rerun()

        for idx_r, r_id in enumerate(ROOM_IDS, start=1):
            r_title = (st.session_state.room_titles.get(r_id, "빈 대화방") or "빈 대화방").strip()
            short_title = r_title[:14] + ("…" if len(r_title) > 14 else "")
            r_msgs = st.session_state.chat_rooms.get(r_id, [])
            q_cnt = (len(r_msgs) + 1) // 2
            f_cnt = len(st.session_state.room_files[r_id].get("items", []))
            t_info = task_mgr["tasks"].get(f"{effective_ws_id}::{r_id}")
            is_run = bool(t_info and t_info.get("status") == "running")
            rev_alerts, _ = get_room_revision_alerts(effective_ws_id, r_id, r_msgs, law_prom_map)

            base_state = "active" if r_id == selected_room else ("hasmsg" if q_cnt > 0 else "empty")
            run_suffix = "_running" if is_run else ""
            btn_key = f"room_btn_{idx_r}_{base_state}{run_suffix}"

            extra_badges = (" 📎" if f_cnt else "") + (" 🔄" if rev_alerts else "") + (" (작성중)" if is_run else "")
            btn_label = f"{idx_r}. {short_title} ({q_cnt}건){extra_badges}"

            if st.button(btn_label, key=btn_key, use_container_width=True):
                if r_id != selected_room:
                    st.session_state.current_room = r_id
                    st.rerun()

    # 📦 [카드 4 — 최하단 고정: [💳 결제] 및 [🔑 API키] 버튼]
    with st.container(border=True):
        col_b1, col_b2 = st.columns(2)
        with col_b1:
            st.link_button("💳 결제", "https://console.cloud.google.com/billing", use_container_width=True)
        with col_b2:
            if st.button("🔑 API키", use_container_width=True, help="연동된 전체 API 리스트 및 갱신 링크 열기"):
                if _dialog_decorator:
                    open_api_key_manager_modal(expiry_all_status, law_prom_map)
                else:
                    st.session_state.show_inline_api_manager = not st.session_state.get("show_inline_api_manager", False)

# ---------------------------------------------------------
# 3. 메인 화면 상단 (슬림 상태바 + 법령 개정 배너 버튼: '📜 개정이유서' / '⚡ 신·구법 비교')
# ---------------------------------------------------------
# 사이드바 렌더링 도중 완료된 백그라운드 작업이 있을 경우에도 즉시 반영
sync_all_rooms_state(effective_ws_id)

current_messages = st.session_state.chat_rooms[selected_room]
current_room_title = st.session_state.room_titles[selected_room]
current_file_info = _normalize_room_file_entry(st.session_state.room_files[selected_room])
current_task_key = f"{effective_ws_id}::{selected_room}"
active_task = task_mgr["tasks"].get(current_task_key)
is_current_room_running = bool(active_task and active_task.get("status") == "running")

if st.session_state.app_top_mode == "CREDIT_CALC":
    if st.session_state.get("show_inline_api_manager"):
        with st.container(border=True):
            st.subheader("🔑 연동 API 현황 및 보안키·만료일 통합 관리")
            _render_api_key_manager_body(expiry_all_status, law_prom_map)
            if st.button("✖️ API 관리창 닫기"):
                st.session_state.show_inline_api_manager = False
                st.rerun()
    import employment_credit
    employment_credit.render_credit_wizard(
        extract_file_fn=extract_file_content,
        log_tokens_fn=_log_custom_tokens,
        active_room_key=current_task_key,
    )
    st.stop()

# 상단 미니 상태바
q_count_now = (len(current_messages) + 1) // 2
mode_badge_html = (
    "👑 <b>Pro 심층 모드 — Argon 직결 (3배 리미트 개방 ON)</b>"
    if pro_toggle_on else
    "⚡ <b>Flash 표준 모드 (고속 검토)</b>"
)
st.markdown(
    f"<div class='top-mini-status-bar'>"
    f"<span>💬 <b>{selected_room} [{current_room_title}]</b> ({q_count_now}건 대화 중)</span>"
    f"<span>💡 현재 엔진: {mode_badge_html}</span>"
    f"</div>",
    unsafe_allow_html=True
)

# 🚨 1) [API 유효기간 1개월(30일) 전 자동 경고 배너]
if expiry_alerts:
    for alert_item in expiry_alerts:
        c_al1, c_al2 = st.columns([5, 1])
        with c_al1:
            if alert_item["level"] == "expired":
                st.error(alert_item["msg"])
            else:
                st.warning(alert_item["msg"])
        with c_al2:
            st.link_button("🔗 연장·재발급", alert_item["url"], use_container_width=True)

# 🔄 2) [개정세법 안내 배너]
current_room_rev_alerts, is_post_revision = get_room_revision_alerts(
    effective_ws_id, selected_room, current_messages, law_prom_map
)
saved_diff_md = st.session_state.get("room_diff_banners", {}).get(selected_room, "")
is_diff_open = st.session_state.get(f"diff_banner_open_{effective_ws_id}_{selected_room}", False)

if current_room_rev_alerts:
    rev_names_str = ", ".join(
        f"**「{it['law_name']}」**(공포:{it['prom_date']}/시행:{it['enf_date']})" for it in current_room_rev_alerts
    )
    first_clean = current_room_rev_alerts[0]["clean_name"]
    c_rv1, c_rv2, c_rv3, c_rv4 = st.columns([3.4, 1.1, 1.2, 0.7])
    with c_rv1:
        if is_post_revision:
            st.warning(
                f"🔄 **[사후 법령 개정 감지]** 이 방의 과거 검토 이후에 {rev_names_str} 규정이 새로 개정 공포되었습니다."
            )
        else:
            st.info(
                f"📜 **[최근 공포·시행예정 법령 인용 안내]** 이 방에서 인용된 {rev_names_str}의 공식 입법이유서 및 신·구 비교표를 바로 열람할 수 있습니다."
            )
    with c_rv2:
        if st.button("📜 개정이유서", key=f"rsn_{selected_room}", use_container_width=True):
            if _dialog_decorator:
                open_official_reason_modal(law_prom_map, effective_ws_id, selected_room, pro_toggle_on, False, first_clean)
            else:
                st.session_state.show_inline_reason_viewer = not st.session_state.get("show_inline_reason_viewer", False)
                st.rerun()
    with c_rv3:
        diff_btn_lbl = "🔼 신·구법 비교 닫기" if (saved_diff_md and is_diff_open) else "⚡ 신·구법 비교"
        if st.button(diff_btn_lbl, key=f"rerev_{selected_room}", use_container_width=True):
            if saved_diff_md:
                st.session_state[f"diff_banner_open_{effective_ws_id}_{selected_room}"] = not is_diff_open
                st.rerun()
            else:
                with st.spinner("⚡ 본문 대화와 분리된 상단 [조문명 / 구법내용 / 개정법내용 / 개정효과] 비교표를 생성 중입니다..."):
                    diff_md = generate_banner_law_diff_table(
                        current_messages, current_room_rev_alerts, room_key=current_task_key
                    )
                    st.session_state.setdefault("room_diff_banners", {})[selected_room] = diff_md
                    st.session_state[f"diff_banner_open_{effective_ws_id}_{selected_room}"] = True
                    save_workspace_to_disk(effective_ws_id)
                st.rerun()
    with c_rv4:
        if st.button("✅ 확인", key=f"ack_{selected_room}", use_container_width=True, help="이 대화방의 법령 알림을 확인 처리하여 숨깁니다."):
            acknowledge_room_laws(effective_ws_id, selected_room, current_room_rev_alerts)
            st.rerun()

if saved_diff_md:
    with st.expander("📊 [신·구법 비교] 조문명 / 구법내용 / 개정법내용 / 개정효과 (클릭하여 열기·닫기)", expanded=is_diff_open):
        st.markdown(f"<div class='law-diff-banner-box'>\n\n{_format_markdown_for_screen(saved_diff_md)}\n\n</div>", unsafe_allow_html=True)
        b_col1, b_col2 = st.columns([1, 5])
        with b_col1:
            if st.button("🔄 비교표 다시 생성", key=f"regen_diff_{selected_room}"):
                with st.spinner("⚡ 신·구법 비교표를 새로 작성 중입니다..."):
                    alerts_for_regen = current_room_rev_alerts or extract_cited_laws_in_room(current_messages, law_prom_map)
                    diff_md = generate_banner_law_diff_table(
                        current_messages, alerts_for_regen, room_key=current_task_key
                    )
                    st.session_state.setdefault("room_diff_banners", {})[selected_room] = diff_md
                    st.session_state[f"diff_banner_open_{effective_ws_id}_{selected_room}"] = True
                    save_workspace_to_disk(effective_ws_id)
                st.rerun()

if st.session_state.get("show_inline_reason_viewer"):
    with st.container(border=True):
        st.subheader("📜 법제처 공식 『제·개정이유 및 주요내용』 원문 뷰어")
        _render_official_reason_viewer_body(law_prom_map, effective_ws_id, selected_room, pro_toggle_on, False)
        if st.button("✖️ 뷰어 닫기"):
            st.session_state.show_inline_reason_viewer = False
            st.rerun()

if st.session_state.get("show_inline_api_manager"):
    with st.container(border=True):
        st.subheader("🔑 연동 API 현황 및 보안키·만료일 통합 관리")
        _render_api_key_manager_body(expiry_all_status, law_prom_map)
        if st.button("✖️ API 관리창 닫기"):
            st.session_state.show_inline_api_manager = False
            st.rerun()

# ---------------------------------------------------------
# 4. 대화방 메시지 목록 렌더링 (방 전환 시 이전 방의 고스트 잔상 원천 차단용 고유 키 컨테이너)
# ---------------------------------------------------------
room_view_Slot = st.empty()
with room_view_Slot.container(key=f"room_main_view_{effective_ws_id}_{selected_room}"):
    if not current_messages and not is_current_room_running:
        st.markdown(
            f"<div style='padding: 18px 20px; background-color: #f8fafc; border: 1px dashed #cbd5e1; "
            f"border-radius: 10px; color: #64748b; font-size: 0.96rem; margin-bottom: 12px;'>"
            f"✨ <b>{selected_room} [{current_room_title}]</b> — 아직 작성된 대화가 없습니다. "
            f"아래 입력창에 새로운 질문을 입력하거나 파일을 첨부해 보세요."
            f"</div>",
            unsafe_allow_html=True
        )

    for idx, message in enumerate(current_messages):
        with st.chat_message(message["role"]):
            display_content = message["content"]
            if message["role"] == "assistant" and hasattr(agent, "_inject_direct_law_links"):
                try:
                    display_content = agent._inject_direct_law_links(display_content)
                except Exception:
                    pass
            display_content = _format_markdown_for_screen(display_content)
            st.markdown(display_content, unsafe_allow_html=True)

            if message["role"] == "assistant":
                excel_bytes = create_excel_from_markdown_tables(message["content"])
                col1, col2, col3, col4 = st.columns(4)
                with col1:
                    st.download_button(
                        label="📕 정식 PDF (.pdf)",
                        data=create_pdf_report(message["content"]),
                        file_name=f"검토보고서_{selected_room}_{current_room_title}_{idx}.pdf",
                        mime="application/pdf",
                        key=f"pdf_{effective_ws_id}_{selected_room}_{idx}",
                        use_container_width=True
                    )
                with col2:
                    st.download_button(
                        label="📘 워드 (.docx)",
                        data=create_word_docx(message["content"]),
                        file_name=f"검토보고서_{selected_room}_{current_room_title}_{idx}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        key=f"docx_{effective_ws_id}_{selected_room}_{idx}",
                        use_container_width=True
                    )
                with col3:
                    if excel_bytes:
                        st.download_button(
                            label="📊 표 엑셀 (.xlsx)",
                            data=excel_bytes,
                            file_name=f"검토표_{selected_room}_{current_room_title}_{idx}.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            key=f"xlsx_{effective_ws_id}_{selected_room}_{idx}",
                            use_container_width=True
                        )
                    else:
                        st.button("📊 표 없음 (.xlsx)", disabled=True, key=f"no_xlsx_{effective_ws_id}_{selected_room}_{idx}", use_container_width=True)
                with col4:
                    st.download_button(
                        label="📝 마크다운 (.md)",
                        data=message["content"].encode("utf-8-sig"),
                        file_name=f"검토보고서_{selected_room}_{current_room_title}_{idx}.md",
                        mime="text/markdown",
                        key=f"md_{effective_ws_id}_{selected_room}_{idx}",
                        use_container_width=True
                    )

                act1, act2, act3 = st.columns(3)
                with act1:
                    if st.button(
                        "💬 카톡 안내문 변환 (Flash)",
                        key=f"kakao_{effective_ws_id}_{selected_room}_{idx}",
                        disabled=is_current_room_running,
                        use_container_width=True
                    ):
                        _trigger_agent_task(
                            effective_ws_id, selected_room,
                            (
                                "방금 작성한 검토 결과를 바탕으로, 거래처 대표님께 카카오톡으로 바로 복사해서 보낼 수 있는 안내문을 작성해 줘. "
                                "[필수 규칙] "
                                "1) 별표(**굵게**), 샵(###), 마크다운 표(| --- |) 등 마크다운 기호는 단 한 개도 쓰지 말고 100% 순수 텍스트와 하이픈(- )만 사용할 것. "
                                "2) 법령명이나 조문번호(제몇조 등), 판례번호는 아예 넣지 말 것. "
                                "3) 어려운 세무 용어를 쓰지 말고 대표님이 스마트폰으로 5초 만에 이해할 수 있게 아주 쉽고 단순한 일상 표현으로 [한줄 결론] -> [핵심 금액/내용] -> [꼭 챙기실 점] 3단 구성(전체 10줄 내외)으로 짧고 친절하게 작성할 것."
                            ),
                            pro_toggle_on, False, force_flash=True
                        )
                with act2:
                    if st.button(
                        "📧 이메일 회신 초안 (Flash)",
                        key=f"email_{effective_ws_id}_{selected_room}_{idx}",
                        disabled=is_current_room_running,
                        use_container_width=True
                    ):
                        _trigger_agent_task(
                            effective_ws_id, selected_room,
                            "방금 작성한 검토 결과를 바탕으로, 거래처에 발송할 정식 비즈니스 이메일 회신 초안(제목, 인사말, 핵심 검토 결과 요약, 실무 유의사항, 맺음말 포함)을 깔끔하게 작성해 줘.",
                            pro_toggle_on, False, force_flash=True
                        )
                with act3:
                    if st.button(
                        "👑 [프로/Argon] 3배 개방 재검토",
                        key=f"repro_{effective_ws_id}_{selected_room}_{idx}",
                        disabled=is_current_room_running,
                        use_container_width=True
                    ):
                        _trigger_agent_task(
                            effective_ws_id, selected_room,
                            "[프로] 직전 질문 사안에 대해 관련 법령·시행령·판례·심판례·예규와 예외 요건까지 데이터·도구 한도를 3배 개방하여 심층 재검토해 줘.",
                            pro_toggle_on, False, force_flash=False
                        )

                with st.expander("📋 카톡·이메일 전송용 텍스트 원클릭 복사 (우측 상단 복사 아이콘 클릭)"):
                    st.code(_extract_clean_copy_text(message["content"]), language="markdown")

# ---------------------------------------------------------
# 5. 메인 화면 최하단 고정: Gemini 스타일 올인원 통합 채팅 입력창 (드래그 앤 드롭 파일 첨부 통합)
# ---------------------------------------------------------
room_file_entry = _normalize_room_file_entry(st.session_state.room_files[selected_room])
st.session_state.room_files[selected_room] = room_file_entry
saved_items = room_file_entry["items"]
reset_cnt = st.session_state.get(f"reset_cnt_{selected_room}", 0)

if saved_items:
    with st.container(border=True):
        c_head1, c_head2 = st.columns([8.5, 1.5])
        with c_head1:
            chips_str = " &nbsp;|&nbsp; ".join(
                f"📄 **{it['name']}** *({len(it['text']):,}자)*" for it in saved_items
            )
            st.markdown(
                f"<div style='font-size:0.88rem; color:#15803d;'>📎 <b>현재 [{selected_room}] 첨부 유지 중 ({len(saved_items)}건):</b> {chips_str}</div>",
                unsafe_allow_html=True
            )
        with c_head2:
            if st.button("🗑️ 첨부 비우기", key=f"clear_all_f_{selected_room}", use_container_width=True):
                st.session_state.room_files[selected_room] = {"name": None, "text": "", "items": []}
                st.session_state[f"reset_cnt_{selected_room}"] = reset_cnt + 1
                save_workspace_to_disk(effective_ws_id)
                st.rerun()

chat_input_placeholder = (
    f"⏳ [{selected_room}] 현재 백그라운드 정밀 검토 진행 중입니다... (좌측 다른 대화방을 클릭해 추가 질문을 하실 수 있습니다)"
    if is_current_room_running else
    f"[{selected_room}] 검토할 이슈를 입력하거나 이 칸에 파일(엑셀·PDF·워드·사진)을 바로 끌어다 놓으세요 (+ 버튼 첨부 가능)"
)

_chat_sig = inspect.signature(st.chat_input).parameters
try:
    if "accept_file" in _chat_sig:
        chat_val = st.chat_input(
            placeholder=chat_input_placeholder,
            disabled=is_current_room_running,
            accept_file="multiple",
            file_type=["xlsx", "xls", "csv", "pdf", "docx", "png", "jpg", "jpeg"],
            key=f"chat_input_{effective_ws_id}_{selected_room}"
        )
    else:
        chat_val = st.chat_input(
            placeholder=chat_input_placeholder,
            disabled=is_current_room_running,
            key=f"chat_input_{effective_ws_id}_{selected_room}"
        )
except Exception:
    chat_val = st.chat_input(
        placeholder=chat_input_placeholder,
        disabled=is_current_room_running,
        key=f"chat_input_{effective_ws_id}_{selected_room}"
    )

if chat_val:
    if isinstance(chat_val, str):
        user_prompt_text = chat_val
        inline_files = []
    else:
        user_prompt_text = getattr(chat_val, "text", "") or ""
        inline_files = list(getattr(chat_val, "files", []) or [])

    if inline_files:
        existing_names = {it["name"] for it in saved_items}
        for uf in inline_files:
            if uf.name not in existing_names:
                with st.spinner(f"📄 '{uf.name}' 내용 추출 및 저장 중..."):
                    extracted_text = extract_file_content(uf, room_key=active_room_key)
                    saved_items.append({"name": uf.name, "text": extracted_text})
                    existing_names.add(uf.name)
        st.session_state.room_files[selected_room] = _normalize_room_file_entry({"items": saved_items})
        save_workspace_to_disk(effective_ws_id)

    if user_prompt_text.strip():
        _trigger_agent_task(
            effective_ws_id=effective_ws_id,
            selected_room=selected_room,
            raw_prompt=user_prompt_text,
            pro_toggle_on=pro_toggle_on,
            budget_exceeded=False,
            force_flash=False
        )
    elif inline_files:
        st.rerun()

# ---------------------------------------------------------
# 6. [비차단(Non-Blocking) 멀티룸 백그라운드 모니터 — st.fragment(run_every=1.0)]
#    - 기존 while True / time.sleep 루프를 완전히 제거하여,
#      다른 대화방이 검색 중이더라도 화면 깜빡임(Dimming)·이전 방 잔상(Ghosting)·입력 포커스 끊김 없이
#      1~10번 어느 방으로든 0.01초 만에 즉시 전환하고 동시에 타이핑·질문할 수 있습니다!
# ---------------------------------------------------------
def _get_ws_task_states(ws_id: str) -> tuple[bool, bool, list[str]]:
    prefix = f"{ws_id}::"
    has_running = False
    has_done = False
    running_rooms = []
    for k, v in list(task_mgr["tasks"].items()):
        if k.startswith(prefix) and isinstance(v, dict):
            st_val = v.get("status")
            r_name = k.split("::", 1)[1] if "::" in k else k
            if st_val == "running":
                has_running = True
                running_rooms.append(r_name)
            elif st_val == "done":
                has_done = True
    return has_running, has_done, running_rooms

_fragment_fn = getattr(st, "fragment", getattr(st, "experimental_fragment", None))

def _render_live_task_monitor(ws_id: str, cur_room: str) -> None:
    has_running, has_done, running_rooms = _get_ws_task_states(ws_id)
    if has_done:
        sync_all_rooms_state(ws_id)
        st.rerun()
        return

    monitor_slot = st.empty()
    cur_task = task_mgr["tasks"].get(f"{ws_id}::{cur_room}")
    with monitor_slot.container(key=f"live_monitor_{ws_id}_{cur_room}"):
        if cur_task and cur_task.get("status") == "running":
            is_pro_running = cur_task.get("is_pro", False)
            status_title = (
                f"👑 [{cur_room}] Pro 심층 모드(Argon 직결 · 3배 개방) 백그라운드 가동 중..."
                if is_pro_running else
                f"⚡ [{cur_room}] 실시간 법령·공시 DB 고속 조회 및 백그라운드 정밀 검토 중..."
            )
            elapsed_sec = max(0, int(time.time() - cur_task.get("started_at", time.time())))
            cur_live_msg = cur_task.get("live_step") or "🔎 **1단계:** 질의 쟁점 분석 및 18대 주요 세법 프리워밍 캐시·공공기관 DB 탐색 중..."
            with st.chat_message("assistant"):
                with st.status(status_title, expanded=True):
                    st.markdown(f"**실시간 진행 상태 ({elapsed_sec}초 경과):** {cur_live_msg}")
                    st.write("🔎 **1단계:** 질의 쟁점 분석 및 7대 국가·공공기관(법제처 현행·구법·국세청·DART·국토부·한은) DB 1턴 직행 탐색")
                    st.write("🧮 **2단계:** 관련 법령·판례 원문 Gzip 고속 파싱 및 파이썬 1원 단위 정밀 수리 계산기 구동")
                    st.write("🛡️ **3단계:** 미조회 근거 강제 차단(`_scan_violations`) 및 듀얼 AI 고속 감수(`_semantic_verify`) 수행")
                    st.caption("💡 **백그라운드 멀티태스킹 작동 중:** 지금 좌측 대화방 리스트를 눌러 다른 대화방으로 이동하셔도 이 방의 AI 검토는 멈추지 않고 끝까지 완료됩니다.")
        elif running_rooms:
            other_str = ", ".join(running_rooms)
            st.caption(f"🌌 현재 백그라운드에서 **{other_str}** 정밀 검토가 진행 중입니다. (현재 **{cur_room}**에서는 멈춤 없이 바로 질문하실 수 있습니다)")

ws_has_running_init, ws_has_done_init, _ = _get_ws_task_states(effective_ws_id)
if ws_has_done_init:
    sync_all_rooms_state(effective_ws_id)
    st.rerun()
elif ws_has_running_init:
    if _fragment_fn is not None:
        @_fragment_fn(run_every=1.0)
        def _auto_poll_fragment():
            _render_live_task_monitor(effective_ws_id, selected_room)
        _auto_poll_fragment()
    else:
        _render_live_task_monitor(effective_ws_id, selected_room)
        time.sleep(1.0)
        st.rerun()

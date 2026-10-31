# -*- coding: utf-8 -*-
"""
인건비 자동화 - 위하고 업로드 엑셀 및 6종 급여명세서·사업소득명세서·대장 PDF/JPG 일괄 생성 엔진 (exporters.py)
1) 위하고(WEHAGO / SmartA) 업로드 엑셀:
   - 근로소득 급여 업로드 양식 (.xlsx, '사업장' 열 자동 제외 후 정식 29개 열 2단 병합 헤더 + 합계행)
   - 사업소득(3.3%) 자료입력 업로드 양식 (.xlsx, 15개 열 표준 서식)
2) 급여대장 PDF (A4 가로, 사원별 3단 행 블록 + 하단 합계 블록)
3) 인별 급여명세서 6종 양식 지원 (`근로기준1`[기본], `근로기준2`, `근로기준3`, `근로기준4`, `근로기준5`, `구`):
   - PDF 일괄 압축 (.zip, 최대 200명, 생년월일 6자리 암호 옵션, 내부 파일명: `근로자명_급여명세서_YYMM.pdf`)
   - 이미지(JPG) 일괄 압축 (.zip, 최대 200명, 내부 파일명: `근로자명_급여명세서_YYMM.jpg`)
4) 사업소득지급대장 PDF (A4 세로, 2단 헤더 + 하단 총계)
5) 인별 사업소득명세서 PDF 일괄 압축 (.zip, 최대 100명, 내부 파일명: `근로자명_사업소득명세서_YYMM.pdf`) 및
   이미지(JPG) 일괄 압축 (.zip, 최대 100명, 내부 파일명: `근로자명_사업소득명세서_YYMM.jpg`)
6) 마우스 오버(Hover) 돋보기 미리보기용 Base64 이미지 캐시 생성기
"""

from __future__ import annotations

import base64
import io
import os
import re
import urllib.request
import zipfile
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image
from pypdf import PdfReader, PdfWriter
import pypdfium2 as pdfium
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .parsers import extract_birth_6digits, rrn_to_birth_dot


PAYSLIP_TEMPLATE_OPTIONS: List[str] = [
    "근로기준1",
    "근로기준2",
    "근로기준3",
    "근로기준4",
    "근로기준5",
    "구",
]

_STORAGE_DIR = ".workspace_storage"
_REGISTERED_FONT_REGULAR = "Helvetica"
_REGISTERED_FONT_BOLD = "Helvetica-Bold"
_FONT_INIT_DONE = False
_FONT_DOWNLOAD_ATTEMPTED = False
_PREVIEW_B64_CACHE: Dict[str, str] = {}


def get_korean_font_names() -> Tuple[str, str]:
    global _REGISTERED_FONT_REGULAR, _REGISTERED_FONT_BOLD, _FONT_INIT_DONE, _FONT_DOWNLOAD_ATTEMPTED
    if _FONT_INIT_DONE:
        return _REGISTERED_FONT_REGULAR, _REGISTERED_FONT_BOLD

    reg_name = "PayKorFont"
    bold_name = "PayKorFontBold"
    registered = pdfmetrics.getRegisteredFontNames()
    if reg_name in registered and bold_name in registered:
        _REGISTERED_FONT_REGULAR = reg_name
        _REGISTERED_FONT_BOLD = bold_name
        _FONT_INIT_DONE = True
        return _REGISTERED_FONT_REGULAR, _REGISTERED_FONT_BOLD

    os.makedirs(_STORAGE_DIR, exist_ok=True)
    cached_reg = os.path.join(_STORAGE_DIR, "NanumGothic.ttf")
    cached_bold = os.path.join(_STORAGE_DIR, "NanumGothicBold.ttf")

    ttf_candidates = [
        ("C:/Windows/Fonts/malgun.ttf", "C:/Windows/Fonts/malgunbd.ttf"),
        (cached_reg, cached_bold),
        ("/usr/share/fonts/truetype/nanum/NanumGothic.ttf", "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf"),
        ("/usr/share/fonts/truetype/unfonts-core/UnDotum.ttf", "/usr/share/fonts/truetype/unfonts-core/UnDotumBold.ttf"),
        ("/System/Library/Fonts/Supplemental/AppleGothic.ttf", "/System/Library/Fonts/Supplemental/AppleGothic.ttf"),
    ]

    chosen_reg, chosen_bold = None, None
    for rp, bp in ttf_candidates:
        if os.path.exists(rp):
            chosen_reg = rp
            chosen_bold = bp if os.path.exists(bp) else rp
            break

    if chosen_reg is None and not _FONT_DOWNLOAD_ATTEMPTED:
        _FONT_DOWNLOAD_ATTEMPTED = True
        try:
            with urllib.request.urlopen(
                "https://github.com/google/fonts/raw/main/ofl/nanumgothic/NanumGothic-Regular.ttf",
                timeout=2.5,
            ) as resp:
                with open(cached_reg, "wb") as out_f:
                    out_f.write(resp.read())
            with urllib.request.urlopen(
                "https://github.com/google/fonts/raw/main/ofl/nanumgothic/NanumGothic-Bold.ttf",
                timeout=3.0,
            ) as resp:
                with open(cached_bold, "wb") as out_f:
                    out_f.write(resp.read())
            if os.path.exists(cached_reg):
                chosen_reg = cached_reg
                chosen_bold = cached_bold if os.path.exists(cached_bold) else cached_reg
        except Exception:
            pass

    if chosen_reg and os.path.exists(chosen_reg):
        try:
            pdfmetrics.registerFont(TTFont(reg_name, chosen_reg))
            pdfmetrics.registerFont(TTFont(bold_name, chosen_bold if (chosen_bold and os.path.exists(chosen_bold)) else chosen_reg))
            pdfmetrics.registerFontFamily(
                reg_name,
                normal=reg_name,
                bold=bold_name,
                italic=reg_name,
                boldItalic=bold_name,
            )
            _REGISTERED_FONT_REGULAR = reg_name
            _REGISTERED_FONT_BOLD = bold_name
            _FONT_INIT_DONE = True
            return _REGISTERED_FONT_REGULAR, _REGISTERED_FONT_BOLD
        except Exception:
            pass

    try:
        pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))
        _REGISTERED_FONT_REGULAR = "HYGothic-Medium"
        _REGISTERED_FONT_BOLD = "HYGothic-Medium"
    except Exception:
        pass

    _FONT_INIT_DONE = True
    return _REGISTERED_FONT_REGULAR, _REGISTERED_FONT_BOLD


def _fmt_won(val: Any) -> str:
    try:
        n = int(round(float(val or 0)))
    except (ValueError, TypeError):
        n = 0
    return f"{n:,}" if n != 0 else ""


def _fmt_won_zero(val: Any) -> str:
    try:
        n = int(round(float(val or 0)))
    except (ValueError, TypeError):
        n = 0
    return f"{n:,}"


def _parse_ym_parts(year_month: str) -> Tuple[int, int]:
    m = re.match(r"^(\d{4})\D*(\d{1,2})", str(year_month or ""))
    if m:
        return int(m.group(1)), int(m.group(2))
    now = datetime.now()
    return now.year, now.month


# ─────────────────────────────────────────────────────────────
# 1. 위하고(WEHAGO) 근로소득 급여 업로드 엑셀 (.xlsx - 정식 29열, '사업장' 열 자동 제외)
# ─────────────────────────────────────────────────────────────

WEHAGO_PAYROLL_ROW1 = [
    "사원코드", "사원명", "부서", "직급", "직종",
    "수당", "", "", "", "", "", "", "", "", "", "", "",
    "공제", "", "", "", "", "", "", "", "", "", "",
    "차인지급액",
]

WEHAGO_PAYROLL_ROW2 = [
    "", "", "", "", "",
    "기본급", "상여", "식대", "자가운전보조금", "약정연장근로수당", "약정휴일근로수당",
    "연장근로수당", "휴일근로수당", "결근공제", "조정금액", "미사용연차수당", "지급액계",
    "국민연금", "건강보험", "고용보험", "장기요양보험료", "소득세", "지방소득세",
    "학자금상환액", "건강보험료정산", "장기요양보험료정산", "국민연금정산", "공제액계",
    "",
]

WEHAGO_PAYROLL_FIELD_KEYS = [
    "기본급", "상여", "식대", "자가운전보조금", "약정연장근로수당", "약정휴일근로수당",
    "연장근로수당", "휴일근로수당", "결근공제", "조정금액", "미사용연차수당", "지급액계",
    "국민연금", "건강보험", "고용보험", "장기요양보험료", "소득세", "지방소득세",
    "학자금상환액", "건강보험료정산", "장기요양보험료정산", "국민연금정산", "공제액계",
]


def export_wehago_payroll_excel(employees: List[Dict[str, Any]]) -> bytes:
    """
    위하고 근로소득 급여자료 업로드용 정식 엑셀 생성 (29컬럼 2단 병합 헤더 + 합계행, '사업장' 열 제외)
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "근로소득_위하고업로드"

    ws.append(WEHAGO_PAYROLL_ROW1)
    ws.append(WEHAGO_PAYROLL_ROW2)

    for col_idx in (1, 2, 3, 4, 5, 29):
        col_letter = get_column_letter(col_idx)
        ws.merge_cells(f"{col_letter}1:{col_letter}2")
    ws.merge_cells("F1:Q1")
    ws.merge_cells("R1:AB1")

    header_fill = PatternFill("solid", fgColor="E8EEF5")
    header_font = Font(name="맑은 고딕", size=10, bold=True, color="1E293B")
    body_font = Font(name="맑은 고딕", size=10, color="0F172A")
    total_fill = PatternFill("solid", fgColor="F8FAFC")
    total_font = Font(name="맑은 고딕", size=10, bold=True, color="0F172A")
    thin_side = Side(style="thin", color="CBD5E1")
    border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for r_i in (1, 2):
        for c_i in range(1, 30):
            cell = ws.cell(row=r_i, column=c_i)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_all

    totals = {k: 0 for k in WEHAGO_PAYROLL_FIELD_KEYS + ["차인지급액"]}

    for idx, emp in enumerate(employees or [], start=1):
        code_str = str(emp.get("사원코드") or f"{idx:03d}").zfill(3)
        row_vals: List[Any] = [
            code_str,
            str(emp.get("성명", "")),
            str(emp.get("부서", "")),
            str(emp.get("직급", "")),
            str(emp.get("직종", "")),
        ]
        for k in WEHAGO_PAYROLL_FIELD_KEYS:
            val = int(emp.get(k, 0) or 0)
            row_vals.append(val)
            totals[k] += val
        net_val = int(emp.get("차인지급액", 0) or 0)
        row_vals.append(net_val)
        totals["차인지급액"] += net_val

        ws.append(row_vals)
        cur_r = ws.max_row
        for c_i in range(1, 30):
            cell = ws.cell(row=cur_r, column=c_i)
            cell.font = body_font
            cell.border = border_all
            if c_i <= 5:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0"

    sum_row: List[Any] = ["합계", "", "", "", ""]
    for k in WEHAGO_PAYROLL_FIELD_KEYS:
        sum_row.append(totals[k])
    sum_row.append(totals["차인지급액"])
    ws.append(sum_row)
    last_r = ws.max_row
    for c_i in range(1, 30):
        cell = ws.cell(row=last_r, column=c_i)
        cell.fill = total_fill
        cell.font = total_font
        cell.border = border_all
        if c_i <= 5:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = "#,##0"

    for c_i in range(1, 30):
        col_letter = get_column_letter(c_i)
        ws.column_dimensions[col_letter].width = 11 if c_i <= 5 else 14

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────
# 2. 위하고(SmartA) 사업소득 업로드 엑셀 (.xlsx)
# ─────────────────────────────────────────────────────────────

def export_wehago_freelancer_excel(
    freelancers: List[Dict[str, Any]],
    default_year_month: str = "",
    default_pay_date: str = "",
) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ws.append([
        "사업소득자료입력     ※ 본 서식은 SmartA 인사급여 사업소득자료입력 메뉴로 Excel DATA를 자동 변환하기 위한 서식입니다.",
        *([""] * 14),
    ])
    guide_text = (
        "1. 귀속년월 및 지급년월일 : YYYYMM(6자리), YYYYMMDD(8자리) 형태로 입력합니다.\n"
        "2. 주민등록번호 : '-'를 제외한 13자리 숫자로 입력합니다.\n"
        "3. 소득구분 : 업종코드 6자리를 입력합니다 (예: 940909 기타자영업).\n"
        "4. 내.외국인구분 : 내국인 '0', 외국인 '1' / 연말정산 : 비대상 '0', 대상 '1'"
    )
    ws.append([guide_text, *([""] * 14)])
    headers = [
        "", "귀속년월", "지급년월일", "소득자명", "주민등록번호",
        "기본주소", "상세주소", "소득구분", "영수일자", "지급총액",
        "세율(%)", "소득세", "지방소득세", "내.외국인구분", "연말정산",
    ]
    ws.append(headers)

    ws.merge_cells("A1:O1")
    ws.merge_cells("A2:O2")
    ws.row_dimensions[2].height = 54

    header_fill = PatternFill("solid", fgColor="EDE9FE")
    header_font = Font(name="맑은 고딕", size=10, bold=True, color="1E1B4B")
    body_font = Font(name="맑은 고딕", size=10, color="0F172A")
    thin_side = Side(style="thin", color="CBD5E1")
    border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for c_i in range(1, 16):
        cell = ws.cell(row=3, column=c_i)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_all

    yyyy, mm_num = _parse_ym_parts(default_year_month)
    default_ym6 = f"{yyyy:04d}{mm_num:02d}"
    default_pay8 = re.sub(r"[^0-9]", "", str(default_pay_date or ""))
    if len(default_pay8) != 8:
        default_pay8 = f"{yyyy:04d}{mm_num:02d}25"

    for fr in freelancers or []:
        ym_digits = re.sub(r"[^0-9]", "", str(fr.get("귀속년월") or ""))
        ym6 = ym_digits[:6] if len(ym_digits) >= 6 else default_ym6

        pay_digits = re.sub(r"[^0-9]", "", str(fr.get("지급년월일") or ""))
        pay8 = pay_digits[:8] if len(pay_digits) >= 8 else default_pay8

        rrn_digits = re.sub(r"[^0-9*]", "", str(fr.get("주민등록번호") or ""))
        gross = int(fr.get("지급총액", 0) or 0)
        itax = int(fr.get("소득세", 0) or 0)
        ltax = int(fr.get("지방소득세", 0) or 0)

        # 업종코드 6자리만 추출 (예: '940909 (기타자영업)' -> '940909')
        raw_ind = str(fr.get("소득구분") or "940909")
        m_ind = re.search(r"\b(\d{6})\b", raw_ind)
        clean_ind = m_ind.group(1) if m_ind else "940909"

        row_vals = [
            "",
            ym6,
            pay8,
            str(fr.get("성명", "")),
            rrn_digits,
            str(fr.get("기본주소", "")),
            str(fr.get("상세주소", "")),
            clean_ind,
            pay8,
            gross,
            int(fr.get("세율", 3) or 3),
            itax,
            ltax,
            "0",
            "0",
        ]
        ws.append(row_vals)
        cur_r = ws.max_row
        for c_i in range(1, 16):
            cell = ws.cell(row=cur_r, column=c_i)
            cell.font = body_font
            cell.border = border_all
            if c_i in (10, 12, 13):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0"
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    col_widths = [4, 12, 14, 14, 18, 20, 16, 12, 14, 15, 10, 14, 14, 14, 12]
    for c_i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(c_i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────
# 3. 급여대장 PDF (A4 가로 - 사원별 3단 블록 + 합계 블록)
# ─────────────────────────────────────────────────────────────

def export_payroll_ledger_pdf(
    employees: List[Dict[str, Any]],
    corp_name: str = "",
    year_month: str = "",
    pay_date: str = "",
) -> bytes:
    font_reg, font_bold = get_korean_font_names()
    yyyy, mm_num = _parse_ym_parts(year_month)
    pay_dt_fmt = pay_date.replace("-", "년 ", 1).replace("-", "월 ") + "일" if "-" in str(pay_date) else f"{yyyy}년{mm_num:02d}월25일"

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )

    title_style = ParagraphStyle("LedgerTitle", fontName=font_bold, fontSize=16, leading=20, alignment=1)
    meta_style = ParagraphStyle("LedgerMeta", fontName=font_reg, fontSize=9, leading=12, alignment=0)
    meta_right_style = ParagraphStyle("LedgerMetaRight", fontName=font_reg, fontSize=9, leading=12, alignment=2)

    elements: List[Any] = []
    elements.append(Paragraph(f"{yyyy}년 {mm_num:02d}월분 급여대장", title_style))
    elements.append(Spacer(1, 3 * mm))

    meta_tbl = Table(
        [[
            Paragraph(f"회사명 : {corp_name or '미입력'} &nbsp;&nbsp;&nbsp; 정렬 : 코드순", meta_style),
            Paragraph(f"[귀속:{yyyy}년{mm_num:02d}월] &nbsp; [지급:{pay_dt_fmt}]", meta_right_style),
        ]],
        colWidths=[138 * mm, 138 * mm],
    )
    meta_tbl.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
    ]))
    elements.append(meta_tbl)
    elements.append(Spacer(1, 2 * mm))

    table_data: List[List[str]] = [
        ["사번", "부서", "기본급", "상여", "식대", "자가운전보조금", "약정연장근로수당", "미사용연차수당", "", "", "건강보험료정산", "장기요양보험료정산", ""],
        ["사원명", "직급", "약정휴일근로수당", "연장근로수당", "휴일근로수당", "결근공제", "조정금액", "소득세", "지방소득세", "학자금상환액", "국민연금정산", "공제액계", ""],
        ["입사일", "퇴사일", "지급액계", "국민연금", "건강보험", "장기요양보험료", "고용보험", "", "", "", "", "차인지급액", ""],
    ]

    totals = {k: 0 for k in WEHAGO_PAYROLL_FIELD_KEYS + ["차인지급액"]}

    for idx, emp in enumerate(employees or [], start=1):
        for k in WEHAGO_PAYROLL_FIELD_KEYS:
            totals[k] += int(emp.get(k, 0) or 0)
        totals["차인지급액"] += int(emp.get("차인지급액", 0) or 0)

        c_code = str(emp.get("사원코드") or f"{idx:03d}")
        c_name = str(emp.get("성명", ""))
        c_hire = str(emp.get("입사일", "")).replace("-", ".")
        c_retire = str(emp.get("퇴사일", "")).replace("-", ".")
        c_dept = str(emp.get("부서", ""))
        c_pos = str(emp.get("직급", ""))

        r1 = [
            c_code, c_dept,
            _fmt_won(emp.get("기본급")), _fmt_won(emp.get("상여")),
            _fmt_won(emp.get("식대")), _fmt_won(emp.get("자가운전보조금")),
            _fmt_won(emp.get("약정연장근로수당")), _fmt_won(emp.get("미사용연차수당")),
            "", "",
            _fmt_won(emp.get("건강보험료정산")), _fmt_won(emp.get("장기요양보험료정산")), "",
        ]
        r2 = [
            c_name, c_pos,
            _fmt_won(emp.get("약정휴일근로수당")), _fmt_won(emp.get("연장근로수당")),
            _fmt_won(emp.get("휴일근로수당")), _fmt_won(emp.get("결근공제")),
            _fmt_won(emp.get("조정금액")), _fmt_won(emp.get("소득세")),
            _fmt_won(emp.get("지방소득세")), _fmt_won(emp.get("학자금상환액")),
            _fmt_won(emp.get("국민연금정산")), _fmt_won_zero(emp.get("공제액계")), "",
        ]
        r3 = [
            c_hire, c_retire,
            _fmt_won_zero(emp.get("지급액계")), _fmt_won(emp.get("국민연금")),
            _fmt_won(emp.get("건강보험")), _fmt_won(emp.get("장기요양보험료")),
            _fmt_won(emp.get("고용보험")), "", "", "", "",
            _fmt_won_zero(emp.get("차인지급액")), "",
        ]
        table_data.extend([r1, r2, r3])

    emp_cnt = len(employees or [])
    s1 = [
        "합 계", "",
        _fmt_won_zero(totals["기본급"]), _fmt_won(totals["상여"]),
        _fmt_won(totals["식대"]), _fmt_won(totals["자가운전보조금"]),
        _fmt_won(totals["약정연장근로수당"]), _fmt_won(totals["미사용연차수당"]),
        "", "",
        _fmt_won(totals["건강보험료정산"]), _fmt_won(totals["장기요양보험료정산"]), "",
    ]
    s2 = [
        f"({emp_cnt}명)", "",
        _fmt_won(totals["약정휴일근로수당"]), _fmt_won(totals["연장근로수당"]),
        _fmt_won(totals["휴일근로수당"]), _fmt_won(totals["결근공제"]),
        _fmt_won(totals["조정금액"]), _fmt_won_zero(totals["소득세"]),
        _fmt_won_zero(totals["지방소득세"]), _fmt_won(totals["학자금상환액"]),
        _fmt_won(totals["국민연금정산"]), _fmt_won_zero(totals["공제액계"]), "",
    ]
    s3 = [
        "", "",
        _fmt_won_zero(totals["지급액계"]), _fmt_won_zero(totals["국민연금"]),
        _fmt_won_zero(totals["건강보험"]), _fmt_won_zero(totals["장기요양보험료"]),
        _fmt_won_zero(totals["고용보험"]), "", "", "", "",
        _fmt_won_zero(totals["차인지급액"]), "",
    ]
    table_data.extend([s1, s2, s3])

    col_widths = [18 * mm, 16 * mm, 24 * mm, 22 * mm, 22 * mm, 22 * mm, 24 * mm, 22 * mm, 20 * mm, 20 * mm, 22 * mm, 25 * mm, 20 * mm]
    tbl = Table(table_data, colWidths=col_widths, repeatRows=3)

    cmds: List[Any] = [
        ("FONTNAME", (0, 0), (-1, -1), font_reg),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (1, -1), "CENTER"),
        ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
        ("ALIGN", (2, 0), (-1, 2), "CENTER"),
        ("BACKGROUND", (0, 0), (-1, 2), colors.HexColor("#E2E8F0")),
        ("FONTNAME", (0, 0), (-1, 2), font_bold),
        ("BOX", (0, 0), (-1, -1), 1.0, colors.HexColor("#334155")),
        ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]

    total_rows = len(table_data)
    for b_start in range(0, total_rows, 3):
        cmds.append(("LINEABOVE", (0, b_start), (-1, b_start), 0.8, colors.HexColor("#475569")))
        cmds.append(("SPAN", (11, b_start + 1), (12, b_start + 1)))
        cmds.append(("SPAN", (11, b_start + 2), (12, b_start + 2)))
        cmds.append(("NOSPLIT", (0, b_start), (-1, min(total_rows - 1, b_start + 2))))

    cmds.append(("BACKGROUND", (0, total_rows - 3), (-1, total_rows - 1), colors.HexColor("#F1F5F9")))
    cmds.append(("FONTNAME", (0, total_rows - 3), (-1, total_rows - 1), font_bold))

    tbl.setStyle(TableStyle(cmds))
    elements.append(tbl)

    doc.build(elements)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────
# 4. 인별 급여명세서 6종 양식 렌더링 (`근로기준1`~`5`, `구`) + PDF/JPG ZIP 압축
# ─────────────────────────────────────────────────────────────

def _extract_nonzero_pay_and_ded_items(emp: Dict[str, Any]) -> Tuple[List[Tuple[str, int]], List[Tuple[str, int]]]:
    """실제 금액이 0원이 아닌 지급 항목과 공제 항목만 순서대로 추출 (업로드된 샘플 6종과 동일하게 빈 줄 생략)"""
    pay_candidates = [
        ("기본급", int(emp.get("기본급", 0) or 0)),
        ("상여", int(emp.get("상여", 0) or 0)),
        ("식대", int(emp.get("식대", 0) or 0)),
        ("자가운전보조금", int(emp.get("자가운전보조금", 0) or 0)),
        ("약정연장근로수당", int(emp.get("약정연장근로수당", 0) or 0)),
        ("약정휴일근로수당", int(emp.get("약정휴일근로수당", 0) or 0)),
        ("연장근로수당", int(emp.get("연장근로수당", 0) or 0)),
        ("휴일근로수당", int(emp.get("휴일근로수당", 0) or 0)),
        ("미사용연차수당", int(emp.get("미사용연차수당", 0) or 0)),
        ("결근공제", int(emp.get("결근공제", 0) or 0)),
        ("조정금액", int(emp.get("조정금액", 0) or 0)),
    ]
    ded_candidates = [
        ("국민연금", int(emp.get("국민연금", 0) or 0)),
        ("건강보험", int(emp.get("건강보험", 0) or 0)),
        ("고용보험", int(emp.get("고용보험", 0) or 0)),
        ("장기요양보험료", int(emp.get("장기요양보험료", 0) or 0)),
        ("소득세", int(emp.get("소득세", 0) or 0)),
        ("지방소득세", int(emp.get("지방소득세", 0) or 0)),
        ("학자금상환액", int(emp.get("학자금상환액", 0) or 0)),
        ("건강보험료정산", int(emp.get("건강보험료정산", 0) or 0)),
        ("장기요양보험료정산", int(emp.get("장기요양보험료정산", 0) or 0)),
        ("국민연금정산", int(emp.get("국민연금정산", 0) or 0)),
    ]
    pay_items = [(k, v) for k, v in pay_candidates if v != 0]
    ded_items = [(k, v) for k, v in ded_candidates if v != 0]
    if not pay_items:
        pay_items = [("기본급", 0)]
    return pay_items, ded_items


def _build_calc_method_table(emp: Dict[str, Any], font_reg: str, font_bold: str, total_w_mm: float = 174.0) -> Table:
    """하단 '계산방법' 표 생성 (`근로기준1` ~ `근로기준5` 공통)"""
    calc_rows: List[List[str]] = [
        ["계산방법", "", ""],
        ["구분", "산출식 또는 산출방법", "지급액"],
    ]
    meal_pay = int(emp.get("식대", 0) or 0)
    fixed_ot = int(emp.get("약정연장근로수당", 0) or 0)
    ot_pay = int(emp.get("연장근로수당", 0) or 0)
    hol_pay = int(emp.get("휴일근로수당", 0) or 0)
    abs_ded = int(emp.get("결근공제", 0) or 0)

    if meal_pay > 0:
        calc_rows.append(["식대", "비과세식대", _fmt_won_zero(meal_pay)])
    if fixed_ot > 0:
        calc_rows.append(["약정연장근로수당", "약정 고정연장수당", _fmt_won_zero(fixed_ot)])
    if ot_pay > 0:
        ot_h = float(emp.get("연장근무시간", 0) or 0)
        hr_txt = f"{ot_h:.1f}시간" if ot_h > 0 else "연장근로시간 × 통상시급 × 1.5"
        calc_rows.append(["연장근로수당", hr_txt, _fmt_won_zero(ot_pay)])
    if hol_pay > 0:
        hol_h = float(emp.get("휴일근무시간", 0) or 0)
        hol_txt = f"{hol_h:.1f}시간" if hol_h > 0 else "휴일근로시간 × 통상시급 × 1.5"
        calc_rows.append(["휴일근로수당", hol_txt, _fmt_won_zero(hol_pay)])
    if abs_ded != 0:
        calc_rows.append(["결근공제", "결근시간 × 약정시급 차감", _fmt_won_zero(abs_ded)])

    if len(calc_rows) == 2:
        calc_rows.append(["기본급", "월 소정근로 기본급", _fmt_won_zero(emp.get("기본급", 0))])

    w1 = 42 * mm
    w3 = 42 * mm
    w2 = (total_w_mm * mm) - w1 - w3
    tbl = Table(calc_rows, colWidths=[w1, w2, w3], rowHeights=[6.8 * mm] * len(calc_rows))
    tbl.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_reg),
        ("FONTNAME", (0, 0), (-1, 1), font_bold),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("SPAN", (0, 0), (2, 0)),
        ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#D9D9D9")),
        ("ALIGN", (0, 0), (-1, 1), "CENTER"),
        ("ALIGN", (0, 2), (1, -1), "LEFT"),
        ("ALIGN", (2, 2), (2, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOX", (0, 0), (-1, -1), 0.7, colors.black),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.black),
        ("LEFTPADDING", (0, 2), (1, -1), 5),
        ("RIGHTPADDING", (2, 2), (2, -1), 5),
    ]))
    return tbl


def render_single_employee_payslip_pdf(
    emp: Dict[str, Any],
    corp_name: str = "",
    year_month: str = "",
    pay_date: str = "",
    password: Optional[str] = None,
    template_type: str = "근로기준1",
) -> bytes:
    """
    단일 근로자 급여명세서 PDF 생성 — 사용자 업로드 6종 양식(`근로기준1`[기본] ~ `근로기준5`, `구`) 완벽 구현
    """
    font_reg, font_bold = get_korean_font_names()
    yyyy, mm_num = _parse_ym_parts(year_month)
    pay_dt_dot = str(pay_date or f"{yyyy}.{mm_num:02d}.25").replace("-", ".")
    tpl = template_type if template_type in PAYSLIP_TEMPLATE_OPTIONS else "근로기준1"

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=15 * mm,
    )

    title_style = ParagraphStyle("SlipTitle", fontName=font_bold, fontSize=17, leading=22, alignment=1)
    meta_left = ParagraphStyle("SlipMetaL", fontName=font_reg, fontSize=9.5, leading=13, alignment=0)
    meta_right = ParagraphStyle("SlipMetaR", fontName=font_reg, fontSize=9.5, leading=13, alignment=2)
    unit_style = ParagraphStyle("SlipUnit", fontName=font_reg, fontSize=8.5, leading=11, alignment=2)
    footer_style = ParagraphStyle("SlipFooter", fontName=font_reg, fontSize=9.5, leading=14, alignment=1)

    emp_code = str(emp.get("사원코드") or "001")
    emp_name = str(emp.get("성명") or "")
    emp_dept = str(emp.get("부서") or "")
    emp_pos = str(emp.get("직급") or "")
    emp_birth = rrn_to_birth_dot(str(emp.get("주민등록번호") or ""))
    emp_hire = str(emp.get("입사일") or "").replace("-", ".")
    ot_h = float(emp.get("연장근무시간", 0) or 0)
    hol_h = float(emp.get("휴일근무시간", 0) or 0)
    hourly_w = int(emp.get("약정시급", 0) or 0)
    if hourly_w <= 0 and int(emp.get("기본급", 0) or 0) > 0:
        hourly_w = int(round(int(emp.get("기본급", 0)) / 209.0))

    pay_items, ded_items = _extract_nonzero_pay_and_ded_items(emp)
    gross_tot_str = _fmt_won_zero(emp.get("지급액계"))
    ded_tot_str = _fmt_won_zero(emp.get("공제액계"))
    net_tot_str = _fmt_won_zero(emp.get("차인지급액"))

    elements: List[Any] = []
    elements.append(Paragraph(f"{yyyy}년 {mm_num:02d}월분  급여명세서", title_style))
    elements.append(Spacer(1, 5 * mm))

    # ── 양식 6: '구' (구양식 - 상단 반페이지 컴팩트형 + 중앙 절취선) ──
    if tpl == "구":
        top_box_data = [
            [f"사원코드:  {emp_code}", f"사 원 명:  {emp_name}", f"입 사 일:  {emp_hire}"],
            [f"부      서:  {emp_dept}", f"직      급:  {emp_pos}", "호      봉:"],
        ]
        top_tbl = Table(top_box_data, colWidths=[58 * mm, 58 * mm, 58 * mm], rowHeights=[7 * mm, 7 * mm])
        top_tbl.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), font_reg),
            ("FONTSIZE", (0, 0), (-1, -1), 9.5),
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#D9D9D9")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BOX", (0, 0), (-1, -1), 0.7, colors.black),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ]))
        elements.append(top_tbl)

        body_rows_cnt = max(8, len(pay_items), len(ded_items))
        gu_data: List[List[str]] = [["지 급 내 역", "지 급 액", "공 제 내 역", "공 제 액"]]
        for i in range(body_rows_cnt):
            p_nm, p_v = pay_items[i] if i < len(pay_items) else ("", 0)
            d_nm, d_v = ded_items[i] if i < len(ded_items) else ("", 0)
            gu_data.append([p_nm, _fmt_won_zero(p_v) if p_nm else "", d_nm, _fmt_won_zero(d_v) if d_nm else ""])
        gu_data.append(["", "", "공 제 액 계", ded_tot_str])
        gu_data.append(["지 급 액 계", gross_tot_str, "차인지급액", net_tot_str])

        gu_tbl = Table(
            gu_data,
            colWidths=[45 * mm, 42 * mm, 45 * mm, 42 * mm],
            rowHeights=[7.5 * mm] + [6.0 * mm] * body_rows_cnt + [7.5 * mm, 7.5 * mm],
        )
        last_r = len(gu_data) - 1
        gu_tbl.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), font_reg),
            ("FONTNAME", (0, 0), (-1, 0), font_bold),
            ("FONTNAME", (0, last_r - 1), (-1, last_r), font_bold),
            ("FONTSIZE", (0, 0), (-1, -1), 9.5),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9D9D9")),
            ("BACKGROUND", (2, last_r - 1), (3, last_r - 1), colors.HexColor("#D9D9D9")),
            ("BACKGROUND", (0, last_r), (3, last_r), colors.HexColor("#D9D9D9")),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("ALIGN", (0, 1), (0, -1), "CENTER"),
            ("ALIGN", (2, 1), (2, -1), "CENTER"),
            ("ALIGN", (1, 1), (1, -1), "RIGHT"),
            ("ALIGN", (3, 1), (3, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BOX", (0, 0), (-1, -1), 0.7, colors.black),
            ("LINEBELOW", (0, 0), (-1, 0), 0.7, colors.black),
            ("LINEBEFORE", (1, 0), (1, -1), 0.5, colors.black),
            ("LINEBEFORE", (2, 0), (2, -1), 0.5, colors.black),
            ("LINEBEFORE", (3, 0), (3, -1), 0.5, colors.black),
            ("LINEABOVE", (2, last_r - 1), (3, last_r - 1), 0.5, colors.black),
            ("LINEABOVE", (0, last_r), (3, last_r), 0.5, colors.black),
            ("RIGHTPADDING", (1, 1), (1, -1), 6),
            ("RIGHTPADDING", (3, 1), (3, -1), 6),
        ]))
        elements.append(gu_tbl)
        elements.append(Spacer(1, 2 * mm))
        foot_tbl = Table(
            [[Paragraph("귀하의 노고에 감사드립니다.", meta_left), Paragraph(corp_name or "", meta_right)]],
            colWidths=[87 * mm, 87 * mm],
        )
        foot_tbl.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
        elements.append(foot_tbl)
        elements.append(Spacer(1, 12 * mm))
        elements.append(HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#64748B"), dash=(2, 2)))
    else:
        # ── 양식 1 ~ 5 공통 상단 회사명 / 지급일 ──
        head_tbl = Table(
            [[
                Paragraph(f"회사명 :  {corp_name or ''}", meta_left),
                Paragraph(f"지급일 :   {pay_dt_dot}", meta_right),
            ]],
            colWidths=[87 * mm, 87 * mm],
        )
        head_tbl.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2)]))
        elements.append(head_tbl)
        elements.append(Spacer(1, 1.5 * mm))

        # ── 상단 사원 인적사항 표 (양식별 분기) ──
        if tpl == "근로기준1":
            info_data = [
                [f"사원코드 :    {emp_code}", f"사원명 :    {emp_name}", f"생년월일 :    {emp_birth}"],
                [f"부  서 :    {emp_dept}", f"직  급 :    {emp_pos}", "호  봉 :"],
            ]
            info_tbl = Table(info_data, colWidths=[58 * mm, 58 * mm, 58 * mm], rowHeights=[6.5 * mm, 6.5 * mm])
            info_tbl.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), font_reg),
                ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOX", (0, 0), (-1, -1), 0.6, colors.black),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ]))
            elements.append(info_tbl)
            elements.append(Spacer(1, 3 * mm))
        elif tpl == "근로기준4":
            info_data = [["성명", emp_name, "사번", emp_code]]
            info_tbl = Table(info_data, colWidths=[22 * mm, 65 * mm, 22 * mm, 65 * mm], rowHeights=[7.5 * mm])
            info_tbl.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), font_reg),
                ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                ("ALIGN", (0, 0), (0, 0), "CENTER"),
                ("ALIGN", (2, 0), (2, 0), "CENTER"),
                ("ALIGN", (1, 0), (1, 0), "LEFT"),
                ("ALIGN", (3, 0), (3, 0), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOX", (0, 0), (-1, -1), 0.6, colors.black),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.black),
                ("LEFTPADDING", (1, 0), (1, 0), 6),
                ("LEFTPADDING", (3, 0), (3, 0), 6),
            ]))
            elements.append(info_tbl)
            elements.append(Spacer(1, 2 * mm))
        else:
            # 근로기준2, 근로기준3, 근로기준5 공통 2행×4열 격자표
            birth_and_code = f"{emp_birth}          ({emp_code})" if emp_birth else f"({emp_code})"
            info_data = [
                ["성명", emp_name, "생년월일(사번)", birth_and_code],
                ["부서", emp_dept, "직급", emp_pos],
            ]
            info_tbl = Table(info_data, colWidths=[22 * mm, 58 * mm, 32 * mm, 62 * mm], rowHeights=[7 * mm, 7 * mm])
            info_tbl.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), font_reg),
                ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                ("ALIGN", (0, 0), (0, -1), "CENTER"),
                ("ALIGN", (2, 0), (2, -1), "CENTER"),
                ("ALIGN", (1, 0), (1, -1), "LEFT"),
                ("ALIGN", (3, 0), (3, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOX", (0, 0), (-1, -1), 0.6, colors.black),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.black),
                ("LEFTPADDING", (1, 0), (1, -1), 6),
                ("LEFTPADDING", (3, 0), (3, -1), 6),
            ]))
            elements.append(info_tbl)
            elements.append(Spacer(1, 3 * mm))

        # ── 근로시간 · 통상시급 표 (근로기준1, 근로기준2, 근로기준5) ──
        if tpl in ("근로기준1", "근로기준2"):
            wh_data = [
                ["연장근로시간", "야간근로시간", "휴일근로시간", "통상시급(원)"],
                [
                    f"{ot_h:g}" if ot_h > 0 else "",
                    "",
                    f"{hol_h:g}" if hol_h > 0 else "",
                    f"{hourly_w:,}" if hourly_w > 0 else "",
                ],
            ]
            wh_tbl = Table(wh_data, colWidths=[43.5 * mm] * 4, rowHeights=[6.8 * mm, 6.8 * mm])
            wh_tbl.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), font_reg),
                ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("ALIGN", (3, 1), (3, 1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOX", (0, 0), (-1, -1), 0.6, colors.black),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.black),
                ("RIGHTPADDING", (3, 1), (3, 1), 6),
            ]))
            elements.append(wh_tbl)
            elements.append(Spacer(1, 1.5 * mm))
        elif tpl == "근로기준5":
            work_days = 20
            tot_hours = int(round(float(emp.get("월근무시간", 160) or 160)))
            wh_data = [
                ["근로일수", "총 근무시간", "연장근로시간", "야간근로시간", "휴일근로시간", "통상시급(원)"],
                [
                    str(work_days),
                    str(tot_hours),
                    f"{ot_h:g}" if ot_h > 0 else "",
                    "",
                    f"{hol_h:g}" if hol_h > 0 else "",
                    f"{hourly_w:,}" if hourly_w > 0 else "",
                ],
            ]
            wh_tbl = Table(wh_data, colWidths=[29 * mm] * 6, rowHeights=[6.8 * mm, 6.8 * mm])
            wh_tbl.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), font_reg),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("ALIGN", (5, 1), (5, 1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOX", (0, 0), (-1, -1), 0.6, colors.black),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.black),
                ("RIGHTPADDING", (5, 1), (5, 1), 6),
            ]))
            elements.append(wh_tbl)
            elements.append(Spacer(1, 1.5 * mm))

        elements.append(Paragraph("(단위, 원)", unit_style))
        elements.append(Spacer(1, 1 * mm))

        # ── 메인 세부내역 표 (양식별 4열 vs 5열 분기) ──
        body_rows_cnt = max(12, len(pay_items), len(ded_items))

        if tpl in ("근로기준1", "근로기준4"):
            detail_data: List[List[str]] = []
            if tpl == "근로기준4":
                detail_data.append(["세부내역", "", "", ""])
                detail_data.append(["지급 항목", "지급 금액", "공제 항목", "공제 금액"])
            else:
                detail_data.append(["지 급 내 역", "지 급 액", "공 제 내 역", "공 제 액"])

            hdr_rows = len(detail_data)
            for i in range(body_rows_cnt):
                p_nm, p_v = pay_items[i] if i < len(pay_items) else ("", 0)
                d_nm, d_v = ded_items[i] if i < len(ded_items) else ("", 0)
                detail_data.append([
                    p_nm,
                    _fmt_won_zero(p_v) if p_nm else "",
                    d_nm,
                    _fmt_won_zero(d_v) if d_nm else "",
                ])

            net_label = "실지급액" if tpl == "근로기준4" else "차인지급액"
            detail_data.append(["", "", "공 제 액 계", ded_tot_str])
            detail_data.append(["지 급 액 계", gross_tot_str, net_label, net_tot_str])

            last_r = len(detail_data) - 1
            row_h = [7 * mm] * hdr_rows + [5.8 * mm] * body_rows_cnt + [7 * mm, 7 * mm]
            det_tbl = Table(detail_data, colWidths=[45 * mm, 42 * mm, 45 * mm, 42 * mm], rowHeights=row_h)

            st_cmds: List[Any] = [
                ("FONTNAME", (0, 0), (-1, -1), font_reg),
                ("FONTNAME", (0, 0), (-1, hdr_rows - 1), font_bold),
                ("FONTNAME", (0, last_r - 1), (-1, last_r), font_bold),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("BACKGROUND", (0, 0), (-1, hdr_rows - 1), colors.HexColor("#D9D9D9")),
                ("BACKGROUND", (2, last_r - 1), (3, last_r - 1), colors.HexColor("#D9D9D9")),
                ("BACKGROUND", (0, last_r), (3, last_r), colors.HexColor("#D9D9D9")),
                ("ALIGN", (0, 0), (-1, hdr_rows - 1), "CENTER"),
                ("ALIGN", (0, hdr_rows), (0, -1), "CENTER"),
                ("ALIGN", (2, hdr_rows), (2, -1), "CENTER"),
                ("ALIGN", (1, hdr_rows), (1, -1), "RIGHT"),
                ("ALIGN", (3, hdr_rows), (3, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOX", (0, 0), (-1, -1), 0.7, colors.black),
                ("LINEBELOW", (0, hdr_rows - 1), (-1, hdr_rows - 1), 0.6, colors.black),
                ("LINEBEFORE", (1, 0), (1, -1), 0.5, colors.black),
                ("LINEBEFORE", (2, 0), (2, -1), 0.5, colors.black),
                ("LINEBEFORE", (3, 0), (3, -1), 0.5, colors.black),
                ("LINEABOVE", (2, last_r - 1), (3, last_r - 1), 0.5, colors.black),
                ("LINEABOVE", (0, last_r), (3, last_r), 0.5, colors.black),
                ("RIGHTPADDING", (1, hdr_rows), (1, -1), 6),
                ("RIGHTPADDING", (3, hdr_rows), (3, -1), 6),
            ]
            if tpl == "근로기준4":
                st_cmds.append(("SPAN", (0, 0), (3, 0)))
                st_cmds.append(("LINEBELOW", (0, 0), (3, 0), 0.5, colors.black))
            det_tbl.setStyle(TableStyle(st_cmds))
            elements.append(det_tbl)
        else:
            # 근로기준2, 근로기준3, 근로기준5 공통 5열 ('구분[매월] | 지급 항목 | 지급 금액 | 공제 항목 | 공제 금액')
            detail_data5: List[List[str]] = [
                ["세부내역", "", "", "", ""],
                ["구분", "지급 항목", "지급 금액", "공제 항목", "공제 금액"],
            ]
            for i in range(body_rows_cnt):
                p_nm, p_v = pay_items[i] if i < len(pay_items) else ("", 0)
                d_nm, d_v = ded_items[i] if i < len(ded_items) else ("", 0)
                detail_data5.append([
                    "매월" if p_nm else "",
                    p_nm,
                    _fmt_won_zero(p_v) if p_nm else "",
                    d_nm,
                    _fmt_won_zero(d_v) if d_nm else "",
                ])
            detail_data5.append(["", "", "", "공 제 액 계", ded_tot_str])
            detail_data5.append(["지 급 액 계", "", gross_tot_str, "실지급액", net_tot_str])

            last_r = len(detail_data5) - 1
            row_h = [6.8 * mm, 6.8 * mm] + [5.6 * mm] * body_rows_cnt + [6.8 * mm, 6.8 * mm]
            det_tbl = Table(
                detail_data5,
                colWidths=[18 * mm, 36 * mm, 36 * mm, 44 * mm, 40 * mm],
                rowHeights=row_h,
            )
            det_tbl.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), font_reg),
                ("FONTNAME", (0, 0), (-1, 1), font_bold),
                ("FONTNAME", (0, last_r - 1), (-1, last_r), font_bold),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("SPAN", (0, 0), (4, 0)),
                ("SPAN", (0, last_r), (1, last_r)),
                ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#D9D9D9")),
                ("BACKGROUND", (3, last_r - 1), (4, last_r - 1), colors.HexColor("#D9D9D9")),
                ("BACKGROUND", (0, last_r), (4, last_r), colors.HexColor("#D9D9D9")),
                ("ALIGN", (0, 0), (-1, 1), "CENTER"),
                ("ALIGN", (0, 2), (1, -1), "CENTER"),
                ("ALIGN", (3, 2), (3, -1), "CENTER"),
                ("ALIGN", (2, 2), (2, -1), "RIGHT"),
                ("ALIGN", (4, 2), (4, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOX", (0, 0), (-1, -1), 0.7, colors.black),
                ("LINEBELOW", (0, 0), (4, 0), 0.5, colors.black),
                ("LINEBELOW", (0, 1), (4, 1), 0.6, colors.black),
                ("LINEBEFORE", (1, 1), (1, last_r - 1), 0.5, colors.black),
                ("LINEBEFORE", (2, 1), (2, -1), 0.5, colors.black),
                ("LINEBEFORE", (3, 1), (3, -1), 0.5, colors.black),
                ("LINEBEFORE", (4, 1), (4, -1), 0.5, colors.black),
                ("LINEABOVE", (3, last_r - 1), (4, last_r - 1), 0.5, colors.black),
                ("LINEABOVE", (0, last_r), (4, last_r), 0.5, colors.black),
                ("RIGHTPADDING", (2, 2), (2, -1), 6),
                ("RIGHTPADDING", (4, 2), (4, -1), 6),
            ]))
            elements.append(det_tbl)

        elements.append(Spacer(1, 1.5 * mm))
        elements.append(Paragraph("(단위, 원)", unit_style))
        elements.append(Spacer(1, 1 * mm))
        elements.append(_build_calc_method_table(emp, font_reg, font_bold))
        elements.append(Spacer(1, 4 * mm))
        elements.append(Paragraph("귀하의 노고에 감사드립니다.", footer_style))

    doc.build(elements)
    raw_pdf = buf.getvalue()

    if password:
        reader = PdfReader(io.BytesIO(raw_pdf))
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        writer.encrypt(password)
        enc_buf = io.BytesIO()
        writer.write(enc_buf)
        return enc_buf.getvalue()

    return raw_pdf


def convert_pdf_bytes_to_jpg_bytes(pdf_bytes: bytes, scale: float = 2.2) -> bytes:
    """단일 페이지 PDF 바이트를 고해상도 JPG 바이트로 변환 (pypdfium2 + Pillow)"""
    pdf_doc = pdfium.PdfDocument(pdf_bytes)
    page = pdf_doc[0]
    bitmap = page.render(scale=scale)
    pil_img: Image.Image = bitmap.to_pil()
    if pil_img.mode != "RGB":
        pil_img = pil_img.convert("RGB")
    img_buf = io.BytesIO()
    pil_img.save(img_buf, format="JPEG", quality=92)
    return img_buf.getvalue()


def export_employee_payslips_zip(
    employees: List[Dict[str, Any]],
    corp_name: str = "",
    year_month: str = "",
    pay_date: str = "",
    encrypt_with_birth: bool = False,
    template_type: str = "근로기준1",
    start_index: int = 0,
    max_count: int = 200,
) -> bytes:
    """
    인별 급여명세서 PDF 일괄 생성 후 단일 ZIP 파일로 압축 반환 (1회 최대 200명):
    - ZIP 내 개별 파일명 규칙: `근로자명_급여명세서_YYMM.pdf` (동명이인 존재 시 `근로자명(사번)_급여명세서_YYMM.pdf`)
    """
    yyyy, mm_num = _parse_ym_parts(year_month)
    yymm = f"{str(yyyy)[2:]}{mm_num:02d}"
    sliced = (employees or [])[start_index : start_index + max_count]

    seen_names: Dict[str, int] = {}
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        for offset, emp in enumerate(sliced, start=start_index + 1):
            name = str(emp.get("성명") or f"사원{offset}").strip()
            safe_name = re.sub(r'[\\/*?:"<>|]', "", name)
            seen_names[safe_name] = seen_names.get(safe_name, 0) + 1
            if seen_names[safe_name] > 1:
                code_sfx = str(emp.get("사원코드") or offset)
                name_tag = f"{safe_name}({code_sfx})"
            else:
                name_tag = safe_name

            pwd: Optional[str] = None
            if encrypt_with_birth:
                birth6 = extract_birth_6digits(str(emp.get("주민등록번호", "")))
                if len(birth6) == 6:
                    pwd = birth6

            pdf_bytes = render_single_employee_payslip_pdf(
                emp=emp,
                corp_name=corp_name,
                year_month=year_month,
                pay_date=pay_date,
                password=pwd,
                template_type=template_type,
            )
            filename = f"{name_tag}_급여명세서_{yymm}.pdf"
            zf.writestr(filename, pdf_bytes)

    return zip_buf.getvalue()


def export_employee_payslips_jpg_zip(
    employees: List[Dict[str, Any]],
    corp_name: str = "",
    year_month: str = "",
    pay_date: str = "",
    template_type: str = "근로기준1",
    start_index: int = 0,
    max_count: int = 200,
) -> bytes:
    """
    인별 급여명세서 이미지(JPG) 일괄 생성 후 단일 ZIP 파일로 압축 반환 (1회 최대 200명):
    - ZIP 내 개별 파일명 규칙: `근로자명_급여명세서_YYMM.jpg`
    """
    yyyy, mm_num = _parse_ym_parts(year_month)
    yymm = f"{str(yyyy)[2:]}{mm_num:02d}"
    sliced = (employees or [])[start_index : start_index + max_count]

    seen_names: Dict[str, int] = {}
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        for offset, emp in enumerate(sliced, start=start_index + 1):
            name = str(emp.get("성명") or f"사원{offset}").strip()
            safe_name = re.sub(r'[\\/*?:"<>|]', "", name)
            seen_names[safe_name] = seen_names.get(safe_name, 0) + 1
            if seen_names[safe_name] > 1:
                code_sfx = str(emp.get("사원코드") or offset)
                name_tag = f"{safe_name}({code_sfx})"
            else:
                name_tag = safe_name

            pdf_bytes = render_single_employee_payslip_pdf(
                emp=emp,
                corp_name=corp_name,
                year_month=year_month,
                pay_date=pay_date,
                password=None,
                template_type=template_type,
            )
            jpg_bytes = convert_pdf_bytes_to_jpg_bytes(pdf_bytes)
            filename = f"{name_tag}_급여명세서_{yymm}.jpg"
            zf.writestr(filename, jpg_bytes)

    return zip_buf.getvalue()


# ─────────────────────────────────────────────────────────────
# 5. 사업소득지급대장 PDF (A4 세로 - 표준 서식)
# ─────────────────────────────────────────────────────────────

def export_freelancer_ledger_pdf(
    freelancers: List[Dict[str, Any]],
    corp_name: str = "",
    year_month: str = "",
) -> bytes:
    font_reg, font_bold = get_korean_font_names()
    yyyy, mm_num = _parse_ym_parts(year_month)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
    )

    title_style = ParagraphStyle("FreeLedgerTitle", fontName=font_bold, fontSize=15, leading=20, alignment=1)
    meta_style = ParagraphStyle("FreeLedgerMeta", fontName=font_reg, fontSize=9.5, leading=13, alignment=0)

    elements: List[Any] = []
    elements.append(Paragraph(f"({yyyy}년 {mm_num:02d}월) 사업소득지급대장(합계)", title_style))
    elements.append(Spacer(1, 4 * mm))
    elements.append(Paragraph(f"회사명 : {corp_name or '미입력'}", meta_style))
    elements.append(Spacer(1, 2 * mm))

    table_data: List[List[str]] = [
        ["NO", "코드", "성명", "귀속년월", "지급액", "소득세", "예술·특고인경비", "학자금상환액", "차인지급액", "영수인"],
        ["", "", "", "지급년월", "", "지방소득세", "고용보험료", "산재보험료", "", ""],
    ]

    tot_gross = 0
    tot_itax = 0
    tot_ltax = 0
    tot_net = 0

    for idx, fr in enumerate(freelancers or [], start=1):
        gross = int(fr.get("지급총액", 0) or 0)
        itax = int(fr.get("소득세", 0) or 0)
        ltax = int(fr.get("지방소득세", 0) or 0)
        net = int(fr.get("차인지급액", 0) or 0)
        tot_gross += gross
        tot_itax += itax
        tot_ltax += ltax
        tot_net += net

        ym_str = str(fr.get("귀속년월") or f"{yyyy}-{mm_num:02d}").replace("-", ".")
        pay_str = str(fr.get("지급년월일") or f"{yyyy}-{mm_num:02d}")[:7].replace("-", ".")

        r1 = [
            str(idx),
            str(fr.get("코드") or f"{idx:04d}"),
            str(fr.get("성명", "")),
            ym_str,
            _fmt_won_zero(gross),
            _fmt_won_zero(itax),
            "",
            "",
            _fmt_won_zero(net),
            "",
        ]
        r2 = [
            "",
            "",
            "",
            pay_str,
            "",
            _fmt_won_zero(ltax),
            "",
            "",
            "",
            "",
        ]
        table_data.extend([r1, r2])

    table_data.append([
        "총 계", "", "", "",
        _fmt_won_zero(tot_gross),
        _fmt_won_zero(tot_itax),
        "", "",
        _fmt_won_zero(tot_net),
        "",
    ])
    table_data.append([
        "", "", "", "",
        "",
        _fmt_won_zero(tot_ltax),
        "", "",
        "",
        "",
    ])

    col_widths = [10 * mm, 16 * mm, 22 * mm, 20 * mm, 26 * mm, 22 * mm, 22 * mm, 20 * mm, 26 * mm, 12 * mm]
    tbl = Table(table_data, colWidths=col_widths, repeatRows=2)

    cmds: List[Any] = [
        ("FONTNAME", (0, 0), (-1, -1), font_reg),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (3, -1), "CENTER"),
        ("ALIGN", (4, 2), (8, -1), "RIGHT"),
        ("ALIGN", (4, 0), (9, 1), "CENTER"),
        ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#EDE9FE")),
        ("FONTNAME", (0, 0), (-1, 1), font_bold),
        ("BOX", (0, 0), (-1, -1), 0.9, colors.HexColor("#334155")),
        ("INNERGRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
    ]

    total_r = len(table_data)
    for r_start in range(0, total_r, 2):
        cmds.append(("LINEABOVE", (0, r_start), (-1, r_start), 0.7, colors.HexColor("#475569")))
        for c_span in (0, 1, 2, 4, 8, 9):
            cmds.append(("SPAN", (c_span, r_start), (c_span, r_start + 1)))

    cmds.append(("SPAN", (0, total_r - 2), (3, total_r - 1)))
    cmds.append(("BACKGROUND", (0, total_r - 2), (-1, total_r - 1), colors.HexColor("#F8FAFC")))
    cmds.append(("FONTNAME", (0, total_r - 2), (-1, total_r - 1), font_bold))

    tbl.setStyle(TableStyle(cmds))
    elements.append(tbl)

    doc.build(elements)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────
# 6. 인별 사업소득명세서 PDF 및 JPG 일괄 압축 (1회 최대 100명, 세로형 심플 양식)
# ─────────────────────────────────────────────────────────────

def render_single_freelancer_statement_pdf(
    fr: Dict[str, Any],
    corp_name: str = "",
    year_month: str = "",
    pay_date: str = "",
) -> bytes:
    font_reg, font_bold = get_korean_font_names()
    yyyy, mm_num = _parse_ym_parts(fr.get("귀속년월") or year_month)
    ym_kor = f"{yyyy}년 {mm_num:02d}월"
    pay_dt_str = str(fr.get("지급년월일") or pay_date or f"{yyyy}-{mm_num:02d}-25")

    gross = int(fr.get("지급총액", 0) or 0)
    itax = int(fr.get("소득세", 0) or 0)
    ltax = int(fr.get("지방소득세", 0) or 0)
    ded_total = int(fr.get("공제액계", itax + ltax) or (itax + ltax))
    net = int(fr.get("차인지급액", gross - ded_total) or (gross - ded_total))

    raw_ind = str(fr.get("소득구분") or "940909")
    m_ind = re.search(r"\b(\d{6})\b", raw_ind)
    clean_ind = m_ind.group(1) if m_ind else "940909"

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=22 * mm,
        rightMargin=22 * mm,
        topMargin=24 * mm,
        bottomMargin=24 * mm,
    )

    title_style = ParagraphStyle("BizSlipTitle", fontName=font_bold, fontSize=20, leading=26, alignment=1)
    sub_style = ParagraphStyle("BizSlipSub", fontName=font_reg, fontSize=10, leading=14, alignment=1, textColor=colors.HexColor("#475569"))
    sec_style = ParagraphStyle("BizSlipSec", fontName=font_bold, fontSize=11, leading=16, alignment=0)
    foot_style = ParagraphStyle("BizSlipFoot", fontName=font_bold, fontSize=11, leading=16, alignment=1)

    elements: List[Any] = []
    elements.append(Paragraph("사 업 소 득 명 세 서", title_style))
    elements.append(Spacer(1, 2 * mm))
    elements.append(Paragraph(f"원천징수의무자(회사명) : {corp_name or '미입력'}", sub_style))
    elements.append(Spacer(1, 8 * mm))

    info_data = [
        ["귀속연월", ym_kor, "지급연월(일자)", pay_dt_str],
        ["근로자명(성명)", str(fr.get("성명", "")), "소득구분", f"사업소득 (3.3% / {clean_ind})"],
    ]
    info_tbl = Table(
        info_data,
        colWidths=[32 * mm, 45 * mm, 32 * mm, 57 * mm],
        rowHeights=[10 * mm, 10 * mm],
    )
    info_tbl.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_reg),
        ("FONTNAME", (0, 0), (0, -1), font_bold),
        ("FONTNAME", (2, 0), (2, -1), font_bold),
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F3E8FF")),
        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F3E8FF")),
        ("BOX", (0, 0), (-1, -1), 0.9, colors.HexColor("#4C1D95")),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C4B5FD")),
    ]))
    elements.append(info_tbl)
    elements.append(Spacer(1, 8 * mm))

    elements.append(Paragraph("지급 및 원천징수 공제 내역", sec_style))
    elements.append(Spacer(1, 2 * mm))

    itax_note = "세전금액 × 3.0% (10원 미만 절사)"
    amt_data = [
        ["구 분", "산출 기준", "금 액 (원)"],
        ["세전금액 (지급총액)", "약정 사업소득 지급총액", f"{gross:,} 원"],
        ["공제가액 - 소득세", itax_note, f"{itax:,} 원"],
        ["공제가액 - 지방소득세", "소득세 × 10% (10원 미만 절사)", f"{ltax:,} 원"],
        ["공제가액 합계", "총 3.3% 원천징수 세액", f"{ded_total:,} 원"],
        ["차인지급액 (실지급액)", "세전금액 - 공제가액 합계", f"{net:,} 원"],
    ]
    amt_tbl = Table(
        amt_data,
        colWidths=[52 * mm, 64 * mm, 50 * mm],
        rowHeights=[10 * mm, 11 * mm, 10 * mm, 10 * mm, 11 * mm, 13 * mm],
    )
    amt_tbl.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_reg),
        ("FONTNAME", (0, 0), (-1, 0), font_bold),
        ("FONTNAME", (0, 1), (0, -1), font_bold),
        ("FONTNAME", (0, 4), (-1, 5), font_bold),
        ("FONTSIZE", (0, 0), (-1, 4), 10),
        ("FONTSIZE", (0, 5), (-1, 5), 11.5),
        ("ALIGN", (0, 0), (1, -1), "CENTER"),
        ("ALIGN", (2, 0), (2, 0), "CENTER"),
        ("ALIGN", (2, 1), (2, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3E8FF")),
        ("BACKGROUND", (0, 4), (-1, 4), colors.HexColor("#F8FAFC")),
        ("BACKGROUND", (0, 5), (-1, 5), colors.HexColor("#EDE9FE")),
        ("BOX", (0, 0), (-1, -1), 1.0, colors.HexColor("#4C1D95")),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("RIGHTPADDING", (2, 1), (2, -1), 10),
    ]))
    elements.append(amt_tbl)
    elements.append(Spacer(1, 14 * mm))
    elements.append(Paragraph("위와 같이 사업소득 금액 및 원천징수 세액을 명세합니다.", foot_style))

    doc.build(elements)
    return buf.getvalue()


def export_freelancer_statements_pdf_zip(
    freelancers: List[Dict[str, Any]],
    corp_name: str = "",
    year_month: str = "",
    pay_date: str = "",
    start_index: int = 0,
    max_count: int = 100,
) -> bytes:
    """
    인별 사업소득명세서 PDF 일괄 생성 후 단일 ZIP 파일로 압축 반환 (1회 최대 100명):
    - ZIP 내 개별 파일명 규칙: `근로자명_사업소득명세서_YYMM.pdf`
    """
    yyyy, mm_num = _parse_ym_parts(year_month)
    yymm = f"{str(yyyy)[2:]}{mm_num:02d}"
    sliced = (freelancers or [])[start_index : start_index + max_count]

    seen_names: Dict[str, int] = {}
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        for offset, fr in enumerate(sliced, start=start_index + 1):
            name = str(fr.get("성명") or f"소득자{offset}").strip()
            safe_name = re.sub(r'[\\/*?:"<>|]', "", name)
            seen_names[safe_name] = seen_names.get(safe_name, 0) + 1
            if seen_names[safe_name] > 1:
                code_sfx = str(fr.get("코드") or offset)
                name_tag = f"{safe_name}({code_sfx})"
            else:
                name_tag = safe_name

            pdf_bytes = render_single_freelancer_statement_pdf(
                fr=fr,
                corp_name=corp_name,
                year_month=year_month,
                pay_date=pay_date,
            )
            zf.writestr(f"{name_tag}_사업소득명세서_{yymm}.pdf", pdf_bytes)
    return zip_buf.getvalue()


def export_freelancer_statements_jpg_zip(
    freelancers: List[Dict[str, Any]],
    corp_name: str = "",
    year_month: str = "",
    pay_date: str = "",
    start_index: int = 0,
    max_count: int = 100,
) -> bytes:
    """
    인별 사업소득명세서 이미지(JPG) 일괄 생성 후 단일 ZIP 파일로 압축 반환 (1회 최대 100명):
    - ZIP 내 개별 파일명 규칙: `근로자명_사업소득명세서_YYMM.jpg`
    """
    yyyy, mm_num = _parse_ym_parts(year_month)
    yymm = f"{str(yyyy)[2:]}{mm_num:02d}"
    sliced = (freelancers or [])[start_index : start_index + max_count]

    seen_names: Dict[str, int] = {}
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        for offset, fr in enumerate(sliced, start=start_index + 1):
            name = str(fr.get("성명") or f"소득자{offset}").strip()
            safe_name = re.sub(r'[\\/*?:"<>|]', "", name)
            seen_names[safe_name] = seen_names.get(safe_name, 0) + 1
            if seen_names[safe_name] > 1:
                code_sfx = str(fr.get("코드") or offset)
                name_tag = f"{safe_name}({code_sfx})"
            else:
                name_tag = safe_name

            pdf_bytes = render_single_freelancer_statement_pdf(
                fr=fr,
                corp_name=corp_name,
                year_month=year_month,
                pay_date=pay_date,
            )
            jpg_bytes = convert_pdf_bytes_to_jpg_bytes(pdf_bytes)
            zf.writestr(f"{name_tag}_사업소득명세서_{yymm}.jpg", jpg_bytes)
    return zip_buf.getvalue()


# ─────────────────────────────────────────────────────────────
# 7. 마우스 오버(Hover) 돋보기 미리보기용 Base64 이미지 캐시 생성기
# ─────────────────────────────────────────────────────────────

def get_payslip_template_preview_b64(template_type: str = "근로기준1") -> str:
    """
    급여명세서 6종 양식 각각의 샘플 미리보기 JPG 이미지를 생성하여 Base64 문자열로 캐싱 반환
    (실제 개인정보 없이 가상 샘플 데이터 사용)
    """
    cache_key = f"emp_slip_{template_type}"
    if cache_key in _PREVIEW_B64_CACHE:
        return _PREVIEW_B64_CACHE[cache_key]

    sample_emp = {
        "사원코드": "001",
        "성명": "홍길동",
        "주민등록번호": "950101-1******",
        "입사일": "2025-01-01",
        "부서": "경영지원팀",
        "직급": "사원",
        "기본급": 2_620_000,
        "식대": 200_000,
        "연장근로수당": 18_000,
        "지급액계": 2_838_000,
        "국민연금": 124_450,
        "건강보험": 94_180,
        "장기요양보험료": 12_370,
        "고용보험": 23_740,
        "소득세": 42_260,
        "지방소득세": 4_220,
        "공제액계": 301_220,
        "차인지급액": 2_536_780,
        "약정시급": 12_536,
        "연장근무시간": 1.0,
        "월근무시간": 160.0,
    }
    pdf_b = render_single_employee_payslip_pdf(
        emp=sample_emp,
        corp_name="주식회사 샘플기업",
        year_month="2026-10",
        pay_date="2026.10.25",
        template_type=template_type,
    )
    jpg_b = convert_pdf_bytes_to_jpg_bytes(pdf_b, scale=1.35)
    b64_str = base64.b64encode(jpg_b).decode("ascii")
    _PREVIEW_B64_CACHE[cache_key] = b64_str
    return b64_str


def get_freelancer_statement_preview_b64() -> str:
    """
    사업소득명세서 세로형 양식의 샘플 미리보기 JPG 이미지를 생성하여 Base64 문자열로 캐싱 반환
    """
    cache_key = "free_slip_default"
    if cache_key in _PREVIEW_B64_CACHE:
        return _PREVIEW_B64_CACHE[cache_key]

    sample_fr = {
        "코드": "0001",
        "성명": "홍길동",
        "소득구분": "940909",
        "귀속년월": "2026-10",
        "지급년월일": "2026-10-25",
        "지급총액": 2_000_000,
        "소득세": 60_000,
        "지방소득세": 6_000,
        "공제액계": 66_000,
        "차인지급액": 1_934_000,
    }
    pdf_b = render_single_freelancer_statement_pdf(
        fr=sample_fr,
        corp_name="주식회사 샘플기업",
        year_month="2026-10",
        pay_date="2026-10-25",
    )
    jpg_b = convert_pdf_bytes_to_jpg_bytes(pdf_b, scale=1.35)
    b64_str = base64.b64encode(jpg_b).decode("ascii")
    _PREVIEW_B64_CACHE[cache_key] = b64_str
    return b64_str


# ─────────────────────────────────────────────────────────────
# 당월 검토보고 PDF (급여 / 사업소득 공통 — 화면 보고서와 같은 내용)
# ─────────────────────────────────────────────────────────────

def export_review_report_pdf(report: Dict[str, Any]) -> bytes:
    """화면의 당월 검토보고(제목·섹션·표·설명·목록)를 A4 가로 PDF로 만듭니다."""
    from xml.sax.saxutils import escape as _esc
    font_reg, font_bold = get_korean_font_names()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                            title=str(report.get("title") or "검토보고"))
    accent = colors.HexColor(str(report.get("accent") or "#1e3a8a"))
    st_title = ParagraphStyle("rt", fontName=font_bold, fontSize=15, leading=20, spaceAfter=8, textColor=colors.HexColor("#0f172a"))
    st_h = ParagraphStyle("rh", fontName=font_bold, fontSize=11.5, leading=16, spaceBefore=10, spaceAfter=5, textColor=colors.HexColor("#0f172a"),
                          borderPadding=(3, 6, 3, 6), backColor=colors.HexColor("#f1f5f9"), borderColor=accent, borderWidth=0, leftIndent=0)
    st_p = ParagraphStyle("rp", fontName=font_reg, fontSize=9.5, leading=14, textColor=colors.HexColor("#1e293b"))
    st_note = ParagraphStyle("rn", parent=st_p, fontSize=8.5, leading=12, textColor=colors.HexColor("#475569"))
    st_cell = ParagraphStyle("rc", fontName=font_reg, fontSize=9, leading=12, alignment=2)
    st_cell_c = ParagraphStyle("rcc", parent=st_cell, alignment=1)
    st_head = ParagraphStyle("rhd", fontName=font_bold, fontSize=9, leading=12, alignment=1, textColor=colors.HexColor("#0f172a"))
    avail = landscape(A4)[0] - 28 * mm
    flow = [Paragraph(_esc(str(report.get("title") or "")), st_title)]
    for sec in report.get("sections", []):
        flow.append(Paragraph(_esc(str(sec.get("heading") or "")), st_h))
        if sec.get("intro"):
            flow.append(Paragraph(_esc(str(sec["intro"])), st_p))
            flow.append(Spacer(1, 4))
        if sec.get("headers"):
            widths = [avail * float(w) / 100.0 for w in sec.get("widths") or [100.0 / len(sec["headers"])] * len(sec["headers"])]
            data = [[Paragraph(_esc(str(h)), st_head) for h in sec["headers"]]]
            st_left = ParagraphStyle("rcl", parent=st_cell, alignment=0)
            for r in sec.get("rows", []):
                data.append([Paragraph(_esc(str(v)), st_left if sec.get("align") == "left" else (st_cell if str(v).endswith("원") else st_cell_c)) for v in r])
            t = Table(data, colWidths=widths, repeatRows=1)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#cbd5e1")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            flow.append(t)
            flow.append(Spacer(1, 4))
        for n in sec.get("notes", []):
            flow.append(Paragraph("* " + _esc(str(n)), st_note))
        for k, v in sec.get("bullets", []):
            flow.append(Paragraph(f"• <b>{_esc(str(k))}:</b> {_esc(str(v))}", st_p))
    doc.build(flow)
    return buf.getvalue()
